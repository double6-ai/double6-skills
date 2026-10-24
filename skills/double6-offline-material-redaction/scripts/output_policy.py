"""CLI与网页共用的原格式默认值、显式转换和实际副本复检。"""
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import engine
import formats
import compat
import image_native

LEGACY = {'.doc':'.docx','.rtf':'.docx','.xls':'.xlsx','.ppt':'.pdf'}
MODERN = {'.odt':'.docx','.ods':'.xlsx','.odp':'.pptx','.xlsm':'.xlsx','.markdown':'.md'}


def category(ext):
    return 'image' if ext in image_native.EXTENSIONS else 'word' if ext=='.docx' else compat.KINDS.get(ext) or formats.EXTENSIONS.get(ext)


def options(ext):
    alternatives=list(image_native.EXTENSIONS)if ext in image_native.EXTENSIONS else [LEGACY.get(ext) or MODERN.get(ext)]
    return [ext,*[x for x in alternatives if x and x!=ext]]


def target(ext, requested=None):
    selected=ext if requested in (None,'','same') else requested.lower()
    if selected not in options(ext):raise engine.Rejected('未支持这一输出格式，不能仅修改扩展名。')
    if ext in LEGACY and selected==ext:
        raise engine.Rejected('该旧格式暂不能安全导出原格式副本；已停止。只有明确指定输出为'+LEGACY[ext]+'才会转换。')
    return selected


def mime(ext):
    if ext in image_native.EXTENSIONS:return 'image/'+{'PNG':'png','JPEG':'jpeg','WEBP':'webp'}[image_native.EXTENSIONS[ext]]
    if ext in compat.ODF:return compat.ODF[ext]
    if ext=='.xlsm':return 'application/vnd.ms-excel.sheet.macroEnabled.12'
    return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' if ext=='.docx' else formats.MIME[category(ext)]


def native_odf(data, ext):
    """仅从已经清理的现代格式构造最小ODF容器，绝不复制原ODF部件。"""
    O,T,A,D,S=compat.O,compat.T,compat.A,compat.D,compat.S
    svg='{urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0}'
    root=ET.Element(O+'document-content',{O+'version':'1.2'})
    styles=ET.Element(O+'document-styles',{O+'version':'1.2'})
    ET.SubElement(styles,O+'styles')
    auto=ET.SubElement(styles,O+'automatic-styles')
    masters=ET.SubElement(styles,O+'master-styles')
    content_styles=ET.SubElement(root,O+'automatic-styles')
    body=ET.SubElement(root,O+'body')
    body=ET.SubElement(body,O+{'.odt':'text','.ods':'spreadsheet','.odp':'presentation'}[ext])
    def paragraph(parent,text):
        node=ET.SubElement(parent,T+'p');last=None
        for chunk in re.split(r'([\t\n]| {2,})',str(text)):
            if not chunk:continue
            if chunk=='\n':last=ET.SubElement(node,T+'line-break')
            elif chunk=='\t':last=ET.SubElement(node,T+'tab')
            elif len(chunk)>1 and not chunk.strip():last=ET.SubElement(node,T+'s',{T+'c':str(len(chunk))})
            elif last is None:node.text=(node.text or '')+chunk
            else:last.tail=(last.tail or '')+chunk
    def table(parent,rows):
        tab=ET.SubElement(parent,A+'table',{A+'name':'业务表'})
        for row in rows:
            tr=ET.SubElement(tab,A+'table-row')
            for value in row:
                attrs={O+'value-type':'string'}
                if isinstance(value,(int,float)) and not isinstance(value,bool):attrs={O+'value-type':'float',O+'value':str(value)}
                cell=ET.SubElement(tr,A+'table-cell',attrs);paragraph(cell,'' if value is None else value)
    if ext=='.odt':
        from docx import Document
        doc=Document(io.BytesIO(data))
        for block in engine.inspect_docx(data)['blocks']:
            if block['kind']=='paragraph':paragraph(body,block['text'])
            else:table(body,block['rows'])
    elif ext=='.ods':
        model,_,_=formats.excel_model(data,include_formats=True)
        cell_style=ET.SubElement(content_styles,S+'style',{S+'name':'Wrap',S+'family':'table-cell'})
        ET.SubElement(cell_style,S+'table-cell-properties',{compat.F+'wrap-option':'wrap',S+'vertical-align':'top'})
        row_style=ET.SubElement(content_styles,S+'style',{S+'name':'AutoRow',S+'family':'table-row'})
        ET.SubElement(row_style,S+'table-row-properties',{S+'use-optimal-row-height':'true'})
        for n,cells in enumerate(model):
            tab=ET.SubElement(body,A+'table',{A+'name':'表'+str(n+1)})
            values={(r,c):formats.cell_display.display(v,f) if f not in {'General','@'} else v for r,c,v,f in cells}
            for c in range(1,max(c for r,c in values)+1):
                name='Col'+str(n)+'_'+str(c)
                length=max((len(str(v))for (r,cc),v in values.items()if cc==c),default=1)
                col_style=ET.SubElement(content_styles,S+'style',{S+'name':name,S+'family':'table-column'})
                ET.SubElement(col_style,S+'table-column-properties',{S+'column-width':str(min(12,max(4,length*.22)))+'cm'})
                ET.SubElement(tab,A+'table-column',{A+'style-name':name,A+'default-cell-style-name':'Wrap'})
            for r in range(1,max(r for r,c in values)+1):
                tr=ET.SubElement(tab,A+'table-row',{A+'style-name':'AutoRow'})
                for c in range(1,max(c for rr,c in values)+1):
                    value=values.get((r,c),'');attrs={O+'value-type':'string'}
                    if isinstance(value,(int,float))and not isinstance(value,bool):attrs={O+'value-type':'float',O+'value':str(value)}
                    cell=ET.SubElement(tr,A+'table-cell',attrs);paragraph(cell,value)
    else:
        from pptx import Presentation
        prs=Presentation(io.BytesIO(data))
        layout=ET.SubElement(auto,S+'page-layout',{S+'name':'PM1'})
        ET.SubElement(layout,S+'page-layout-properties',{compat.F+'page-width':'33.867cm',compat.F+'page-height':'19.05cm',S+'print-orientation':'landscape'})
        ET.SubElement(masters,S+'master-page',{S+'name':'Default',S+'page-layout-name':'PM1'})
        for n,slide in enumerate(prs.slides):
            page=ET.SubElement(body,D+'page',{D+'name':'第'+str(n+1)+'页',D+'master-page-name':'Default'})
            for shape in slide.shapes:
                if not (shape.has_text_frame or shape.has_table):continue
                cm=lambda value:str(round(value/360000,4))+'cm'
                frame=ET.SubElement(page,D+'frame',{svg+'x':cm(shape.left),svg+'y':cm(shape.top),svg+'width':cm(shape.width),svg+'height':cm(shape.height)})
                if shape.has_table:table(frame,[[c.text for c in row.cells]for row in shape.table.rows])
                else:
                    box=ET.SubElement(frame,D+'text-box')
                    for p in shape.text_frame.paragraphs:paragraph(box,p.text)
    manifest_ns='{urn:oasis:names:tc:opendocument:xmlns:manifest:1.0}'
    manifest=ET.Element(manifest_ns+'manifest',{manifest_ns+'version':'1.2'})
    for name,media in [('/',compat.ODF[ext]),('content.xml','text/xml'),('styles.xml','text/xml')]:
        ET.SubElement(manifest,manifest_ns+'file-entry',{manifest_ns+'full-path':name,manifest_ns+'media-type':media})
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED)as z:
        z.writestr('mimetype',compat.ODF[ext],compress_type=zipfile.ZIP_STORED)
        for name,node in [('content.xml',root),('styles.xml',styles),('META-INF/manifest.xml',manifest)]:z.writestr(name,ET.tostring(node,encoding='utf-8',xml_declaration=True))
    return out.getvalue()


def xlsm_container(data):
    parts=formats.package(data)
    parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED)as z:
        for name,value in parts.items():z.writestr(name,value)
    return out.getvalue()


def verify(data, ext, plan, source=None):
    kind=category(ext)
    if ext in {'.xlsm','.xlsx'}:
        content=formats.package(data)['[Content_Types].xml']
        expected=b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' if ext=='.xlsm' else b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'
        if expected not in content:raise engine.Rejected('Excel实际容器类型与扩展名不符。')
    if kind=='image':return image_native.verify(data,plan.get('rectangles'),source,image_native.EXTENSIONS[ext])
    if ext in compat.ODF:
        if set(formats.package(data))!={'mimetype','content.xml','styles.xml','META-INF/manifest.xml'}:raise engine.Rejected('ODF副本含非本工具的额外载体。')
        normalized,detail=compat.normalize(data,ext)
        info=engine.inspect_docx(normalized) if kind=='word' else formats.inspect(normalized,kind)
        preview=info['preview'];terms=engine.terms_checked(plan.get('terms',[]));keep=engine.keep_checked(plan.get('keep',[]))
        raw_text=' '.join(formats.package(data)[n].decode('utf-8') for n in ('content.xml','styles.xml'))
        checks={'selected_strings_absent':not any(engine.term_remains(raw_text, x) for x in terms),'odf_clean_container':not detail.get('hidden_nodes_removed'),'odf_reopened':True}
        source_preview=None
        if source is not None:
            try:
                source_preview=(engine.inspect_docx(source) if kind=='word' else formats.inspect(source,kind))['preview']
            except engine.Rejected:
                source_preview=None
        if keep:checks['keep_phrases_preserved']=engine.phrases_kept(source_preview,preview,keep)
        if not all(checks.values()):raise engine.Rejected('ODF实际副本隐藏内容或保留项复检失败。')
        return {'checks':checks,'not_checked':[] if keep else ['keep_phrases_preserved'],'output_sha256':engine.digest(data),'scope':'最小ODF原格式副本，实际容器及可见内容已重新读取。'}
    if kind=='word':
        checks=engine.verify_docx(data,engine.terms_checked(plan.get('terms',[])))
        keep=engine.keep_checked(plan.get('keep',[]))
        source_preview=engine.inspect_docx(source)['preview'] if source is not None else None
        if keep:checks['keep_phrases_preserved']=engine.phrases_kept(source_preview,engine.inspect_docx(data)['preview'],keep)
        if not all(checks.values()):raise engine.Rejected('Word副本复检未通过。')
        return {'checks':checks,'not_checked':[] if keep else ['keep_phrases_preserved'],'output_sha256':engine.digest(data),'scope':'已选原文、支持的隐藏载体与声明保留项复检。'}
    return formats.verify(data,kind,plan,source)


def export(raw, ext, plan, requested=None):
    output_ext=target(ext,requested or plan.get('output_format'))
    kind=category(ext)
    if kind is None:raise engine.Rejected('输入扩展名不支持。')
    if kind=='image' and image_native.image_format(raw)!=image_native.EXTENSIONS[ext]:raise engine.Rejected('图片实际编码与原扩展名不符。')
    data,compatibility=compat.normalize(raw,ext)
    if kind=='image':out,_=engine.export_image(data,plan.get('rectangles'),image_native.EXTENSIONS[output_ext])
    elif kind=='word':out,_=engine.export_docx(data,plan.get('terms',[]),plan.get('acknowledge_loss'),plan.get('clean_only',False),plan.get('keep',[]))
    else:
        if kind in {'csv','tsv'}:
            decoded=formats.decode(data,plan.get('input_encoding'));plan={**plan,'_line_ending':'\r\n' if '\r\n' in decoded else '\n'}
        out=formats.export(data,kind,plan)
    if kind in {'text','markdown','csv','tsv'}:
        value=out.decode('utf-8-sig')
        if data.startswith(b'\xff\xfe'):out=b'\xff\xfe'+value.encode('utf-16-le')
        elif data.startswith(b'\xfe\xff'):out=b'\xfe\xff'+value.encode('utf-16-be')
        elif data.startswith(b'\xef\xbb\xbf'):out=value.encode('utf-8-sig')
        else:out=value.encode(plan.get('input_encoding') or 'utf-8')
    if output_ext in compat.ODF:out=native_odf(out,output_ext)
    if output_ext=='.xlsm':out=xlsm_container(out)
    report=verify(out,output_ext,plan,data)
    if len(out)>engine.MAX_BYTES:raise engine.Rejected('输出超过20MB。')
    report['checks']['output_format_matches_request']=True
    report.update(input_format=ext,output_format=output_ext,format_preserved=ext==output_ext,conversion_explicit=ext!=output_ext,compatibility=compatibility)
    report['output_sha256']=engine.digest(out)
    return out,report,output_ext
