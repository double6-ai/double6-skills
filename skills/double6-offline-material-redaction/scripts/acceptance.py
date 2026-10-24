"""安装后在本机自动执行合成正反例；不需要浏览器、网络或真实材料。"""
import argparse
import io
import json
import platform
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

import engine
import formats
import image_native
import output_policy as policy

SECRET='DEMO_PRIVATE_782'
KEEP='BUSINESS 120'


def word():
    d=Document();d.add_paragraph(SECRET+' '+KEEP);d.core_properties.author=SECRET
    b=io.BytesIO();d.save(b);return b.getvalue()


def excel():
    w=Workbook();w.active['A1']=SECRET;w.active['B1']=KEEP;w.active['C1']=.15;w.active['C1'].number_format='0%'
    w.create_sheet('Hidden').sheet_state='hidden';w['Hidden']['A1']=SECRET
    b=io.BytesIO();w.save(b);return b.getvalue()


def slides():
    d=Presentation();d.slides.add_slide(d.slide_layouts[6]).shapes.add_textbox(Inches(1),Inches(1), Inches(7), Inches(2)).text=SECRET+' '+KEEP
    b=io.BytesIO();d.save(b);return b.getvalue()


def picture(fmt):
    im=Image.new('RGB',(160,100),'white');draw=ImageDraw.Draw(im);draw.rectangle((10,10,60,35),fill='red');draw.rectangle((20,60,130,80),fill='black')
    b=io.BytesIO();im.save(b,format=fmt);return b.getvalue()


def pdf(hidden=False, encrypted=False):
    f=__import__('pdf_preserve').backend();d=f.open();p=d.new_page(width=300,height=200)
    p.insert_text((20,35),SECRET);p.insert_text((20,120),KEEP)
    if hidden:p.insert_text((20,180),'INVISIBLE_PRIVATE',render_mode=3)
    d.set_metadata({'author':SECRET});d.embfile_add('hidden.txt',SECRET.encode())
    b=d.tobytes(encryption=f.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password')if encrypted else d.tobytes();d.close();return b


def plan(kind):
    return {'version':1,'kind':kind,'reviewed':True,'terms':[SECRET],'keep':[KEEP],'acknowledge_loss':True}


def pdf_plan():
    return {'version':1,'kind':'pdf','reviewed':True,'reviewed_pages':[1],'pages':[{'page':1,'rectangles':[[20,30,290,95]]}],'keep':[KEEP]}


def rejected(fn, part):
    try:fn()
    except engine.Rejected as exc:return part in str(exc)
    return False


def run():
    cases=[]
    def case(name,kind,fn):
        try:ok=fn() is True;reason='' if ok else '行为与本例要求不符'
        except Exception as exc:ok=False;reason='意外异常：'+type(exc).__name__
        cases.append({'name':name,'kind':kind,'passed':ok,'reason':reason})
    def roundtrip(ext,raw,p):
        out,report,target=policy.export(raw,ext,p)
        assert target==ext and report['format_preserved'] and all(report['checks'].values())
        assert raw!=out
        return True
    for ext,kind,raw in [('.docx','word',word()),('.txt','text',(SECRET+' '+KEEP).encode()),('.md','markdown',(SECRET+' '+KEEP).encode()),('.markdown','markdown',(SECRET+' '+KEEP).encode()),('.csv','csv',(SECRET+','+KEEP).encode()),('.tsv','tsv',(SECRET+'\t'+KEEP).encode()),('.xlsx','excel',excel()),('.xlsm','excel',policy.xlsm_container(excel())),('.pptx','slides',slides())]:
        case('原格式'+ext,'positive',lambda e=ext,k=kind,b=raw:roundtrip(e,b,plan(k)))
    def encoded_text(codec):
        raw=(SECRET+' '+KEEP+'\r\n').encode(codec);p=plan('text')
        if codec=='gb18030':p['input_encoding']=codec
        out,r,ext=policy.export(raw,'.txt',p)
        return '【已隐去1】' in out.decode(codec) and out.decode(codec).endswith('\r\n') and KEEP in out.decode(codec) and SECRET not in out.decode(codec)
    case('原编码UTF16文本','positive',lambda:encoded_text('utf-16'))
    case('原编码GB18030文本','positive',lambda:encoded_text('gb18030'))
    for ext,fmt in [('.png','PNG'),('.jpg','JPEG'),('.jpeg','JPEG'),('.webp','WEBP')]:
        case('原格式'+ext,'positive',lambda e=ext,f=fmt:roundtrip(e,picture(f),{'rectangles':[[5,5,70,45]]}))
    for ext,kind,raw in [('.odt','word',word()),('.ods','excel',excel()),('.odp','slides',slides())]:
        case('原格式'+ext,'positive',lambda e=ext,k=kind,b=raw:roundtrip(e,policy.native_odf(b,e),plan(k)))
    def native_pdf():
        raw=pdf();out,r,ext=policy.export(raw,'.pdf',pdf_plan());f=__import__('pdf_preserve').backend()
        with f.open(stream=out,filetype='pdf')as d:
            assert KEEP in d[0].get_text() and SECRET not in d[0].get_text() and not d.embfile_count()
        return ext=='.pdf' and r['checks']['pdf_not_rasterized']
    case('PDF保留可搜索业务文字','positive',native_pdf)
    def raster():
        p=pdf_plan();p.update(pdf_mode='rasterize',acknowledge_loss=True,keep=[])
        out,r,ext=policy.export(pdf(),'.pdf',p)
        return ext=='.pdf' and r['checks']['image_only_pdf']
    case('明确选择PDF图像化','positive',raster)
    def convert():
        out,r,ext=policy.export(picture('JPEG'),'.jpg',{'rectangles':[[5,5,70,45]]},'.png')
        return ext=='.png' and r['conversion_explicit'] and image_native.image_format(out)=='PNG'
    case('明确选择JPG转PNG','positive',convert)
    case('反例缺失保留项','negative',lambda:rejected(lambda:policy.export(word(),'.docx',{**plan('word'),'keep':['MISSING BUSINESS']}),'复检'))
    case('反例补选原词不存在','negative',lambda:rejected(lambda:policy.export(word(),'.docx',{**plan('word'),'terms':['MISSING PRIVATE']}),'未在预览'))
    case('反例未框选图片','negative',lambda:rejected(lambda:policy.export(picture('PNG'),'.png',{'rectangles':[]}),'框选'))
    case('反例图片越界框选','negative',lambda:rejected(lambda:policy.export(picture('PNG'),'.png',{'rectangles':[[0,0,9999,10]]}),'边界'))
    case('反例伪装图片扩展名','negative',lambda:rejected(lambda:policy.export(picture('JPEG'),'.png',{'rectangles':[[0,0,20,20]]}),'编码'))
    for ext in policy.LEGACY:
        case('反例旧格式'+ext+'未授权转换','negative',lambda e=ext:rejected(lambda:policy.export(b'not executed',e,plan(policy.category(e))),'原格式'))
    case('反例PDF漏核对页码','negative',lambda:rejected(lambda:policy.export(pdf(),'.pdf',{**pdf_plan(),'reviewed_pages':[]}),'逐页'))
    def cropped_rotated_pdf():
        f=__import__('pdf_preserve').backend();d=f.open(stream=pdf(),filetype='pdf');d[0].set_cropbox(f.Rect(5,5,295,195));d[0].set_rotation(90);raw=d.tobytes();d.close();return raw
    case('反例PDF裁剪并旋转','negative',lambda:rejected(lambda:policy.export(cropped_rotated_pdf(),'.pdf',pdf_plan()),'裁剪框和旋转'))
    case('反例PDF加密','negative',lambda:rejected(lambda:policy.export(pdf(encrypted=True),'.pdf',pdf_plan()),'加密'))
    case('反例PDF不可见文字','negative',lambda:rejected(lambda:policy.export(pdf(hidden=True),'.pdf',pdf_plan()),'不可见'))
    def ambiguous_pdf(covered):
        f=__import__('pdf_preserve').backend();d=f.open();p=d.new_page(width=300,height=200);p.insert_text((20,35),SECRET);p.insert_text((20,120),KEEP)
        p.insert_text((20,180),'HIDDEN PRIVATE',color=(0,0,0)if covered else(1,1,1))
        if covered:p.draw_rect(f.Rect(10,160,200,190),color=None,fill=(1,1,1))
        raw=d.tobytes();d.close();return rejected(lambda:policy.export(raw,'.pdf',pdf_plan()),'原生路线')
    case('反例PDF白色不可见文字','negative',lambda:ambiguous_pdf(False))
    case('反例PDF文字被图形盖住','negative',lambda:ambiguous_pdf(True))
    case('反例图像化未接受损耗','negative',lambda:rejected(lambda:policy.export(pdf(),'.pdf',{**pdf_plan(),'pdf_mode':'rasterize','keep':[]}),'明确选择'))
    def no_overwrite():
        import cli
        with tempfile.TemporaryDirectory(prefix='redaction-acceptance-')as tmp:
            p=Path(tmp)/'existing.txt';p.write_bytes(b'ORIGINAL')
            try:cli.write_new(p,b'CHANGED')
            except cli.Failure:return p.read_bytes()==b'ORIGINAL' and len(list(Path(tmp).iterdir()))==1
            return False
    def image_pdf(mode='visible'):
        f=__import__('pdf_preserve').backend();d=f.open();page=d.new_page(width=300,height=200)
        if mode=='clipped':page.insert_image(f.Rect(20,20,180,120),stream=picture('PNG'));xref=page.get_contents()[0];content=d.xref_stream(xref);d.update_stream(xref,b'q 20 20 50 50 re W n '+content+b' Q')
        else:
            image=picture('PNG')
            if mode=='alpha':
                im=Image.new('RGBA',(160,100),(255,0,0,100));b=io.BytesIO();im.save(b,format='PNG');image=b.getvalue()
            page.insert_image(f.Rect(20,20,180,120),stream=image)
        if mode=='stroke':page.draw_line(f.Point(50,20),f.Point(50,110),color=(1,1,1),width=40)
        if mode=='covered':page.draw_rect(f.Rect(20,20,80,60),color=None,fill=(1,1,1))
        page.insert_text((20,180),KEEP);raw=d.tobytes();d.close();return raw
    image_plan={'reviewed_pages':[1],'pages':[{'page':1,'rectangles':[[40,40,185,145]]}],'keep':[KEEP]}
    def visible_image_pdf():
        out,r,ext=policy.export(image_pdf(),'.pdf',image_plan)
        return ext=='.pdf' and r['checks']['pdf_hidden_objects_absent'] and r['checks']['keep_phrases_preserved']
    case('PDF可见图片原生删改','positive',visible_image_pdf)
    for mode in ('clipped','covered','alpha','stroke'):
        case('反例PDF隐藏图片'+mode,'negative',lambda m=mode:rejected(lambda:policy.export(image_pdf(m),'.pdf',image_plan),'原生路线'))
    def nested_or_clipped_text(nested):
        f=__import__('pdf_preserve').backend();original=f.open(stream=pdf(),filetype='pdf')
        if nested:
            d=f.open();page=d.new_page(width=300,height=200);page.show_pdf_page(page.rect,original,0);raw=d.tobytes();d.close()
        else:
            page=original[0]
            for xref in page.get_contents():original.update_stream(xref,b'q 10 10 20 20 re W n '+original.xref_stream(xref)+b' Q')
            raw=original.tobytes()
        original.close();return rejected(lambda:policy.export(raw,'.pdf',pdf_plan()),'原生路线')
    def safe_form():
        f=__import__('pdf_preserve').backend();original=f.open(stream=pdf(),filetype='pdf');d=f.open();page=d.new_page(width=300,height=200);page.show_pdf_page(page.rect,original,0);raw=d.tobytes();d.close();original.close()
        out,r,ext=policy.export(raw,'.pdf',pdf_plan())
        return ext=='.pdf'and r['checks']['keep_phrases_preserved']and r['checks']['pdf_not_rasterized']
    case('正常嵌套FormXObject','positive',safe_form)
    def clipped_form():
        f=__import__('pdf_preserve').backend();original=f.open(stream=image_pdf('clipped'),filetype='pdf');d=f.open();page=d.new_page(width=300,height=200);page.show_pdf_page(page.rect,original,0);raw=d.tobytes();d.close();original.close()
        return rejected(lambda:policy.export(raw,'.pdf',image_plan),'原生路线')
    case('反例Form内隐藏裁剪图片','negative',clipped_form)
    def safe_rect_clip():
        f=__import__('pdf_preserve').backend();d=f.open(stream=pdf(),filetype='pdf');page=d[0]
        for xref in page.get_contents():d.update_stream(xref,b'q 0 0 300 200 re W n '+d.xref_stream(xref)+b' Q')
        raw=d.tobytes();d.close();out,r,ext=policy.export(raw,'.pdf',pdf_plan());return ext=='.pdf'and r['checks']['keep_phrases_preserved']
    case('正常矩形裁剪内完整文字','positive',safe_rect_clip)
    case('反例PDF被裁剪的文字','negative',lambda:nested_or_clipped_text(False))
    def forged_image_overlay():
        f=__import__('pdf_preserve').backend();raw=image_pdf();d=f.open(stream=raw,filetype='pdf');d[0].draw_rect(f.Rect(20,20,92.5,72.5),color=None,fill=(0,0,0));fake=d.tobytes();d.close()
        return rejected(lambda:policy.verify(fake,'.pdf',image_plan,raw),'本工具')
    case('反例PDF黑盖保留原图','negative',forged_image_overlay)
    case('同原件计划确定性PDF','positive',lambda:policy.export(pdf(),'.pdf',pdf_plan())[0]==policy.export(pdf(),'.pdf',pdf_plan())[0])
    def offpage_text(crop):
        f=__import__('pdf_preserve').backend();d=f.open(stream=pdf(),filetype='pdf');page=d[0]
        page.insert_text((20,180)if crop else(330,100),'OFFPAGE_PRIVATE_456')
        if crop:page.set_cropbox(f.Rect(0,0,300,150))
        raw=d.tobytes();d.close();return rejected(lambda:policy.export(raw,'.pdf',pdf_plan()),'原生路线')
    case('反例PDF页面之外隐藏文字','negative',lambda:offpage_text(False))
    case('反例PDF裁剪之外隐藏文字','negative',lambda:offpage_text(True))
    case('反例已有文件不得覆盖','negative',no_overwrite)
    def neutral_download():
        import server,threading,urllib.request,urllib.parse
        service=server.LocalServer(0);thread=threading.Thread(target=service.serve_forever);thread.start();base=service.base+'/'+service.nonce+'/'
        opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        def post(route,data,headers):return opener.open(urllib.request.Request(base+route,data,{'Origin':service.base,**headers},method='POST'))
        try:
            job=__import__('json').load(post('api/import',word(),{'X-File-Kind':'word','X-File-Extension':'.docx','X-File-Name':urllib.parse.quote(SECRET+'_13800000000.docx')}))
            result=__import__('json').load(post('api/export',__import__('json').dumps({**plan('word'),'id':job['id']}).encode(),{'Content-Type':'application/json'}))
            header=urllib.parse.unquote(opener.open(base+result['download']).headers['Content-Disposition'])
            return SECRET not in header and '13800000000'not in header and '.docx'in header and result['report']['checks']['neutral_output_filename']
        finally:service.shutdown();thread.join();service.server_close()
    case('下载中性新文件名','positive',neutral_download)
    def object_metadata(private=False):
        import pdf_preserve
        f=pdf_preserve.backend();original=f.open(stream=pdf(),filetype='pdf');doc=f.open();page=doc.new_page(width=300,height=200);page.show_pdf_page(page.rect,original,0);original.close()
        xref=next(x for x in range(1,doc.xref_length())if doc.xref_get_key(x,'Subtype')[1]=='/Form')
        if private:doc.xref_set_key(xref,'UnsupportedPrivate','(PRIVATE_FORM_782)')
        else:
            meta=doc.get_new_xref();doc.update_object(meta,'<</Type /Metadata /Subtype /XML>>');doc.update_stream(meta,b'<xmp>PRIVATE_FORM_782</xmp>');doc.xref_set_key(xref,'Metadata',str(meta)+' 0 R');doc.xref_set_key(xref,'PieceInfo','<</History (PRIVATE_FORM_782)>>')
        raw=doc.tobytes();doc.close();return raw
    def cleaned_object_metadata():
        import pdf_preserve
        output,_,_=policy.export(object_metadata(),'.pdf',pdf_plan())
        with pdf_preserve.backend().open(stream=output,filetype='pdf')as doc:
            raw=b' '.join(doc.xref_object(x).encode()+(doc.xref_stream(x)or b'')for x in range(1,doc.xref_length()))
            return b'PRIVATE_FORM_782'not in raw
    case('PDF嵌套对象元数据与历史清理','positive',cleaned_object_metadata)
    case('PDF未知私有绘制字段','negative',lambda:rejected(lambda:policy.export(object_metadata(True),'.pdf',pdf_plan()),'私有字段'))
    def group_private():
        import pdf_preserve
        f=pdf_preserve.backend();raw=object_metadata();doc=f.open(stream=raw,filetype='pdf');xref=next(x for x in range(1,doc.xref_length())if doc.xref_get_key(x,'Subtype')[1]=='/Form');doc.xref_set_key(xref,'Group','<</S /Transparency /CS /DeviceRGB /Private (HIDDEN_GROUP_782)>>');raw=doc.tobytes();doc.close();return raw
    case('PDF混合组隐藏字段','negative',lambda:rejected(lambda:policy.export(group_private(),'.pdf',pdf_plan()),'混合组'))
    files=sorted(Path(__file__).parent.rglob('*.py'))+sorted((Path(__file__).parent/'web').glob('*'))
    source_sha=engine.digest('\n'.join(p.relative_to(Path(__file__).parents[1]).as_posix()+':'+engine.digest(p.read_bytes())for p in files if p.is_file()).encode())
    return {'ok':all(c['passed']for c in cases),'version':engine.VERSION,'source_sha256':source_sha,'lock_sha256':engine.digest((Path(__file__).parents[1]/'requirements-runtime.lock').read_bytes()),'python':platform.python_version(),'platform':platform.system(),'machine':platform.machine(),
            'positive':{'passed':sum(c['passed']for c in cases if c['kind']=='positive'),'total':sum(c['kind']=='positive'for c in cases)},
            'negative':{'passed':sum(c['passed']for c in cases if c['kind']=='negative'),'total':sum(c['kind']=='negative'for c in cases)},
            'cases':cases,'synthetic_only':True,'not_tested':['旧DOC/XLS/PPT/RTF转换依赖的原生Office执行','任意真实材料的补选完整性','可见桌面双击行为']}


def main():
    for stream in (sys.stdout,sys.stderr):
        if hasattr(stream,'reconfigure'):stream.reconfigure(encoding='utf-8')
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--report',type=Path);a=p.parse_args();result=run()
    text=json.dumps(result,ensure_ascii=False,indent=2)
    if a.report:
        import cli
        cli.write_new(a.report,text.encode())
    print(text);return 0 if result['ok'] else 4


if __name__=='__main__':raise SystemExit(main())
