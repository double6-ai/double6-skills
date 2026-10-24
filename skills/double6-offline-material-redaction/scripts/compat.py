"""Legacy/ODF import boundary. Rebuild only supported visible content into OOXML.

ODF is parsed without running office code. Binary/RTF import uses a private,
macro-disabled LibreOffice profile; its output is never the delivered document.
"""
import io
import os
from pathlib import Path
import base64
import subprocess
import sys
import tempfile
from datetime import datetime
from xml.etree import ElementTree as ET

import engine
import cell_display
from engine import Rejected

KINDS={'.doc':'word','.rtf':'word','.odt':'word','.xls':'excel','.ods':'excel','.ppt':'pdf','.odp':'slides'}
ODF={'.odt':'application/vnd.oasis.opendocument.text','.ods':'application/vnd.oasis.opendocument.spreadsheet','.odp':'application/vnd.oasis.opendocument.presentation'}
NS={p:'{urn:oasis:names:tc:opendocument:xmlns:'+v+':1.0}'for p,v in [('office','office'),('text','text'),('table','table'),('draw','drawing'),('style','style'),('presentation','presentation')]}
NS['fo']='{urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0}'
O,T,A,D,S,P,F=[NS[x]for x in ['office','text','table','draw','style','presentation','fo']]
NOTICE='兼容导入仅保留支持的可见文字、数值和表格，再输出DOCX/XLSX/PPTX；格式化表格数值可能按源显示缓存转为文字，不保留原计算类型。原属性、隐藏内容、备注、修订历史和原容器不复制。图片、图表、版式及交互不保留；请核对转换后的全文与关键业务内容。旧二进制/RTF还需要本机LibreOffice，转换可能影响公式计算结果。'


def converter():
    value=os.environ.get('REDACTION_LIBREOFFICE')
    candidates=[Path(value)]if value else ([Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')]if sys.platform=='darwin'else [Path(os.environ.get('LOCALAPPDATA',''))/'Double6/LibreOffice/program/soffice.exe',Path('C:/Program Files/LibreOffice/program/soffice.exe')])
    for p in candidates:
        if p.is_absolute()and p.is_file():return p.resolve()
    raise Rejected('旧DOC/XLS/PPT和RTF需要本机LibreOffice；先按Skill安装并设置REDACTION_LIBREOFFICE。ODT/ODS/ODP无需此依赖。')


def status():
    try:
        exe=converter()
        return {'available':True,'separate_dependency':True,'launcher_sha256':engine.digest(exe.read_bytes()),'execution_verified_by_doctor':False}
    except (Rejected,OSError):return {'available':False,'separate_dependency':True,'execution_verified_by_doctor':False}


def _convert(data,ext):
    if ext=='.rtf':
        if not data.lstrip().startswith(b'{\\rtf'):raise Rejected('RTF文件头不符，未猜测格式。')
    elif not data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):raise Rejected('旧Office文件头不符；不接受改名的HTML/XML或加密伪装文件。')
    exe=converter();target={'.doc':'odt','.rtf':'odt','.xls':'ods','.ppt':'pdf'}[ext]
    # Separate profile avoids attaching to any existing user Office session.
    with tempfile.TemporaryDirectory(prefix='redaction-compat-')as tmp:
        root=Path(tmp);profile=root/'profile';(profile/'user').mkdir(parents=True)
        os.chmod(root,0o700)
        config='''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop><prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop><prop oor:name="DisableOLEAutomation" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop><prop oor:name="Field" oor:op="fuse"><value>false</value></prop><prop oor:name="Chart" oor:op="fuse"><value>false</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item><item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="FirstRun" oor:op="fuse"><value>false</value></prop></item></oor:items>'''
        (profile/'user/registrymodifications.xcu').write_text(config,encoding='utf-8')
        source=root/('input'+ext);source.write_bytes(data);source.chmod(0o600)
        runner=exe.with_suffix('.com')if sys.platform=='win32'and exe.with_suffix('.com').is_file()else exe
        args=[str(runner),'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--norestore','--convert-to',target,'--outdir',str(root),str(source)]
        if sys.platform=='darwin':
            sandbox=Path('/usr/bin/sandbox-exec')
            if not sandbox.is_file():raise Rejected('本机缺少兼容转换网络隔离入口，已停止。')
            args=[str(sandbox),'-p','(version 1)(allow default)(deny network-outbound)(allow network-outbound (remote ip "localhost:*"))',*args]
        elif sys.platform=='win32':
            # A dedicated block rule is installed by the documented dependency setup,
            # never silently edit global firewall/profile or attach to existing Office.
            script="""$ErrorActionPreference='Stop';$paths=@($args[0],(Join-Path (Split-Path $args[0]) 'soffice.bin'));foreach($p in $paths){$found=@(Get-NetFirewallApplicationFilter -Program $p -ErrorAction Stop | Get-NetFirewallRule | Where-Object {$_.Enabled -eq 'True' -and $_.Direction -eq 'Outbound' -and $_.Action -eq 'Block' -and $_.Profile -eq 'Any'});if($found.Count -lt 1){exit 2}};exit 0"""
            # Command is constant; executable path travels as an environment value.
            script="if(@(Get-NetFirewallProfile|Where-Object {-not $_.Enabled}).Count -gt 0){exit 3};"+script
            script=script.replace('$args[0]','$env:REDACTION_CONVERTER_PROBE')
            guard_env=os.environ.copy();guard_env['REDACTION_CONVERTER_PROBE']=str(exe)
            probe=subprocess.run([str(Path(os.environ.get('SystemRoot','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'),'-NoProfile','-EncodedCommand',base64.b64encode(script.encode('utf-16le')).decode()],capture_output=True,timeout=15,env=guard_env)
            if probe.returncode!=0:raise Rejected('Windows兼容转换出站隔离核对失败，已停止。')
        else:raise Rejected('旧格式兼容转换尚未验证此系统。')
        env=os.environ.copy()
        for key in ('PYTHONHOME','PYTHONPATH','URE_BOOTSTRAP','UNO_PATH'):env.pop(key,None)
        proc=subprocess.Popen(args,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,env=env,start_new_session=sys.platform!='win32',creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if sys.platform=='win32'else 0)
        try:code=proc.wait(timeout=60)
        except subprocess.TimeoutExpired:
            if sys.platform=='win32':subprocess.run(['taskkill.exe','/PID',str(proc.pid),'/T','/F'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            else:
                import signal
                try:os.killpg(proc.pid,signal.SIGKILL)
                except ProcessLookupError:pass
            proc.wait();raise Rejected('本地兼容转换超时，未导出；仅结束本次转换进程。')
        result=root/('input.'+target)
        if code!=0 or not result.is_file()or not 0<result.stat().st_size<=engine.MAX_BYTES:raise Rejected('本地兼容转换未完成；可能加密、损坏或含未支持对象。')
        return result.read_bytes(),'.'+target


def _roots(data,ext):
    import formats
    parts=formats.package(data)
    if parts.get('mimetype',b'').decode('ascii','replace')!=ODF[ext]:raise Rejected('ODF实际类型与扩展名不符。')
    if 'content.xml'not in parts:raise Rejected('ODF缺少正文。')
    if any('Scripts/'in n or n.startswith('Basic/')or n.startswith('Object ')for n in parts):raise Rejected('ODF含宏或嵌入对象，已停止兼容导入。')
    manifest=engine.parse_xml(parts.get('META-INF/manifest.xml',b'<manifest/>'))
    if any(x.tag.endswith('}encryption-data')for x in manifest.iter()):raise Rejected('加密ODF尚未支持。')
    roots=[engine.parse_xml(parts[n])for n in ('styles.xml','content.xml')if n in parts]
    return roots,roots[-1]


class Visible:
    def __init__(self,roots):
        self.styles={};self.defaults={}
        self.data_styles={e.get(S+'name'):e for root in roots for e in root.iter()if e.tag.startswith('{urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0}')and e.tag.endswith('-style')}
        for root in roots:
            for e in root.iter(S+'style'):
                key=(e.get(S+'name'),e.get(S+'family'))
                if key in self.styles:raise Rejected('ODF同类型样式名冲突，已停止。')
                self.styles[key]=e
            for e in root.iter(S+'default-style'):
                self.defaults[e.get(S+'family')]={k:v for c in e for k,v in c.attrib.items()}
        self.layers={e.get(D+'name'):dict(e.attrib)for root in roots for e in root.iter(D+'layer')}
        self.hidden_removed=0
        for root in roots:
            if any(e.tag in {T+'hidden-text',T+'hidden-paragraph'}or e.get(T+'condition')is not None for e in root.iter()):raise Rejected('ODF条件隐藏或条件样式尚未支持，已停止。')
            if any(e.find(S+'map')is not None for e in root.iter(S+'style')):raise Rejected('ODF条件内容样式尚未支持，已停止。')
    def props(self,name,family,seen=None):
        seen=set()if seen is None else seen
        result=self.defaults.get(family,{}).copy()
        if not name:return result
        key=(name,family)
        if key in seen:raise Rejected('ODF样式循环引用。')
        el=self.styles.get(key)
        if el is None:return result
        seen.add(key);result.update(self.props(el.get(S+'parent-style-name'),family,seen))
        if el.get(S+'data-style-name')is not None:result[S+'data-style-name']=el.get(S+'data-style-name')
        for child in el:
            if child.tag.endswith('-properties'):result.update(child.attrib)
        return result
    def cell_value(self,el,value,paragraphs,fallback=None):
        props=self.props(el.get(A+'style-name')or fallback,'table-cell')
        color=props.get(F+'color','').lower();background=props.get(F+'background-color','transparent').lower()
        if color and (color==background or color=='#ffffff'and background=='transparent'):
            raise Rejected('ODF单元格文字与背景同色，当前无法安全判断可见性，已停止。')
        name=props.get(S+'data-style-name');style=self.data_styles.get(name)
        if name and style is None:raise Rejected('ODF单元格显示样式引用缺失。')
        if style is not None:
            number_ns='{urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0}'
            allowed={number_ns+x for x in ['number','text','currency-symbol','year','month','day','day-of-week','hours','minutes','seconds','am-pm','text-content']}
            if any(c.tag not in allowed for c in style):raise Rejected('ODF单元格含未支持的条件或显示格式，已停止。')
            if not any(c.tag!=number_ns+'text'or (c.text or '').strip()for c in style):
                self.hidden_removed+=1;return None
            if not paragraphs and value is not None:raise Rejected('ODF格式化单元格缺少可见显示缓存，已停止。')
        if isinstance(value,(int,float,datetime))and not isinstance(value,bool)and paragraphs:
            shown='\n'.join(paragraphs)
            if not shown:self.hidden_removed+=1;return None
            # ODF按源文件显示缓存保留业务读法；格式化数值可转为普通文本。
            if shown!=cell_display.plain(value):return shown
        return value
    def hidden(self,el):
        if el.tag in {O+'annotation',O+'annotation-end',T+'tracked-changes',P+'notes'}:return True
        props={}
        text_family='paragraph'if el.tag in {T+'p',T+'h'}else'text'
        table_family={A+'table':'table',A+'table-row':'table-row',A+'table-column':'table-column'}.get(el.tag,'table-cell')
        draw_family='drawing-page'if el.tag==D+'page'else'graphic'
        for attr,family in ((T+'style-name',text_family),(A+'style-name',table_family),(D+'style-name',draw_family),(P+'style-name','presentation')):
            if el.get(attr)is not None:props.update(self.props(el.get(attr),family))
        if el.tag in {T+'p',T+'h',T+'span'}:props={**self.props(None,text_family),**props}
        attrs={**self.layers.get(el.get(D+'layer'),{}),**props,**el.attrib}
        return attrs.get(T+'display')=='none'or attrs.get(A+'display')in {'false','0'}or attrs.get(A+'visibility')in {'collapse','filter'}or attrs.get(P+'visibility')=='hidden'or attrs.get(D+'display')=='none'or attrs.get(F+'visibility')=='hidden'
    def text(self,el):
        if self.hidden(el):self.hidden_removed+=1;return ''
        if el.tag in {D+'frame',D+'image',D+'object',D+'object-ole',O+'binary-data'}:return ''
        if el.tag in {T+'note',T+'ruby'}:raise Rejected('ODF行内脚注/注音结构尚未支持。')
        if el.tag==T+'s':
            n=number(el.get(T+'c','1'),10000);return ' '*n
        if el.tag==T+'tab':return '\t'
        if el.tag==T+'line-break':return '\n'
        result=el.text or ''
        for child in el:result+=self.text(child)+(child.tail or '')
        return result


def number(value,limit):
    try:n=int(value)
    except (ValueError,TypeError):raise Rejected('ODF重复范围无效。')
    if not 1<=n<=limit:raise Rejected('ODF重复范围超限。')
    return n


def _table(el,view):
    rows=[];total=0
    def row_nodes(node):
        for child in node:
            if child.tag==A+'table-row':yield child
            elif child.tag in {A+'table-header-rows',A+'table-row-group',A+'table-rows'}:
                if view.hidden(child):view.hidden_removed+=1
                else:yield from row_nodes(child)
    hidden_cols=set();column_styles={};col=0
    def columns(node,inherited_hidden=False):
        for c in node:
            if c.tag==A+'table-column':yield c,inherited_hidden or view.hidden(c)
            elif c.tag in {A+'table-column-group',A+'table-columns',A+'table-header-columns'}:yield from columns(c,inherited_hidden or view.hidden(c))
    for c,is_hidden in columns(el):
        repeat=number(c.get(A+'number-columns-repeated','1'),16384)
        if is_hidden:hidden_cols.update(range(col,col+repeat))
        if c.get(A+'default-cell-style-name'):
            for offset in range(repeat):column_styles[col+offset]=c.get(A+'default-cell-style-name')
        col+=repeat
        if col>16384:raise Rejected('ODF列范围超限。')
    for row in row_nodes(el):
        if view.hidden(row):view.hidden_removed+=1;continue
        nrows=number(row.get(A+'number-rows-repeated','1'),1048576);values=[];col=0
        for cell in row:
            if cell.tag not in {A+'table-cell',A+'covered-table-cell'}:continue
            repeat=number(cell.get(A+'number-columns-repeated','1'),16384)
            if any(x.tag==A+'table'for x in cell.iter()):raise Rejected('ODF嵌套表格尚未支持。')
            typ=cell.get(O+'value-type');value=None
            if not view.hidden(cell):
                paragraphs=[view.text(p)for p in cell if p.tag in {T+'p',T+'h'}]
                if typ in {'float','percentage','currency'}:
                    try:value=float(cell.get(O+'value'))
                    except (ValueError,TypeError):raise Rejected('ODF数值缓存无效。')
                    import math
                    if not math.isfinite(value):raise Rejected('ODF数值非有限数。')
                elif typ=='boolean':value=cell.get(O+'boolean-value')=='true'
                elif typ=='date':
                    try:value=datetime.fromisoformat(cell.get(O+'date-value'))
                    except (ValueError,TypeError):raise Rejected('ODF日期缓存无效。')
                elif typ=='time':value=cell.get(O+'time-value')
                else:value=('\n'.join(paragraphs)or None)if paragraphs else cell.get(O+'string-value')
                if cell.get(A+'formula')and value is None:raise Rejected('ODF公式缺少缓存值，已停止。')
                fallback=row.get(A+'default-cell-style-name')or column_styles.get(col)or el.get(A+'default-cell-style-name')
                value=view.cell_value(cell,value,paragraphs,fallback)
            else:view.hidden_removed+=1
            for offset in range(repeat):
                values.append(None if col+offset in hidden_cols else value)
            col+=repeat
            if len(values)>100000:raise Rejected('ODF单元格范围超限。')
        while values and values[-1] is None:values.pop()
        if not values:continue
        total+=len(values)*nrows
        if total>100000:raise Rejected('ODF单元格范围超过10万格。')
        rows.extend([values.copy()for _ in range(nrows)])
    return rows


def _word(body,roots,view):
    from docx import Document
    doc=Document();count=0
    def visit(container):
        nonlocal count
        for e in container:
            if view.hidden(e):view.hidden_removed+=1;continue
            if e.tag in {T+'p',T+'h'}:
                text=view.text(e)
                if text:doc.add_paragraph(text);count+=len(text)
            elif e.tag==A+'table':
                rows=_table(e,view)
                if not rows:continue
                if max(map(len,rows))>30:raise Rejected('ODT表格列数超过30列。')
                table=doc.add_table(rows=0,cols=max(map(len,rows)))
                for row in rows:
                    cells=table.add_row().cells
                    for c,value in zip(cells,row):c.text=''if value is None else str(value);count+=len(c.text)
            elif e.tag==O+'forms' and len(e)==0:continue
            elif e.tag in {T+'section',T+'list',T+'list-item',T+'list-header'}:visit(e)
            elif e.tag not in {T+'sequence-decls',T+'tracked-changes',T+'user-field-decls',T+'variable-decls',T+'dde-connection-decls'}:raise Rejected('ODT正文含未支持对象，已停止而非静默丢失。')
    visit(body)
    masters=[x for root in roots for x in root.iter(S+'master-page')]
    with_headers=[x for x in masters if any(c.tag in {S+'header',S+'footer',S+'header-left',S+'footer-left',S+'header-first',S+'footer-first'}for c in x)]
    if with_headers:
        if len(masters)!=1 or any(c.tag.endswith(('-left','-first'))for x in with_headers for c in x):raise Rejected('ODT多母版/首页或奇偶页眉页脚尚未支持。')
        for x in with_headers:
            for child in x:
                if child.tag in {S+'header',S+'footer'}:visit(child)
    if not 0<count<=400000:raise Rejected('ODT无支持的可见内容或文字过多。')
    for key in ('author','title','subject','keywords','comments','last_modified_by','category','identifier','language','version','content_status'):setattr(doc.core_properties,key,'')
    out=io.BytesIO();doc.save(out);return out.getvalue()


def _excel(body,view):
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active);total=0
    for table in body:
        if table.tag!=A+'table':continue
        if view.hidden(table):view.hidden_removed+=1;continue
        rows=_table(table,view);ws=wb.create_sheet('表'+str(len(wb.worksheets)+1))
        for r,row in enumerate(rows,1):
            for c,value in enumerate(row,1):
                if value is not None:
                    cell=ws.cell(r,c,value)
                    if isinstance(value,str):cell.data_type='s'
                    total+=1
    if not wb.worksheets or not 0<total<=100000:raise Rejected('ODS无可见单元格或内容超限。')
    out=io.BytesIO();wb.save(out);return out.getvalue()


def _slides(body,view):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation();pages=0
    for page in body:
        if page.tag!=D+'page':continue
        if view.hidden(page):view.hidden_removed+=1;continue
        slide=prs.slides.add_slide(prs.slide_layouts[6]);pages+=1;items=0
        def visit(node):
            nonlocal items
            for e in node:
                if view.hidden(e):view.hidden_removed+=1;continue
                if e.tag==A+'table':
                    rows=_table(e,view)
                    if not rows:continue
                    cols=max(map(len,rows))
                    if not 0<cols<=10 or len(rows)>100:raise Rejected('ODP表格范围超限。')
                    table=slide.shapes.add_table(len(rows),cols,Inches(1),Inches(1),Inches(8),Inches(4)).table
                    for r,row in enumerate(rows):
                        for c,v in enumerate(row):table.cell(r,c).text=''if v is None else str(v)
                    items+=1
                elif e.tag in {T+'p',T+'h'}:
                    value=view.text(e)
                    if value:slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text=value;items+=1
                elif e.tag in {D+'frame',D+'text-box',D+'g',D+'custom-shape',T+'list',T+'list-header',T+'list-item'}:visit(e)
                elif e.tag not in {P+'notes',O+'annotation',O+'annotation-end',D+'image',D+'page-thumbnail',D+'enhanced-geometry'}:raise Rejected('ODP含未支持的形状/业务对象，已停止。')
        visit(page)
        if not items:raise Rejected('ODP有可见页缺少支持的文字/表格，已停止。')
    if not 1<=pages<=100:raise Rejected('ODP可见页数无效。')
    out=io.BytesIO();prs.save(out);return out.getvalue()


def normalize(data,ext):
    if ext not in KINDS:return data,None
    if len(data)>engine.MAX_BYTES:raise Rejected('文件超过20MB。')
    original_ext=ext
    if ext not in ODF:data,ext=_convert(data,ext)
    if original_ext=='.ppt':
        import formats
        formats.pdf_info(data)
        return data,{'input_extension':original_ext,'normalized_kind':'pdf','notice':'旧PPT保留页面外观转为PDF，再逐页人工框选；导出纯图像PDF，文字不可编辑/搜索。隐藏页、备注和原容器不进入最终副本。','conversion_used':True}
    roots,root=_roots(data,ext);view=Visible(roots);body=root.find(O+'body')
    inner=body.find(O+{'.odt':'text','.ods':'spreadsheet','.odp':'presentation'}[ext])if body is not None else None
    if inner is None:raise Rejected('ODF正文类型不符。')
    output=_word(inner,roots,view)if ext=='.odt'else _excel(inner,view)if ext=='.ods'else _slides(inner,view)
    if len(output)>engine.MAX_BYTES:raise Rejected('兼容重建文件超过20MB。')
    return output,{'input_extension':original_ext,'normalized_kind':KINDS[original_ext],'notice':NOTICE,'hidden_nodes_removed':view.hidden_removed,'conversion_used':original_ext not in ODF}


def selftest():
    import zipfile
    bodies={'.odt':'<office:text><text:p>TEST BUSINESS120</text:p></office:text>','.ods':'<office:spreadsheet><table:table><table:table-row><table:table-cell office:value-type="string"><text:p>TEST BUSINESS120</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet>','.odp':'<office:presentation><draw:page><draw:frame><draw:text-box><text:p>TEST BUSINESS120</text:p></draw:text-box></draw:frame></draw:page></office:presentation>'}
    result={}
    for ext,body in bodies.items():
        out=io.BytesIO();ns=' '.join('xmlns:'+p+'="'+v[1:-1]+'"'for p,v in NS.items())
        with zipfile.ZipFile(out,'w')as z:
            z.writestr('mimetype',ODF[ext]);z.writestr('content.xml','<office:document-content '+ns+'><office:body>'+body+'</office:body></office:document-content>')
        raw,_=normalize(out.getvalue(),ext)
        if KINDS[ext]=='word':report=engine.inspect_docx(raw)
        else:
            import formats
            report=formats.inspect(raw,KINDS[ext])
        result[ext[1:]+'_visible_rebuild']='BUSINESS120'in report['preview']
    return result
