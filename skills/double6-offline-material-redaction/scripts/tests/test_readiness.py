"""制课前反例：隐藏值不暴露、业务读法和真实检查范围。"""
import io,json,sys,unittest,threading,urllib.request,urllib.error
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,Color,PatternFill
from openpyxl.formatting.rule import CellIsRule
from docx import Document
sys.path.insert(0,str(Path(__file__).parents[1]))
import formats,engine,server,compat
from test_compat import odf


def data_book():
    w=Workbook();s=w.active;s['A1']='测试姓名 业务120';s['A2']=.15;s['A2'].number_format='0%';s['A3']=123;s['A3'].number_format='00000';s['A4']=1200.5;s['A4'].number_format='"¥"#,##0.00';s['A5']=datetime(2026,10,24);s['A5'].number_format='yyyy-mm-dd';s['B4']='HIDDEN_STYLE';s['B4'].number_format=';;;';b=io.BytesIO();w.save(b);return b.getvalue()


def plan(**change):
    return {'terms':['测试姓名'],'acknowledge_loss':True,'keep':['业务120','15%','00123','¥1,200.50','2026-10-24'],**change}


class ReadinessTests(unittest.TestCase):
    def test_hidden_format_absent_and_business_formats_preserved(self):
        raw=data_book();out=formats.export(raw,'excel',plan());w=load_workbook(io.BytesIO(out));s=w.active
        self.assertIsNone(s['B4'].value);self.assertEqual((s['A2'].value,s['A2'].number_format),(.15,'0%'));self.assertEqual((s['A3'].value,s['A3'].number_format),(123,'00000'));self.assertEqual(s['A4'].number_format,'"¥"#,##0.00');self.assertEqual(s['A5'].number_format,'yyyy-mm-dd')
        text=formats.inspect(out,'excel')['preview'];self.assertNotIn('HIDDEN_STYLE',text);self.assertTrue(all(x in text for x in plan()['keep']))
    def test_selected_formatted_numeric_value_is_actually_replaced(self):
        out=formats.export(data_book(),'excel',plan(terms=['00123'],keep=['15%']));s=load_workbook(io.BytesIO(out)).active;self.assertIsInstance(s['A3'].value,str);self.assertIn('已隐去',s['A3'].value);self.assertEqual(s['A3'].number_format,'General')
    def test_empty_keep_is_not_reported_as_pass(self):
        out=formats.export(data_book(),'excel',plan(keep=[]));r=formats.verify(out,'excel',plan(keep=[]));self.assertNotIn('keep_phrases_preserved',r['checks']);self.assertEqual(r['not_checked'],['keep_phrases_preserved'])
    def test_missing_keep_rejects_export(self):
        with self.assertRaises(engine.Rejected):formats.export(data_book(),'excel',plan(keep=['应保留但不存在']))
    def test_unknown_number_format_conditional_and_invisible_font_stop(self):
        for change in ['format','conditional','font']:
            w=Workbook();s=w.active;s['A1']='业务120';s['A2']=12
            if change=='format':s['A2'].number_format='0;[Red]-0;;@'
            if change=='conditional':s.conditional_formatting.add('A2',CellIsRule(operator='equal',formula=['12'],font=Font(color='FFFFFF')))
            if change=='font':s['A2'].font=Font(color='FFFFFF')
            b=io.BytesIO();w.save(b)
            with self.assertRaises(engine.Rejected):formats.inspect(b.getvalue(),'excel')
    def test_theme_and_indexed_invisible_fonts_stop(self):
        for color in [Color(theme=0),Color(indexed=1)]:
            w=Workbook();w.active['A1']='业务120';w.active['A2']='HIDDEN_PRIVATE';w.active['A2'].font=Font(color=color);b=io.BytesIO();w.save(b)
            with self.assertRaises(engine.Rejected):formats.inspect(b.getvalue(),'excel')
    def test_theme_background_checked_and_readable_contrast_kept(self):
        for hidden in [True,False]:
            w=Workbook();s=w.active;s['A1']='测试姓名 业务120';s['A1'].fill=PatternFill('solid',fgColor=Color(theme=1));s['A1'].font=Font(color=Color(theme=1 if hidden else 0));b=io.BytesIO();w.save(b)
            if hidden:
                with self.assertRaises(engine.Rejected):formats.inspect(b.getvalue(),'excel')
            else:self.assertIn('业务120',formats.inspect(formats.export(b.getvalue(),'excel',plan(keep=['业务120'])),'excel')['preview'])
    def test_tinted_colors_fail_closed(self):
        w=Workbook();s=w.active;s['A1']='业务120';s['A1'].font=Font(color=Color(theme=1,tint=1));b=io.BytesIO();w.save(b)
        with self.assertRaises(engine.Rejected):formats.inspect(b.getvalue(),'excel')
    def test_word_keep_checked_and_missing_rejected(self):
        d=Document();d.add_paragraph('姓名 业务120');b=io.BytesIO();d.save(b)
        out,r=engine.export_docx(b.getvalue(),['姓名'],True,keep=['业务120']);self.assertTrue(r['checks']['keep_phrases_preserved'])
        with self.assertRaises(engine.Rejected):engine.export_docx(b.getvalue(),['姓名'],True,keep=['应保留但不存在'])
        out,r=engine.export_docx(b.getvalue(),['姓名'],True);self.assertNotIn('keep_phrases_preserved',r['checks']);self.assertIn('keep_phrases_preserved',r['not_checked'])
    def test_odf_formatted_values_follow_display_cache_and_empty_style_drops(self):
        ns='urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0'
        styles='<style:style style:name="P" style:family="table-cell" style:data-style-name="percent"/><number:percentage-style xmlns:number="'+ns+'" style:name="percent"><number:number/><number:text>%</number:text></number:percentage-style><style:style style:name="H" style:family="table-cell" style:data-style-name="empty"/><number:number-style xmlns:number="'+ns+'" style:name="empty"/>'
        body='<office:spreadsheet><table:table><table:table-row><table:table-cell table:style-name="P" office:value-type="percentage" office:value="0.15"><text:p>15%</text:p></table:table-cell><table:table-cell table:style-name="H" office:value-type="string"><text:p>HIDDEN_STYLE</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet>'
        out,_=compat.normalize(odf('.ods',body,styles),'.ods');s=load_workbook(io.BytesIO(out)).active;self.assertEqual(s['A1'].value,'15%');self.assertIsNone(s['B1'].value)


class ServerReadinessTests(unittest.TestCase):
    def setUp(self):
        self.server=server.LocalServer(0);self.thread=threading.Thread(target=self.server.serve_forever);self.thread.start();self.url=self.server.base+'/'+self.server.nonce+'/'
    def tearDown(self):self.server.shutdown();self.thread.join();self.server.server_close()
    def post(self,path,data,headers=None):
        req=urllib.request.Request(self.url+path,data=data,headers={'Origin':self.server.base,**(headers or {})},method='POST');return json.loads(urllib.request.urlopen(req).read())
    def test_gui_scope_and_failed_export_revoke_previous_download(self):
        job=self.post('api/import',data_book(),{'X-File-Kind':'excel','X-File-Extension':'.xlsx'});payload={**plan(keep=[]),'id':job['id'],'reviewed':True}
        r=self.post('api/export',json.dumps(payload).encode());self.assertNotIn('keep_phrases_preserved',r['report']['checks']);self.assertIn('keep_phrases_preserved',r['report']['not_checked']);urllib.request.urlopen(self.url+r['download']).read()
        payload['keep']=['不存在的业务条件']
        with self.assertRaises(urllib.error.HTTPError):self.post('api/export',json.dumps(payload).encode())
        with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(self.url+r['download']).read()
    def test_plan_change_revokes_previous_download_before_reexport(self):
        job=self.post('api/import',data_book(),{'X-File-Kind':'excel','X-File-Extension':'.xlsx'});payload={**plan(keep=[]),'id':job['id'],'reviewed':True}
        out=self.post('api/export',json.dumps(payload).encode());urllib.request.urlopen(self.url+out['download']).read()
        self.assertTrue(self.post('api/invalidate',json.dumps({'id':job['id']}).encode())['invalidated'])
        with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(self.url+out['download']).read()

    def test_served_version_and_support_from_current_runtime(self):
        html=urllib.request.urlopen(self.url).read().decode();self.assertIn(engine.VERSION,html);self.assertNotIn('__VERSION__',html);self.assertNotIn('尚未完成适配',html);self.assertIn('.ods',html)

if __name__=='__main__':unittest.main()
