"""命令行的交付、安全失败和内容保留；只用自编材料。"""
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
sys.path.insert(0, str(Path(__file__).parents[1]))
import cli
import engine


class CliTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.source = self.root / '原始 材料.docx'
        d = Document();d.add_paragraph('测试姓名 业务数字120仍需保留');d.core_properties.author='测试姓名';d.save(self.source)
        self.original = self.source.read_bytes()
        self.plan = self.root / '计划.json'
        self.values = {'version':1,'kind':'word','reviewed':True,'terms':['测试姓名'],'acknowledge_loss':True,'keep':['业务数字120仍需保留']}
        self.save_plan()
    def tearDown(self): self.tmp.cleanup()
    def save_plan(self): self.plan.write_text(json.dumps(self.values,ensure_ascii=False),encoding="utf-8")
    def run_cli(self,*args):
        p = subprocess.run([sys.executable, str(Path(cli.__file__)), *map(str,args)],capture_output=True,text=True,encoding="utf-8")
        return p.returncode,json.loads(p.stdout)
    def test_roundtrip_without_sensitive_stdout(self):
        out=self.root/'干净 副本.docx'
        code,report=self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)
        self.assertEqual(code,0);self.assertNotIn('测试姓名',json.dumps(report,ensure_ascii=False))
        code,report=self.run_cli('verify','--input',out,'--plan',self.plan);self.assertEqual(code,0)
        self.assertEqual(self.source.read_bytes(),self.original)
    def test_no_overwrite_input_or_existing_output(self):
        for out in [self.source,self.plan]:
            self.assertNotEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],0)
        out=self.root/'existing.docx';out.write_bytes(b'preserve')
        self.assertNotEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],0)
        self.assertEqual(out.read_bytes(),b'preserve');self.assertEqual(self.source.read_bytes(),self.original)
    def test_bad_plan_and_keep_failure_write_nothing(self):
        for change in [{'reviewed':False},{'term':['测试姓名']},{'keep':['必须保留而输入没有的内容']}]:
            self.values.update(change);self.save_plan();out=self.root/'not-created.docx'
            self.assertNotEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],0)
            self.assertFalse(out.exists())
            self.values={'version':1,'kind':'word','reviewed':True,'terms':['测试姓名'],'acknowledge_loss':True}
    def test_empty_terms_require_explicit_clean_only(self):
        self.values['terms']=[];self.save_plan();out=self.root/'clean-only.docx'
        self.assertNotEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],0)
        self.values['clean_only']=True;self.save_plan();self.assertEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],0)
    def test_inspect_does_not_print_original_values_by_default(self):
        code,result=self.run_cli('inspect','--input',self.source);self.assertEqual(code,0)
        self.assertNotIn('测试姓名',json.dumps(result,ensure_ascii=False));self.assertNotIn('preview',result)
    def test_style_hidden_text_is_blocked(self):
        d=Document();style=d.styles.add_style('PrivateStyle',1);style.font.hidden=True;d.add_paragraph('隐藏的姓名',style='PrivateStyle');b=io.BytesIO();d.save(b)
        with self.assertRaises(engine.Rejected):engine.inspect_docx(b.getvalue())
    def test_dtd_and_broken_zip_are_blocked(self):
        with zipfile.ZipFile(io.BytesIO(self.original)) as z:parts={n:z.read(n) for n in z.namelist()}
        parts['word/document.xml']=b'<!DOCTYPE a [<!ENTITY x "secret">]>'+parts['word/document.xml']
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as z:
            for n,value in parts.items():z.writestr(n,value)
        # UTF-16的DTD同样要拒绝，不能只靠UTF-8字节匹配。
        parts['word/document.xml']='<?xml version="1.0" encoding="utf-16"?><!DOCTYPE a [<!ENTITY x "secret">]><a>&x;</a>'.encode('utf-16')
        utf16=io.BytesIO()
        with zipfile.ZipFile(utf16,'w') as z:
            for n,value in parts.items():z.writestr(n,value)
        for raw in [b.getvalue(),utf16.getvalue(),b'PK-not-a-document']:
            with self.assertRaises(engine.Rejected):engine.inspect_docx(raw)
    def test_image_verify_requires_source_and_detects_metadata(self):
        im=Image.new('RGB',(50,50),'white');b=io.BytesIO();im.save(b,format='PNG')
        plan={'rectangles':[[0,0,10,10]],'keep':[]};out,_=engine.export_image(b.getvalue(),plan['rectangles'])
        with self.assertRaises(cli.Failure):cli.check_output(out,'image',plan)
        self.assertTrue(all(cli.check_output(out,'image',plan,b.getvalue())['checks'].values()))
    def test_unreferenced_header_is_not_promoted_to_body(self):
        d=Document();d.add_paragraph('当前正文');d.sections[0].header.paragraphs[0].text='当前页眉';b=io.BytesIO();d.save(b)
        with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z:parts={n:z.read(n) for n in z.namelist()}
        parts['word/header-old.xml']=parts['word/header1.xml'].replace('当前页眉'.encode(),'历史隐私'.encode())
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as z:
            for n,value in parts.items():z.writestr(n,value)
        text=engine.inspect_docx(b.getvalue())['preview'];self.assertIn('当前页眉',text);self.assertNotIn('历史隐私',text)
    def test_custom_header_name_resolved_by_relationship(self):
        d=Document();d.add_paragraph('当前正文');d.sections[0].header.paragraphs[0].text='自定义页眉';b=io.BytesIO();d.save(b)
        with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z:parts={n:z.read(n) for n in z.namelist()}
        parts['word/custom-header.xml']=parts.pop('word/header1.xml')
        parts['word/_rels/document.xml.rels']=parts['word/_rels/document.xml.rels'].replace(b'header1.xml',b'custom-header.xml')
        parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'header1.xml',b'custom-header.xml')
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as z:
            for n,value in parts.items():z.writestr(n,value)
        self.assertIn('自定义页眉',engine.inspect_docx(b.getvalue())['preview'])
    def test_first_page_header_and_nested_control_are_blocked(self):
        d=Document();d.add_paragraph('当前正文');d.sections[0].different_first_page_header_footer=True;d.sections[0].first_page_header.paragraphs[0].text='首页专用';b=io.BytesIO();d.save(b)
        with self.assertRaises(engine.Rejected):engine.inspect_docx(b.getvalue())
        d=Document();p=d.add_paragraph('内容控件');p._p.append(OxmlElement('w:sdt'));b=io.BytesIO();d.save(b)
        with self.assertRaises(engine.Rejected):engine.inspect_docx(b.getvalue())
    def test_too_deep_xml_and_boolean_version_rejected(self):
        raw=('<r>'*150+'content'+'</r>'*150).encode()
        with self.assertRaises(engine.Rejected):engine.parse_xml(raw)
        self.values['version']=True;self.save_plan()
        self.assertNotEqual(self.run_cli('verify','--input',self.source,'--plan',self.plan)[0],0)
    @unittest.skipUnless(hasattr(__import__('os'), 'mkfifo'), 'POSIX-only FIFO')
    def test_non_regular_input_does_not_block(self):
        import os
        fifo=self.root/'named-pipe.docx';os.mkfifo(fifo)
        with self.assertRaises(cli.Failure):cli.load_file(fifo)
    def test_large_png_export_stops_before_unverifiable_output(self):
        from unittest.mock import patch
        import random
        im=Image.frombytes('RGB',(100,100),random.Random(5).randbytes(30000));b=io.BytesIO();im.save(b,format='JPEG')
        with patch.object(engine,'MAX_BYTES',len(b.getvalue())+100):
            with self.assertRaises(engine.Rejected):engine.export_image(b.getvalue(),[[0,0,4,4]])
    def test_unmatched_term_is_not_silently_accepted(self):
        self.values['terms']=['输入中没有的名字'];self.save_plan();out=self.root/'not-produced.docx'
        self.assertEqual(self.run_cli('redact','--input',self.source,'--plan',self.plan,'--output',out)[0],2)
        self.assertFalse(out.exists())
    def test_failed_verification_with_original_still_present(self):
        self.assertNotEqual(self.run_cli('verify','--input',self.source,'--plan',self.plan)[0],0)

if __name__=='__main__':unittest.main()
