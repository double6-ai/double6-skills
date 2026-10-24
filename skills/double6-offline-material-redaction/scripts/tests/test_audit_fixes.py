"""回归：0.7.0 离线实测里确认过的脱敏问题。材料全部是虚构示例。"""
import io
import math
import sys
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parents[1]))
import engine
import formats
import output_policy as policy
import pdf_preserve as pdf


PHONE = "13800000000"
KEEP = "整理120条活动反馈"
NOTE = "PDF批注甲乙丙13700000016"


def _page(note=True):
    fitz = pdf.backend()
    doc = fitz.open()
    page = doc.new_page(width=400, height=280)
    page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
    page.insert_text((20, 140), PHONE, fontname="helv", fontsize=14)
    if note:
        page.add_text_annot(fitz.Point(300, 200), NOTE)
    raw = doc.tobytes()
    found = page.search_for(PHONE)[0]
    doc.close()
    return raw, found


def _pdf_plan(box):
    return {
        "version": 1,
        "kind": "pdf",
        "reviewed": True,
        "reviewed_pages": [1],
        "keep": [KEEP],
        "pages": [{"page": 1, "rectangles": [box]}],
    }


def _pixel_box(rect):
    return [(rect.x0 - 1) * 2, (rect.y0 - 1) * 2, (rect.x1 + 1) * 2, (rect.y1 + 1) * 2]


def _cover_words(message):
    """遮盖拒绝要说明原因，不再把任何像素变化都说成画面变了。"""
    return any(part in message for part in (
        "盖住了",
        "看不一致",
        "看不到这些字的边缘",
        "没法对齐",
        "线条或图片",
        "接近字迹的深色",
    ))


def _preview_box(page, text_rect):
    """用 pdfium 的未缩放显示页乘 2 得到预览框。不经过产品自己的旋转函数。

    PyMuPDF 的文字坐标含 UserUnit，它自带的旋转平移不含。先除掉这个倍数，再乘未缩放旋转，才是渲染上的位置。
    """
    fitz = pdf.backend()
    crop = page.cropbox
    cw = abs(crop.x1 - crop.x0)
    ch = abs(crop.y1 - crop.y0)
    if page.rotation in (90, 270):
        unit = page.rect.width / ch
    else:
        unit = page.rect.width / cw
    unscaled = fitz.Rect(text_rect.x0 / unit, text_rect.y0 / unit, text_rect.x1 / unit, text_rect.y1 / unit)
    mapped = unscaled * page.rotation_matrix
    return [mapped.x0 * 2 - 2, mapped.y0 * 2 - 2, mapped.x1 * 2 + 2, mapped.y1 * 2 + 2]


class PdfAnnotationTests(unittest.TestCase):
    def test_visible_note_on_blank_is_removed(self):
        raw, found = _page()
        out, report, _ext = policy.export(raw, ".pdf", _pdf_plan(_pixel_box(found)))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(KEEP, text)
            self.assertNotIn(PHONE, text)
            self.assertNotIn("13700000016", text)
            self.assertIsNone(doc[0].first_annot)

    def test_annotation_does_not_excuse_a_cut_glyph(self):
        raw, found = _page(note=False)
        # 只盖住电话左半边。框外剩下的字形被删改时必须仍然失败。
        partial = [(found.x0 - 1) * 2, (found.y0 - 1) * 2, (found.x0 + found.width / 2) * 2, (found.y1 + 1) * 2]
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", _pdf_plan(partial))
        self.assertIn("框外", str(caught.exception))


class TableKeepTests(unittest.TestCase):
    def _office_plan(self, kind, terms, keep):
        return {"version": 1, "kind": kind, "reviewed": True, "terms": terms, "keep": keep, "acknowledge_loss": True}

    def test_plain_negative_stays_and_real_formulas_are_quoted(self):
        raw = "姓名,金额\n张三,-120\n公式,=1+1\n危险,-2+3\n命令,@SUM(A1)\n加号,+1+1\n".encode()
        for kind in ("csv", "tsv"):
            text = raw.decode().replace(",", "\t") if kind == "tsv" else raw.decode()
            out = formats.export(text.encode(), kind, self._office_plan(kind, ["张三"], ["-120"]))
            rows = formats.csv_rows(out, kind)
            self.assertEqual(rows[1][1], "-120")
            self.assertEqual(rows[2][1], "'=1+1")
            self.assertEqual(rows[3][1], "'-2+3")
            self.assertEqual(rows[4][1], "'@SUM(A1)")
            self.assertEqual(rows[5][1], "'+1+1")
            self.assertNotIn("张三", str(rows))

    def test_keep_sees_quoted_formula_as_changed(self):
        raw = "姓名,公式\n张三,=1+1\n".encode()
        with self.assertRaises(engine.Rejected) as caught:
            formats.export(raw, "csv", self._office_plan("csv", ["张三"], ["=1+1"]))
        self.assertIn("完整单元格", str(caught.exception))
        self.assertIn("=1+1", str(caught.exception))


class WrappedTermTests(unittest.TestCase):
    TEXT = "1380000\n0000\n编号0000\n整理120条活动反馈\n"

    def test_split_short_terms_are_rejected(self):
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["1380000", "0000"], "keep": ["整理120条活动反馈"]}
        with self.assertRaises(engine.Rejected) as caught:
            formats.export(self.TEXT.encode(), "text", plan)
        message = str(caught.exception)
        self.assertNotIn("编号0000", message)
        self.assertNotIn("13800000000", message)
        self.assertIn("完整", message)

    def test_newline_inside_one_term_does_not_change_other_numbers(self):
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["1380000\n0000"], "keep": ["整理120条活动反馈", "编号0000"]}
        out = formats.export(self.TEXT.encode(), "text", plan).decode()
        self.assertIn("编号0000", out)
        self.assertIn("整理120条活动反馈", out)
        self.assertNotIn("1380000\n0000", out)

    def test_web_client_keeps_a_wrapped_selection_as_one_term(self):
        script = Path(__file__).parents[1].joinpath("web", "app.js").read_text(encoding="utf-8")
        self.assertIn("function decodeTerms", script)
        self.assertIn("terms:decodeTerms($('terms').value)", script)
        self.assertNotIn("terms:$('terms').value.split('\\n')", script)
        # 与网页 encode/decode 同一规则：跨行选区写成一行里的 \\n，读回仍是一项。
        shown = "1380000\\n0000\n整理120条活动反馈"
        items = []
        for line in shown.split("\n"):
            out = []
            chars = list(line)
            i = 0
            while i < len(chars):
                if chars[i] == "\\" and i + 1 < len(chars) and chars[i + 1] == "n":
                    out.append("\n")
                    i += 2
                elif chars[i] == "\\" and i + 1 < len(chars) and chars[i + 1] == "\\":
                    out.append("\\")
                    i += 2
                else:
                    out.append(chars[i])
                    i += 1
            text = "".join(out).strip()
            if text:
                items.append(text)
        self.assertEqual(items, ["1380000\n0000", "整理120条活动反馈"])


class MarkerCollisionTests(unittest.TestCase):
    def test_marker_word_uses_a_marker_that_does_not_contain_the_term(self):
        raw = "请已隐去此段 整理120条活动反馈\n".encode()
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["已隐去"], "keep": ["整理120条活动反馈"]}
        out = formats.export(raw, "text", plan).decode()
        self.assertNotIn("已隐去", out)
        self.assertIn("〔已处理A〕", out)
        self.assertIn("整理120条活动反馈", out)

    def test_both_markers_collide_explains_why(self):
        raw = "已隐去并且已处理\n".encode()
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["已隐去", "已处理"], "keep": []}
        with self.assertRaises(engine.Rejected) as caught:
            formats.export(raw, "text", plan)
        message = str(caught.exception)
        self.assertIn("替换标记", message)
        self.assertIn("已隐去", message)
        self.assertNotIn("导出复检失败", message)


class SuggestionAndRepeatTests(unittest.TestCase):
    SAMPLE = "\n".join([
        "联系人张三 Zhang San",
        "半角电话13800000000",
        "空格电话138 0000 0000",
        "连字符电话138-0000-0000",
        "全角电话１３８００００００００",
        "零宽电话138\u200b00000000",
        "邮箱zhangsan@example.com",
        "全角邮箱ｚｈａｎｇｓａｎ＠ｅｘａｍｐｌｅ．ｃｏｍ",
        "身份证110101199001010015",
        "银行卡6228480000000007",
        "订单号13600000000",
        "长订单136000000001234",
        "整理120条活动反馈",
    ]) + "\n"

    def test_suggestions_cover_variants_and_skip_order_numbers(self):
        found = engine.suggest_terms(self.SAMPLE)
        self.assertEqual(set(found), {
            "110101199001010015",
            "138 0000 0000",
            "138-0000-0000",
            "13800000000",
            "138\u200b00000000",
            "6228480000000007",
            "zhangsan@example.com",
            "ｚｈａｎｇｓａｎ＠ｅｘａｍｐｌｅ．ｃｏｍ",
            "１３８００００００００",
        })
        self.assertEqual(found, sorted(found))
        self.assertNotIn("13600000000", found)
        self.assertNotIn("136000000001234", found)
        self.assertNotIn("张三", found)

    def test_short_order_inside_a_longer_number_is_refused(self):
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["13600000000"], "keep": ["整理120条活动反馈"]}
        with self.assertRaises(engine.Rejected) as caught:
            formats.export(self.SAMPLE.encode(), "text", plan)
        message = str(caught.exception)
        self.assertNotIn("136000000001234", message)
        self.assertIn("更长", message)

    def test_repeat_on_redacted_text_is_refused(self):
        raw = "电话13800000000\n整理120条活动反馈\n".encode()
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["13800000000"], "keep": ["整理120条活动反馈"]}
        first, report, _ = policy.export(raw, ".txt", plan)
        self.assertNotIn("13800000000", first.decode())
        self.assertTrue(report["checks"]["selected_strings_absent"])
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(first, ".txt", plan)
        self.assertIn("未在预览", str(caught.exception))

    def test_missing_term_on_fresh_text_still_stops(self):
        plan = {"version": 1, "kind": "text", "reviewed": True, "acknowledge_loss": True, "terms": ["13800000000"], "keep": ["整理120条活动反馈"]}
        with self.assertRaises(engine.Rejected) as caught:
            formats.export("整理120条活动反馈\n".encode(), "text", plan)
        self.assertIn("未在预览", str(caught.exception))


MARKER = "【已隐去1】"
ALT = "〔已处理A〕"
WRONG = "13800000001"
LONG = "138000000001234"


def _office_plan(kind, terms, keep=None):
    return {
        "version": 1,
        "kind": kind,
        "reviewed": True,
        "acknowledge_loss": True,
        "terms": terms,
        "keep": [KEEP] if keep is None else keep,
    }


def _docx(text):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx(rows):
    book = Workbook()
    sheet = book.active
    for index, row in enumerate(rows, start=1):
        for col, value in enumerate(row, start=1):
            sheet.cell(index, col, value)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


class ForgedMarkerTests(unittest.TestCase):
    def _assert_refused(self, export):
        with self.assertRaises(engine.Rejected) as caught:
            export()
        message = str(caught.exception)
        self.assertIn("未在预览", message)
        self.assertNotIn(PHONE, message)

    def test_wrong_phone_with_a_handwritten_marker_is_refused(self):
        body = MARKER + "\n电话" + PHONE + "\n" + KEEP + "\n"
        self._assert_refused(lambda: policy.export(body.encode(), ".txt", _office_plan("text", [WRONG])))
        self._assert_refused(lambda: policy.export(body.encode(), ".csv", _office_plan("csv", [WRONG], [])))
        self._assert_refused(lambda: policy.export(_docx(body), ".docx", _office_plan("word", [WRONG])))
        self._assert_refused(lambda: policy.export(_xlsx([[MARKER], [PHONE], [KEEP]]), ".xlsx", _office_plan("excel", [WRONG])))

    def test_missing_name_with_a_marker_is_refused(self):
        body = MARKER + "\n电话" + PHONE + "\n" + KEEP + "\n"
        self._assert_refused(lambda: policy.export(body.encode(), ".txt", _office_plan("text", ["张三"])))

    def test_zero_width_phone_with_a_marker_is_refused(self):
        body = MARKER + "\n138\u200b00000000\n" + KEEP + "\n"
        self._assert_refused(lambda: policy.export(body.encode(), ".txt", _office_plan("text", [PHONE])))


class LeftoverNumberTests(unittest.TestCase):
    def test_marker_does_not_excuse_a_number_inside_a_longer_one(self):
        body = MARKER + "\n" + LONG + "\n" + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(body.encode(), ".txt", _office_plan("text", [PHONE]))
        self.assertNotIn(LONG, str(caught.exception))
        self.assertTrue(engine.term_remains(body, PHONE))

    def test_standalone_and_prefix_of_a_longer_number_are_refused(self):
        body = "电话" + PHONE + "\n长编号" + LONG + "\n" + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(body.encode(), ".txt", _office_plan("text", [PHONE]))
        self.assertNotIn(LONG, str(caught.exception))

    def test_eight_digit_prefix_left_inside_a_longer_number_is_refused(self):
        short = "13800000"
        longer = "138000009999"
        body = short + "\n" + longer + "\n" + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(body.encode(), ".txt", _office_plan("text", [short]))
        message = str(caught.exception)
        self.assertNotIn(longer, message)
        self.assertNotIn(short, message)
        self.assertTrue(engine.term_remains("【已隐去1】" + longer, short))

    def test_zero_width_sibling_is_still_present(self):
        body = PHONE + "\n138\u200b00000000\n" + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(body.encode(), ".txt", _office_plan("text", [PHONE]))
        self.assertNotIn(PHONE, str(caught.exception))
        self.assertTrue(engine.term_remains("138\u200b00000000", PHONE))


class CoveredAnnotationTests(unittest.TestCase):
    def _pdf(self, draw):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=280)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 140), "13800008888", fontname="helv", fontsize=14)
        page.insert_text((20, 200), "DELETEME", fontname="helv", fontsize=14)
        draw(fitz, page)
        raw = doc.tobytes()
        target = page.search_for("DELETEME")[0]
        doc.close()
        return raw, target

    def _plan_for(self, target):
        box = [(target.x0 - 1) * 2, (target.y0 - 1) * 2, (target.x1 + 1) * 2, (target.y1 + 1) * 2]
        return _office_plan("pdf", []) | {
            "reviewed_pages": [1],
            "pages": [{"page": 1, "rectangles": [box]}],
            "terms": [],
        }

    def test_white_square_covering_text_is_refused(self):
        def draw(fitz, page):
            phone = page.search_for("13800008888")[0]
            annot = page.add_rect_annot(fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2))
            annot.set_colors(stroke=(1, 1, 1), fill=(1, 1, 1))
            annot.set_opacity(1)
            annot.update()

        raw, target = self._pdf(draw)
        plan = self._plan_for(target)
        leaked = pdf.export(raw, plan, _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
            self.assertIn("13800008888", doc[0].get_text())
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), message)
        self.assertIn("第1页", message)
        self.assertIn("去掉", message)

    def test_hidden_note_and_link_away_from_text_are_still_removed(self):
        def draw(fitz, page):
            page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(250, 150, 370, 190), "uri": "https://example.com/note"})
            annot = page.add_text_annot(fitz.Point(340, 40), "隐藏便签甲乙丙")
            annot.set_flags(fitz.PDF_ANNOT_IS_HIDDEN | fitz.PDF_ANNOT_IS_NO_VIEW)
            annot.update()
            # 默认弹出框会压到左侧正文。把它挪到右侧空白，像素不变且不压未选对象时才可以去掉。
            value = page.parent.xref_get_key(page.xref, "Annots")[1]
            tokens = value.replace("[", " ").replace("]", " ").split()
            for index in range(0, len(tokens), 3):
                if index + 2 >= len(tokens) or tokens[index + 2] != "R":
                    continue
                xref = int(tokens[index])
                if page.parent.xref_get_key(xref, "Subtype")[1] == "/Popup":
                    page.parent.xref_set_key(xref, "Rect", "[220 140 390 220]")

        raw, target = self._pdf(draw)
        out, report, _ = policy.export(raw, ".pdf", self._plan_for(target))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        self.assertTrue(report["checks"]["keep_phrases_preserved"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            self.assertIn(KEEP, doc[0].get_text())
            self.assertNotIn("DELETEME", doc[0].get_text())
            self.assertIsNone(doc[0].first_annot)
            self.assertFalse(doc[0].get_links())

    def test_link_over_visible_text_is_removed(self):
        def draw(fitz, page):
            page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(20, 28, 160, 44), "uri": "https://example.com/note"})
            annot = page.add_text_annot(fitz.Point(300, 230), "隐藏便签甲乙丙")
            annot.set_flags(fitz.PDF_ANNOT_IS_HIDDEN | fitz.PDF_ANNOT_IS_NO_VIEW)
            annot.update()

        raw, target = self._pdf(draw)
        out, report, _ext = policy.export(raw, ".pdf", self._plan_for(target))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(KEEP, text)
            self.assertIn("13800008888", text)
            self.assertNotIn("DELETEME", text)
            self.assertFalse(doc[0].get_links())
            self.assertIsNone(doc[0].first_annot)

    def test_freetext_and_stamp_on_blank_are_removed(self):
        fitz = pdf.backend()

        def freetext(_fitz, page):
            page.add_freetext_annot(fitz.Rect(250, 210, 380, 250), "自由文本甲乙丙", fontsize=11, fontname="china-s", fill_color=(1, 1, 0.8), text_color=(0, 0, 0)).update()

        def stamp(_fitz, page):
            page.add_stamp_annot(fitz.Rect(250, 180, 370, 250)).update()

        for draw in (freetext, stamp):
            raw, target = self._pdf(draw)
            out, report, _ext = policy.export(raw, ".pdf", self._plan_for(target))
            self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
            with pdf.backend().open(stream=out, filetype="pdf") as doc:
                text = doc[0].get_text()
                self.assertIn(KEEP, text)
                self.assertIn("13800008888", text)
                self.assertNotIn("DELETEME", text)
                self.assertNotIn("自由文本甲乙丙", text)
                self.assertIsNone(doc[0].first_annot)


class SeparatorAndKeepTests(unittest.TestCase):
    def test_short_term_separated_from_a_label_is_refused(self):
        for sep in ("\n", "\u200b", "\u00a0", "\u3000"):
            body = "编号" + sep + "0000\n" + KEEP + "\n"
            with self.assertRaises(engine.Rejected) as caught:
                policy.export(body.encode(), ".txt", _office_plan("text", ["0000"]))
            message = str(caught.exception)
            self.assertNotIn("编号0000", message)
            self.assertNotIn("0000", message)

    def test_formula_after_exotic_space_is_quoted_and_plain_negative_stays(self):
        for lead in ("\u00a0", "\u3000", "\u2028", "\f", "\v"):
            self.assertEqual(formats.safe_cell(lead + "=1+1"), "'" + lead + "=1+1")
        self.assertEqual(formats.safe_cell("-120"), "-120")
        self.assertEqual(formats.safe_cell(" -120"), " -120")
        raw = "姓名,公式\n张三,\u00a0=1+1\n李四,\u3000=1+1\n".encode()
        out = formats.export(raw, "csv", _office_plan("csv", ["张三", "李四"], []))
        rows = formats.csv_rows(out, "csv")
        self.assertEqual(rows[1][1], "'\u00a0=1+1")
        self.assertEqual(rows[2][1], "'\u3000=1+1")

    def test_keep_does_not_accept_a_quoted_formula_with_trailing_space(self):
        raw = "姓名,公式\n张三,=1+1 \n".encode()
        with self.assertRaises(engine.Rejected) as caught:
            formats.export(raw, "csv", _office_plan("csv", ["张三"], ["=1+1"]))
        self.assertIn("空格", str(caught.exception))

    def test_keep_of_a_new_marker_is_refused(self):
        raw = "请已隐去此段 " + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw.encode(), ".txt", _office_plan("text", ["已隐去"], [ALT]))
        self.assertIn("复检", str(caught.exception))

    def test_existing_alternate_marker_collides(self):
        raw = "已有" + ALT + " 请已隐去此段 " + KEEP + "\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw.encode(), ".txt", _office_plan("text", ["已隐去"]))
        message = str(caught.exception)
        self.assertIn("替换标记", message)
        self.assertNotIn("请已隐去此段", message)

    def test_existing_default_marker_switches_style(self):
        raw = "备注" + MARKER + " 电话" + PHONE + "\n" + KEEP + "\n"
        out = policy.export(raw.encode(), ".txt", _office_plan("text", [PHONE]))[0].decode()
        self.assertIn(ALT, out)
        self.assertIn(MARKER, out)
        self.assertNotIn(PHONE, out)
        self.assertIn(KEEP, out)


COVERED = "13800006661"


def _digits(phone, base):
    return "".join(chr(base + int(ch)) for ch in phone)


def _superscript(phone):
    table = {"0": "\u2070", "1": "\u00b9", "2": "\u00b2", "3": "\u00b3", "4": "\u2074", "5": "\u2075", "6": "\u2076", "7": "\u2077", "8": "\u2078", "9": "\u2079"}
    return "".join(table[ch] for ch in phone)


class DarkCoverTests(unittest.TestCase):
    def _pdf(self, color, cover):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=280)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 130), COVERED, fontname="helv", fontsize=16, color=color)
        page.insert_text((20, 230), "OTHERTEXT", fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        other = page.search_for("OTHERTEXT")[0]
        cover(fitz, page, phone)
        raw = doc.tobytes()
        doc.close()
        return raw, other, phone

    def _plan(self, target):
        return _pdf_plan(_pixel_box(target))

    def _square(self, fill, opacity):
        def draw(fitz, page, phone):
            annot = page.add_rect_annot(fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2))
            annot.set_colors(stroke=fill, fill=fill)
            annot.set_opacity(opacity)
            annot.update()

        return draw

    def test_dark_cover_over_lighter_or_darker_text_is_refused(self):
        colors = {
            "black": (0, 0, 0),
            "red": (0.8, 0.1, 0.1),
            "yellow": (0.85, 0.75, 0.1),
            "gray": (0.45, 0.45, 0.45),
            "light": (0.75, 0.75, 0.75),
            "near_white": (0.92, 0.92, 0.92),
        }
        covers = (
            self._square((0, 0, 0), 1),
            self._square((0, 0, 0), 0.98),
            self._square((0.25, 0.25, 0.25), 1),
        )

        def freetext(fitz, page, phone):
            page.add_freetext_annot(
                fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 4),
                "盖",
                fontsize=12,
                fontname="china-s",
                fill_color=(0, 0, 0),
                text_color=(0, 0, 0),
            ).update()

        covers = covers + (freetext,)
        for name, color in colors.items():
            for draw in covers:
                raw, other, _phone = self._pdf(color, draw)
                leaked = pdf.export(raw, self._plan(other), _verify=False)
                with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
                    self.assertIn(COVERED, doc[0].get_text(), name)
                with self.assertRaises(engine.Rejected, msg=name) as caught:
                    policy.export(raw, ".pdf", self._plan(other))
                message = str(caught.exception)
                self.assertTrue(_cover_words(message), message)
                self.assertNotIn(COVERED, message, name)

    def test_black_square_on_blank_space_is_removed(self):
        def draw(fitz, page, _phone):
            annot = page.add_rect_annot(fitz.Rect(300, 30, 360, 80))
            annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
            annot.set_opacity(1)
            annot.update()

        raw, other, _phone = self._pdf((0, 0, 0), draw)
        out, report, _ext = policy.export(raw, ".pdf", self._plan(other))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(KEEP, text)
            self.assertIn(COVERED, text)
            self.assertNotIn("OTHERTEXT", text)
            self.assertIsNone(doc[0].first_annot)

    def test_boxing_the_covered_number_still_refuses(self):
        raw, _other, phone = self._pdf((0, 0, 0), self._square((0, 0, 0), 1))
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", self._plan(phone))
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), message)
        self.assertIn("盖住了", message)
        self.assertIn("去掉", message)
        self.assertNotIn(COVERED, message)

    def test_highlight_over_visible_text_is_removed(self):
        def draw(fitz, page, _phone):
            found = page.search_for(KEEP)[0]
            page.add_highlight_annot(found).update()

        raw, other, _phone = self._pdf((0, 0, 0), draw)
        out, report, _ext = policy.export(raw, ".pdf", self._plan(other))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(KEEP, text)
            self.assertIn(COVERED, text)
            self.assertNotIn("OTHERTEXT", text)
            self.assertIsNone(doc[0].first_annot)

    def test_thick_line_over_text_is_refused(self):
        def line(fitz, page, phone):
            y = (phone.y0 + phone.y1) / 2
            annot = page.add_line_annot(fitz.Point(phone.x0 - 6, y), fitz.Point(phone.x1 + 18, y))
            annot.set_border(width=12)
            annot.set_colors(stroke=(0, 0, 0))
            annot.update()

        raw, other, _phone = self._pdf((0, 0, 0), line)
        leaked = pdf.export(raw, self._plan(other), _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
            self.assertIn(COVERED, doc[0].get_text())
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", self._plan(other))
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), message)
        self.assertIn("直线", message)
        self.assertIn("去掉", message)
        self.assertNotIn(COVERED, message)

    def test_vector_strokes_under_a_black_square_are_refused(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=280)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 230), "OTHERTEXT", fontname="helv", fontsize=14)
        shape = page.new_shape()
        for dx in range(0, 120, 8):
            shape.draw_line(fitz.Point(30 + dx, 100), fitz.Point(32 + dx, 140))
        shape.finish(color=(0, 0, 0), width=1.5)
        shape.commit()
        annot = page.add_rect_annot(fitz.Rect(24, 92, 160, 148))
        annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
        annot.set_opacity(1)
        annot.update()
        other = page.search_for("OTHERTEXT")[0]
        raw = doc.tobytes()
        doc.close()
        plan = self._plan(other)
        leaked = pdf.export(raw, plan, _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as opened:
            self.assertNotIn(COVERED, opened[0].get_text())
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), message)
        self.assertIn("去掉", message)

    def test_faint_square_and_thin_stroke_leave_the_digits_readable(self):
        def stroke(fitz, page, phone):
            annot = page.add_rect_annot(fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2))
            annot.set_colors(stroke=(0, 0, 0))
            annot.set_border(width=2)
            annot.update()

        raw, other, phone = self._pdf((0, 0, 0), self._square((0, 0, 0), 0.08))
        span, _mean = _crop_luma(formats.pdf_image(raw, 0), [phone.x0 * 2 + 2, phone.y0 * 2 + 2, phone.x1 * 2 - 2, phone.y1 * 2 - 2])
        self.assertGreater(span, 60)
        out, report, _ext = policy.export(raw, ".pdf", self._plan(other))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            self.assertIn(COVERED, doc[0].get_text())
            self.assertNotIn("OTHERTEXT", doc[0].get_text())
        raw, other, phone = self._pdf((0, 0, 0), stroke)
        span, _mean = _crop_luma(formats.pdf_image(raw, 0), [phone.x0 * 2 + 2, phone.y0 * 2 + 2, phone.x1 * 2 - 2, phone.y1 * 2 - 2])
        self.assertGreater(span, 60)
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", self._plan(other))
        message = str(caught.exception)
        self.assertIn("接近字迹的深色", message)
        self.assertNotIn(COVERED, message)


class ShortNumberSideTests(unittest.TestCase):
    def test_right_side_and_hyphen_are_refused(self):
        samples = ("0000合同", "0000 合同", "0000\n合同", "ID-0000", "ID\u20140000", "0000-ID")
        for sample in samples:
            body = sample + "\n" + KEEP + "\n"
            with self.assertRaises(engine.Rejected, msg=sample) as caught:
                policy.export(body.encode(), ".txt", _office_plan("text", ["0000"]))
            message = str(caught.exception)
            self.assertNotIn("0000", message, sample)
            self.assertNotIn("合同", message, sample)
            self.assertNotIn("ID", message, sample)

    def test_standalone_short_number_still_delivers(self):
        body = KEEP + "。\n0000\n120\n"
        out, report, _ext = policy.export(body.encode(), ".txt", _office_plan("text", ["0000"], [KEEP]))
        text = out.decode()
        self.assertNotIn("0000", text)
        self.assertIn("【已隐去1】", text)
        self.assertIn(KEEP, text)
        self.assertTrue(report["checks"]["selected_strings_absent"])

    def test_underscore_slash_and_hyphenated_digits_are_refused(self):
        samples = ("ID_0000", "0000_合同", "ID/0000", "0000/合同", "ID／0000", "0000-1234", "1234-0000")
        for sample in samples:
            body = sample + "\n" + KEEP + "\n"
            with self.assertRaises(engine.Rejected, msg=sample) as caught:
                policy.export(body.encode(), ".txt", _office_plan("text", ["0000"]))
            message = str(caught.exception)
            self.assertIn("整段", message, sample)
            self.assertNotIn("0000", message, sample)
            self.assertNotIn("1234", message, sample)
            self.assertNotIn("合同", message, sample)
            self.assertNotIn("ID", message, sample)

    def test_short_number_beside_words_tells_user_to_select_the_span(self):
        body = "价格 12 元\n活动反馈说明\n"
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(body.encode(), ".txt", _office_plan("text", ["12"], ["活动反馈说明"]))
        message = str(caught.exception)
        self.assertIn("整段", message)
        self.assertIn("一起选上", message)
        self.assertNotIn("12", message)
        self.assertNotIn("价格", message)
        self.assertNotIn("元", message)


def _circled(phone):
    return "".join("\u24ea" if ch == "0" else chr(0x245F + int(ch)) for ch in phone)


class AlternateSpellingTests(unittest.TestCase):
    def _variants(self):
        return [
            "138 0000 0000",
            "138-0000-0000",
            "138\u00a00000\u00a00000",
            "138\u30000000\u30000000",
            "138.0000.0000",
            "138\u00b70000\u00b70000",
            _digits(PHONE, 0x0660),
            _digits(PHONE, 0x06F0),
            _digits(PHONE, 0x1D7CE),
            _superscript(PHONE),
            _circled(PHONE),
            "138\u20090000\u20090000",
            "138\u20070000\u20070000",
            "138\u202f0000\u202f0000",
            "138,0000,0000",
            "138/0000/0000",
            "138_0000_0000",
            "138 \u200b0000 \u200b0000",
            "1380000\ufe0e0000",
            "138\ufe0f00000000",
            "138" + "\U000E0100" + "00000000",
            "138\u20dd00000000",
            "138\u0300" + "00000000",
        ]

    def test_other_spelling_of_the_same_number_is_refused(self):
        for variant in self._variants():
            body = "电话" + PHONE + "\n另一种" + variant + "\n" + KEEP + "\n"
            with self.assertRaises(engine.Rejected) as caught:
                policy.export(body.encode(), ".txt", _office_plan("text", [PHONE]))
            message = str(caught.exception)
            self.assertIn("另一种写法", message)
            self.assertNotIn(PHONE, message)
            self.assertNotIn(variant, message)

    def test_a_different_number_can_remain_and_is_not_called_absent(self):
        other = "139 0000 0000"
        body = "电话" + PHONE + "\n其它" + other + "\n" + KEEP + "\n"
        out, report, _ext = policy.export(body.encode(), ".txt", _office_plan("text", [PHONE]))
        text = out.decode()
        self.assertNotIn(PHONE, text)
        self.assertIn(other, text)
        self.assertTrue(report["checks"]["selected_strings_absent"])
        self.assertIn(KEEP, text)

    def test_selecting_only_the_other_spelling_still_refuses(self):
        for variant in self._variants():
            body = "电话" + PHONE + "\n另一种" + variant + "\n" + KEEP + "\n"
            with self.assertRaises(engine.Rejected) as caught:
                policy.export(body.encode(), ".txt", _office_plan("text", [variant]))
            message = str(caught.exception)
            self.assertIn("另一种写法", message)
            self.assertNotIn(PHONE, message)
            self.assertNotIn(variant, message)


class PdfKeepMessageTests(unittest.TestCase):
    def test_missing_keep_names_the_keep_item(self):
        raw, found = _page(note=False)
        plan = _pdf_plan(_pixel_box(found))
        plan["keep"] = ["原文里没有的保留句"]
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertIn("保留项", message)
        self.assertIn("没有对上", message)
        self.assertNotIn("框外", message)
        self.assertNotIn("原文里没有的保留句", message)


class RenderGeometryTests(unittest.TestCase):
    def _refuse_covered(self, raw, plan):
        leaked = pdf.export(raw, plan, _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
            self.assertIn(COVERED, doc[0].get_text())
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), message)
        self.assertIn("去掉", message)
        self.assertNotIn(COVERED, message)

    def test_light_text_under_a_black_square_is_refused(self):
        colors = (
            (0.975, 0.975, 0.975),
            (0.980, 0.980, 0.980),
            (0.985, 0.985, 0.985),
            (1, 0.97, 0.97),
        )
        fitz = pdf.backend()
        for color in colors:
            doc = fitz.open()
            page = doc.new_page(width=400, height=300)
            page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
            page.insert_text((20, 90), COVERED, fontname="helv", fontsize=16, color=color)
            page.insert_text((20, 220), "DELETEME", fontname="helv", fontsize=14)
            phone = page.search_for(COVERED)[0]
            other = page.search_for("DELETEME")[0]
            annot = page.add_rect_annot(fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2))
            annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
            annot.set_opacity(1)
            annot.update()
            raw = doc.tobytes()
            doc.close()
            self._refuse_covered(raw, _pdf_plan(_pixel_box(other)))

    def test_text_on_each_edge_of_a_square_is_refused(self):
        fitz = pdf.backend()
        square = fitz.Rect(30, 140, 340, 270)
        spots = {
            "top": (48, 141.2),
            "bottom": (48, 269.2),
            "left": (30.6, 200),
            "right": (338.4, 200),
        }
        for name, point in spots.items():
            doc = fitz.open()
            page = doc.new_page(width=400, height=300)
            page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
            page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
            page.insert_text(point, COVERED, fontname="helv", fontsize=1)
            other = page.search_for("DELETEME")[0]
            annot = page.add_rect_annot(square)
            annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
            annot.set_opacity(1)
            annot.update()
            raw = doc.tobytes()
            doc.close()
            self._refuse_covered(raw, _pdf_plan(_pixel_box(other)))

    def _covered_geometry(self, angle=0, crop=None, userunit=None, mediabox=None, word="DELETEME"):
        fitz = pdf.backend()
        doc = fitz.open()
        if crop is None:
            page = doc.new_page(width=400, height=300)
            origin = (24, 90)
        else:
            page = doc.new_page(width=500, height=400)
            page.set_cropbox(fitz.Rect(*crop))
            origin = (80, 110)
        if mediabox is not None:
            page.set_mediabox(fitz.Rect(*mediabox))
        if userunit is not None:
            page.parent.xref_set_key(page.xref, "UserUnit", str(userunit))
        page.insert_text((origin[0], origin[1] - 50), KEEP, fontname="china-s", fontsize=14)
        page.insert_text(origin, COVERED, fontname="helv", fontsize=14)
        page.insert_text((origin[0], origin[1] + 100), "DELETEME", fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        chosen = page.search_for(word)[0]
        annot = page.add_rect_annot(fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2))
        annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
        annot.set_opacity(1)
        annot.update()
        if angle:
            page.set_rotation(angle)
        # 先在页面还开着时算出批注的文字坐标和显示矩阵。关掉后再用 pdfium，不能回头读 annot.rect。
        area = pdf._annot_area(page.first_annot, page)
        matrix = pdf._display_matrix(page)
        shown = fitz.Rect(page.rect)
        box = _preview_box(page, chosen)
        raw = doc.tobytes()
        doc.close()
        image = formats.pdf_image(raw, 0)
        sx = image.size[0] / shown.width
        sy = image.size[1] / shown.height
        mapped = fitz.Rect(area) * matrix
        cover = (
            max(0, math.floor(mapped.x0 * sx)),
            max(0, math.floor(mapped.y0 * sy)),
            min(image.width, math.ceil(mapped.x1 * sx)),
            min(image.height, math.ceil(mapped.y1 * sy)),
        )
        self.assertGreater(cover[2] - cover[0], 1)
        self.assertGreater(cover[3] - cover[1], 1)
        _span, mean = _crop_luma(image, cover)
        self.assertLess(mean, 40, cover)
        return raw, _pdf_plan(box)

    def test_rotated_black_square_uses_the_same_transform(self):
        for angle in (0, 90, 180, 270):
            raw, plan = self._covered_geometry(angle=angle)
            self._refuse_covered(raw, plan)

    def test_crop_offset_and_user_unit_use_the_same_scale(self):
        cases = (
            {"crop": (40, 20, 440, 320)},
            {"userunit": 2},
            {"userunit": 2, "angle": 90},
            {"userunit": 2, "angle": 180},
            {"userunit": 2, "angle": 270},
            {"mediabox": (100, 50, 500, 350)},
        )
        for case in cases:
            raw, plan = self._covered_geometry(**case)
            self._refuse_covered(raw, plan)

    def test_user_unit_two_without_annotation_redacts_the_number(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.parent.xref_set_key(page.xref, "UserUnit", "2")
        page.insert_text((24, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((24, 240), "KEEPME", fontname="helv", fontsize=14)
        page.insert_text((24, 90), COVERED, fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        box = _preview_box(page, phone)
        raw = doc.tobytes()
        doc.close()
        out, report, _ext = policy.export(raw, ".pdf", _pdf_plan(box))
        self.assertTrue(report["checks"]["selected_text_removed"])
        self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
        self.assertTrue(report["checks"]["keep_phrases_preserved"])
        with pdf.backend().open(stream=out, filetype="pdf") as opened:
            text = opened[0].get_text()
            self.assertNotIn(COVERED, text)
            self.assertIn(KEEP, text)
            self.assertIn("KEEPME", text)

    def test_user_unit_black_square_still_names_the_cover(self):
        raw, plan = self._covered_geometry(userunit=2, word=COVERED)
        leaked = pdf.export(raw, plan, _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
            self.assertNotIn(COVERED, doc[0].get_text())
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertIn("盖住了", message)
        self.assertIn("方形", message)
        self.assertNotIn("可能切到字形", message)
        self.assertNotIn(COVERED, message)

    def test_crop_and_rotation_together_stay_refused(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=500, height=400)
        page.set_cropbox(fitz.Rect(40, 20, 440, 320))
        page.insert_text((80, 60), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((80, 140), "DELETEME", fontname="helv", fontsize=14)
        page.set_rotation(90)
        other = page.search_for("DELETEME")[0]
        box = _preview_box(page, other)
        raw = doc.tobytes()
        doc.close()
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", _pdf_plan(box))
        message = str(caught.exception)
        self.assertIn("裁剪框和旋转", message)
        self.assertNotIn(COVERED, message)

    def test_header_and_table_shading_still_deliver(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        shape = page.new_shape()
        shape.draw_rect(fitz.Rect(0, 0, 400, 36))
        shape.finish(fill=(0.12, 0.14, 0.2), color=None, width=0)
        shape.commit()
        shape = page.new_shape()
        shape.draw_rect(fitz.Rect(30, 70, 200, 110))
        shape.finish(fill=(0.93, 0.93, 0.86), color=(0.7, 0.7, 0.7), width=0.6)
        shape.commit()
        page.insert_text((36, 96), "CELL", fontname="helv", fontsize=12)
        page.insert_text((20, 170), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((20, 230), KEEP, fontname="china-s", fontsize=14)
        other = page.search_for("DELETEME")[0]
        raw = doc.tobytes()
        doc.close()
        out, report, _ext = policy.export(raw, ".pdf", _pdf_plan(_pixel_box(other)))
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
        self.assertTrue(report["checks"]["keep_phrases_preserved"])
        with pdf.backend().open(stream=out, filetype="pdf") as opened:
            text = opened[0].get_text()
            self.assertIn("CELL", text)
            self.assertIn(KEEP, text)
            self.assertNotIn("DELETEME", text)
            self.assertIsNone(opened[0].first_annot)
        image = formats.pdf_image(out, 0).convert("L")
        self.assertLess(image.getpixel((20, 20)), 80)
        self.assertLess(image.getpixel((320, 160)), 250)
        self.assertGreater(image.getpixel((320, 160)), 180)

    def test_ink_and_file_attachment_name_their_own_marks(self):
        fitz = pdf.backend()

        def ink(page):
            page.add_ink_annot([[(280, 40), (350, 70), (300, 110)]]).update()

        def attachment(page):
            page.add_file_annot(fitz.Point(300, 60), b"note", filename="a.txt").update()

        for draw in (ink, attachment):
            doc = fitz.open()
            page = doc.new_page(width=400, height=280)
            page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
            page.insert_text((20, 200), "DELETEME", fontname="helv", fontsize=14)
            draw(page)
            other = page.search_for("DELETEME")[0]
            raw = doc.tobytes()
            doc.close()
            out, report, _ext = policy.export(raw, ".pdf", _pdf_plan(_pixel_box(other)))
            self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
            with pdf.backend().open(stream=out, filetype="pdf") as doc:
                text = doc[0].get_text()
                self.assertIn(KEEP, text)
                self.assertNotIn("DELETEME", text)
                self.assertIsNone(doc[0].first_annot)

    def test_one_point_glyph_in_the_selection_halo_is_not_left_readable(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=280)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 140), "DELETEME", fontname="helv", fontsize=14)
        target = page.search_for("DELETEME")[0]
        box = fitz.Rect(target.x0 - 1, target.y0 - 1, target.x1 + 1, target.y1 + 1)
        # 字形框跨进选区，墨迹主要落在框外两像素里。涂黑会删掉这个字，交付后正文里没有它。
        page.insert_text((box.x1 - 0.4, box.y1 - 0.4), "9", fontname="helv", fontsize=1)
        raw = doc.tobytes()
        doc.close()
        plan = _pdf_plan([box.x0 * 2, box.y0 * 2, box.x1 * 2, box.y1 * 2])
        out, report, _ext = policy.export(raw, ".pdf", plan)
        self.assertTrue(report["checks"]["selected_text_removed"])
        with pdf.backend().open(stream=out, filetype="pdf") as opened:
            self.assertNotIn("9", opened[0].get_text())
            self.assertIn(KEEP, opened[0].get_text())
        outside = fitz.open()
        page = outside.new_page(width=400, height=280)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 140), "DELETEME", fontname="helv", fontsize=14)
        target = page.search_for("DELETEME")[0]
        box = fitz.Rect(target.x0 - 1, target.y0 - 1, target.x1 + 1, target.y1 + 1)
        # 完全落在框外的同一点字仍留下，不因为晕圈被删掉。
        page.insert_text((box.x1 + 0.6, box.y1 - 0.4), "9", fontname="helv", fontsize=1)
        raw = outside.tobytes()
        outside.close()
        out, _report, _ext = policy.export(raw, ".pdf", plan)
        with pdf.backend().open(stream=out, filetype="pdf") as opened:
            self.assertIn("9", opened[0].get_text())
            self.assertIn(KEEP, opened[0].get_text())


def _crop_luma(image, box):
    x0, y0, x1, y1 = (int(v) for v in box)
    crop = image.crop((x0, y0, x1, y1)).convert("L")
    low, high = crop.getextrema()
    raw = crop.tobytes()
    return high - low, sum(raw) / len(raw)


def _page_plan(box, keep=True):
    plan = _pdf_plan(box)
    if not keep:
        plan["keep"] = []
    return plan


class MarkupAppearanceTests(unittest.TestCase):
    """批注看底下的字还在不在。类型名、颜色和线宽都不能放行。"""

    def _covered(self, border=None):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((20, 180), COVERED, fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        other = page.search_for("DELETEME")[0]
        rect = fitz.Rect(phone.x0 - 2, phone.y0 - 2, phone.x1 + 2, phone.y1 + 2)
        annot = page.add_rect_annot(rect)
        annot.set_colors(stroke=(0, 0, 0), fill=(0, 0, 0))
        annot.set_opacity(1)
        if border is not None:
            annot.set_border(width=border)
        annot.update()
        raw = doc.tobytes()
        doc.close()
        return raw, other, phone, rect

    def _rename(self, raw, subtype):
        fitz = pdf.backend()
        doc = fitz.open(stream=raw, filetype="pdf")
        page = doc[0]
        value = doc.xref_get_key(page.xref, "Annots")[1]
        tokens = value.replace("[", " ").replace("]", " ").split()
        refs = [int(tokens[i]) for i in range(0, len(tokens), 3) if tokens[i + 2] == "R"]
        self.assertEqual(len(refs), 1, value)
        doc.xref_set_key(refs[0], "Subtype", subtype)
        renamed = doc.tobytes()
        doc.close()
        check = fitz.open(stream=renamed, filetype="pdf")
        page = check[0]
        value = check.xref_get_key(page.xref, "Annots")[1]
        tokens = value.replace("[", " ").replace("]", " ").split()
        xref = int(tokens[0])
        self.assertEqual(check.xref_get_key(xref, "Subtype")[1], subtype)
        check.close()
        return renamed

    def _phone_pixels(self, phone):
        return [phone.x0 * 2 + 2, phone.y0 * 2 + 2, phone.x1 * 2 - 2, phone.y1 * 2 - 2]

    def test_pdf_user_rect_lands_on_the_annot_rect(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        rect = fitz.Rect(20, 160, 120, 190)
        annot = page.add_rect_annot(rect)
        annot.set_colors(fill=(0, 0, 0), stroke=(0, 0, 0))
        annot.update()
        kind, value = doc.xref_get_key(annot.xref, "Rect")
        self.assertEqual(kind, "array")
        pdf_rect = tuple(float(token) for token in value.replace("[", " ").replace("]", " ").split())
        bound = pdf._bound_from_pdf_rect(page, pdf_rect)
        self.assertLess(abs(bound.x0 - annot.rect.x0), 0.2)
        self.assertLess(abs(bound.y0 - annot.rect.y0), 0.2)
        self.assertLess(abs(bound.x1 - annot.rect.x1), 0.2)
        self.assertLess(abs(bound.y1 - annot.rect.y1), 0.2)
        doc.close()

    def test_black_markup_appearance_is_refused(self):
        raw, other, phone, _rect = self._covered()
        plan = _page_plan(_pixel_box(other))
        phone_box = self._phone_pixels(phone)
        for subtype in ("/Highlight", "/Underline", "/StrikeOut", "/Squiggly", "/Text", "/Link"):
            renamed = self._rename(raw, subtype)
            image = formats.pdf_image(renamed, 0)
            span, mean = _crop_luma(image, phone_box)
            self.assertLess(span, 8, subtype)
            self.assertLess(mean, 40, subtype)
            leaked = pdf.export(renamed, plan, _verify=False)
            with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
                self.assertIn(COVERED, doc[0].get_text(), subtype)
            with self.assertRaises(engine.Rejected, msg=subtype) as caught:
                policy.export(renamed, ".pdf", plan)
            message = str(caught.exception)
            self.assertTrue(_cover_words(message), message)
            self.assertNotIn(COVERED, message, subtype)
            self.assertNotIn("DELETEME", message, subtype)

    def test_unpainted_popup_over_visible_text_is_removed(self):
        raw, other, _phone, _rect = self._covered()
        renamed = self._rename(raw, "/Popup")
        plan = _page_plan(_pixel_box(other))
        out, report, _ext = policy.export(renamed, ".pdf", plan)
        self.assertTrue(report["checks"]["covered_text_not_revealed"])
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(COVERED, text)
            self.assertIn(KEEP, text)
            self.assertNotIn("DELETEME", text)
            self.assertIsNone(doc[0].first_annot)

    def test_highlight_underline_squiggly_and_note_are_removed_and_strikeout_is_refused(self):
        fitz = pdf.backend()

        def highlight(page):
            page.add_highlight_annot(page.search_for(KEEP)[0]).update()

        def underline(page):
            page.add_underline_annot(page.search_for(KEEP)[0]).update()

        def strikeout(page):
            page.add_strikeout_annot(page.search_for(KEEP)[0]).update()

        def squiggly(page):
            page.add_squiggly_annot(page.search_for(KEEP)[0]).update()

        def note(page):
            page.add_text_annot(fitz.Point(320, 50), "便签甲乙丙").update()

        for draw in (highlight, underline, squiggly, note):
            doc = fitz.open()
            page = doc.new_page(width=400, height=300)
            page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
            page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
            page.insert_text((20, 180), COVERED, fontname="helv", fontsize=14)
            draw(page)
            other = page.search_for("DELETEME")[0]
            raw = doc.tobytes()
            doc.close()
            out, report, _ext = policy.export(raw, ".pdf", _page_plan(_pixel_box(other)))
            self.assertTrue(report["checks"]["covered_text_not_revealed"])
            with pdf.backend().open(stream=out, filetype="pdf") as doc:
                text = doc[0].get_text()
                self.assertIn(KEEP, text)
                self.assertIn(COVERED, text)
                self.assertNotIn("DELETEME", text)
                self.assertIsNone(doc[0].first_annot)
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((20, 180), COVERED, fontname="helv", fontsize=14)
        strikeout(page)
        other = page.search_for("DELETEME")[0]
        raw = doc.tobytes()
        doc.close()
        with self.assertRaises(engine.Rejected) as caught:
            policy.export(raw, ".pdf", _page_plan(_pixel_box(other)))
        message = str(caught.exception)
        self.assertIn("接近字迹的深色", message)
        self.assertIn("删除线", message)
        self.assertNotIn(COVERED, message)

    def test_zero_width_black_square_is_refused_even_when_boxed(self):
        raw, other, _phone, rect = self._covered(border=0)
        width, height = formats.pdf_info(raw)[0]
        plans = (
            _page_plan(_pixel_box(other)),
            _page_plan(_pixel_box(rect)),
            _page_plan([0, 0, width, height], keep=False),
        )
        for plan in plans:
            with self.assertRaises(engine.Rejected) as caught:
                policy.export(raw, ".pdf", plan)
            message = str(caught.exception)
            self.assertIn("盖住了", message)
            self.assertIn("去掉", message)
            self.assertNotIn("画面变了", message)
            self.assertNotIn(COVERED, message)

    def test_stamp_on_blank_is_removed_even_when_fully_boxed(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((20, 180), COVERED, fontname="helv", fontsize=14)
        stamp = fitz.Rect(250, 200, 360, 270)
        page.add_stamp_annot(stamp).update()
        other = page.search_for("DELETEME")[0]
        raw = doc.tobytes()
        doc.close()
        width, height = formats.pdf_info(raw)[0]
        plans = (
            _page_plan(_pixel_box(other)),
            _page_plan(_pixel_box(stamp)),
            _page_plan([0, 0, width, height], keep=False),
        )
        for plan in plans:
            out, report, _ext = policy.export(raw, ".pdf", plan)
            self.assertTrue(report["checks"]["outside_page_pixels_preserved"])
            with pdf.backend().open(stream=out, filetype="pdf") as doc:
                text = doc[0].get_text()
                self.assertIsNone(doc[0].first_annot)
                if plan["keep"]:
                    self.assertIn(KEEP, text)
                    self.assertIn(COVERED, text)


class StrictPixelCompareTests(unittest.TestCase):
    """白边灰底和其它真遮盖仍然拒绝。不按类型、颜色或线宽放行。"""

    def _assert_refused(self, raw, plan, label):
        leaked = pdf.export(raw, plan, _verify=False)
        with pdf.backend().open(stream=leaked, filetype="pdf") as doc:
            self.assertIn(COVERED, doc[0].get_text(), label)
        with self.assertRaises(engine.Rejected, msg=label) as caught:
            policy.export(raw, ".pdf", plan)
        message = str(caught.exception)
        self.assertTrue(_cover_words(message), label + "\n" + message)
        self.assertIn("第1页", message, label)
        self.assertIn("去掉", message, label)
        self.assertNotIn(COVERED, message, label)
        self.assertNotIn("DELETEME", message, label)

    def _interior(self, phone):
        return [phone.x0 * 2 + 4, phone.y0 * 2 + 4, phone.x1 * 2 - 4, phone.y1 * 2 - 4]

    def _base(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.insert_text((20, 40), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((20, 90), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((20, 180), COVERED, fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        other = page.search_for("DELETEME")[0]
        return fitz, doc, page, phone, other

    def _square(self, page, phone, gray, border, opacity=1, fill=True, pad=8):
        rect = pdf.backend().Rect(phone.x0 - pad, phone.y0 - pad, phone.x1 + pad, phone.y1 + pad)
        annot = page.add_rect_annot(rect)
        if fill:
            annot.set_colors(stroke=(1, 1, 1), fill=(gray, gray, gray))
        else:
            annot.set_colors(stroke=(gray, gray, gray))
        annot.set_border(width=border)
        annot.set_opacity(opacity)
        annot.update()
        return annot

    def _forge(self, raw, subtype):
        fitz = pdf.backend()
        doc = fitz.open(stream=raw, filetype="pdf")
        value = doc.xref_get_key(doc[0].xref, "Annots")[1]
        tokens = value.replace("[", " ").replace("]", " ").split()
        refs = [int(tokens[i]) for i in range(0, len(tokens), 3) if i + 2 < len(tokens) and tokens[i + 2] == "R"]
        self.assertEqual(len(refs), 1, value)
        doc.xref_set_key(refs[0], "Subtype", subtype)
        renamed = doc.tobytes()
        doc.close()
        return renamed

    def _appearance_xref(self, doc, annot):
        value = doc.xref_get_key(annot.xref, "AP")[1]
        tokens = value.replace("<<", " ").replace(">>", " ").split()
        return int(tokens[tokens.index("/N") + 1])

    def _paint(self, doc, annot, kind):
        xref = self._appearance_xref(doc, annot)
        _kind, bbox = doc.xref_get_key(xref, "BBox")
        x0, y0, x1, y1 = (float(token) for token in bbox.replace("[", " ").replace("]", " ").split())
        width = x1 - x0
        height = y1 - y0
        parts = []
        if kind == "gradient":
            steps = 8
            band = height / steps
            for index in range(steps):
                level = 0.2 + 0.65 * index / (steps - 1)
                parts.append(f"{level:.3f} {level:.3f} {level:.3f} rg")
                parts.append(f"{x0:.3f} {y0 + index * band:.3f} {width:.3f} {band + 0.4:.3f} re f")
        elif kind == "two-color":
            mid = y0 + height / 2
            parts.append(f"0.280 0.280 0.280 rg {x0:.3f} {y0:.3f} {width:.3f} {height / 2 + 0.3:.3f} re f")
            parts.append(f"0.720 0.450 0.300 rg {x0:.3f} {mid:.3f} {width:.3f} {height / 2 + 0.3:.3f} re f")
        elif kind == "texture":
            cell = 3.5
            row = 0
            top = y0
            while top < y1 - 0.05:
                col = 0
                left = x0
                while left < x1 - 0.05:
                    level = 0.22 if (row + col) % 2 == 0 else 0.78
                    box_w = min(cell, x1 - left)
                    box_h = min(cell, y1 - top)
                    parts.append(f"{level:.2f} {level:.2f} {level:.2f} rg {left:.2f} {top:.2f} {box_w:.2f} {box_h:.2f} re f")
                    left += cell
                    col += 1
                top += cell
                row += 1
        else:
            raise AssertionError(kind)
        parts.append(f"1 1 1 RG\n2 w\n{x0:.3f} {y0:.3f} {width:.3f} {height:.3f} re S")
        doc.update_stream(xref, "\n".join(parts).encode("ascii"))

    def test_opaque_gray_white_border_covers_are_refused(self):
        samples = []
        for kind in ("square", "circle", "polygon"):
            for gray in (0.4, 0.7):
                for border in (1, 6):
                    samples.append((kind, gray, border, None))
        for subtype in ("/FreeText", "/Stamp", "/Highlight", "/Ink"):
            for gray in (0.4, 0.7):
                for border in (1, 6):
                    samples.append(("square", gray, border, subtype))
        samples.extend((
            ("square", 0.55, 3, None),
            ("square", 0.85, 8, None),
            ("square", 0.92, 2, "/Highlight"),
        ))
        fitz = pdf.backend()
        for kind, gray, border, subtype in samples:
            label = f"{subtype or kind} gray={gray} border={border}"
            _fitz, doc, page, phone, other = self._base()
            if kind == "square":
                self._square(page, phone, gray, border)
            elif kind == "circle":
                center_x = (phone.x0 + phone.x1) / 2
                center_y = (phone.y0 + phone.y1) / 2
                radius = max(phone.width, phone.height) / 2 + 18
                annot = page.add_circle_annot(fitz.Rect(center_x - radius, center_y - radius, center_x + radius, center_y + radius))
                annot.set_colors(stroke=(1, 1, 1), fill=(gray, gray, gray))
                annot.set_border(width=border)
                annot.set_opacity(1)
                annot.update()
            else:
                rect = fitz.Rect(phone.x0 - 8, phone.y0 - 8, phone.x1 + 8, phone.y1 + 8)
                annot = page.add_polygon_annot([rect.tl, rect.tr, rect.br, rect.bl])
                annot.set_colors(stroke=(1, 1, 1), fill=(gray, gray, gray))
                annot.set_border(width=border)
                annot.set_opacity(1)
                annot.update()
            raw = doc.tobytes()
            doc.close()
            if subtype:
                raw = self._forge(raw, subtype)
            span, mean = _crop_luma(formats.pdf_image(raw, 0), self._interior(phone))
            self.assertLess(span, 30, label)
            self.assertGreater(mean, 50, label)
            self._assert_refused(raw, _page_plan(_pixel_box(other)), label)

    def test_gradient_two_color_and_texture_covers_are_refused(self):
        for kind in ("gradient", "two-color", "texture"):
            _fitz, doc, page, phone, other = self._base()
            annot = self._square(page, phone, 0.5, 2)
            self._paint(doc, annot, kind)
            raw = doc.tobytes()
            doc.close()
            self._assert_refused(raw, _page_plan(_pixel_box(other)), kind)

    def _assert_still_readable(self, raw, plan, label):
        with pdf.backend().open(stream=raw, filetype="pdf") as opened:
            phone = opened[0].search_for(COVERED)[0]
            interior = self._interior(phone)
        span, _mean = _crop_luma(formats.pdf_image(raw, 0), interior)
        self.assertGreater(span, 60, label)
        out, report, _ext = policy.export(raw, ".pdf", plan)
        self.assertTrue(report["checks"]["covered_text_not_revealed"], label)
        with pdf.backend().open(stream=out, filetype="pdf") as doc:
            text = doc[0].get_text()
            self.assertIn(COVERED, text, label)
            self.assertNotIn("DELETEME", text, label)
            self.assertIsNone(doc[0].first_annot, label)

    def test_translucent_cover_is_refused_and_strokes_that_leave_digits_deliver(self):
        _fitz, doc, page, phone, other = self._base()
        self._square(page, phone, 0.45, 3, opacity=0.45)
        raw = doc.tobytes()
        doc.close()
        self._assert_refused(raw, _page_plan(_pixel_box(other)), "translucent")

        _fitz, doc, page, phone, other = self._base()
        self._square(page, phone, 0.25, 4, fill=False)
        raw = doc.tobytes()
        doc.close()
        self._assert_still_readable(raw, _page_plan(_pixel_box(other)), "stroke")

        _fitz, doc, page, phone, other = self._base()
        self._square(page, phone, 1, 3, fill=False)
        raw = doc.tobytes()
        doc.close()
        self._assert_still_readable(raw, _page_plan(_pixel_box(other)), "white-stroke")

    def test_rotated_gray_border_is_refused(self):
        for angle in (90, 180, 270):
            _fitz, doc, page, phone, other = self._base()
            self._square(page, phone, 0.45, 4)
            page.set_rotation(angle)
            box = _preview_box(page, other)
            raw = doc.tobytes()
            doc.close()
            self._assert_refused(raw, _page_plan(box), f"rotation-{angle}")

    def test_cropbox_offset_gray_border_is_refused(self):
        fitz = pdf.backend()
        doc = fitz.open()
        page = doc.new_page(width=500, height=400)
        page.set_cropbox(fitz.Rect(40, 20, 440, 320))
        page.insert_text((80, 80), KEEP, fontname="china-s", fontsize=14)
        page.insert_text((80, 150), "DELETEME", fontname="helv", fontsize=14)
        page.insert_text((80, 230), COVERED, fontname="helv", fontsize=14)
        phone = page.search_for(COVERED)[0]
        other = page.search_for("DELETEME")[0]
        self._square(page, phone, 0.5, 5)
        box = _preview_box(page, other)
        raw = doc.tobytes()
        doc.close()
        self._assert_refused(raw, _page_plan(box), "crop-offset")


if __name__ == "__main__":
    unittest.main()
