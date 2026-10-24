"""安全边界和材料保留的回归，全部输入在内存中合成。"""
import io
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1]))
from docx import Document
from PIL import Image
import engine


def document_bytes():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("林芩")
    p.add_run("澄 整理120条活动反馈")
    doc.add_paragraph("课程练习中做过筛选与连接，尚无业务取数经历。")
    doc.sections[0].header.paragraphs[0].text = "qin.test@example.invalid"
    doc.core_properties.author = "林芩澄"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class EngineTests(unittest.TestCase):
    def test_word_keeps_business_and_removes_cross_run_and_header_values(self):
        source = document_bytes()
        output, report = engine.export_docx(source, ["林芩澄", "qin.test@example.invalid"], True)
        self.assertTrue(all(report["checks"].values()))
        doc = Document(io.BytesIO(output))
        text = "".join(p.text for p in doc.paragraphs)
        self.assertIn("整理120条活动反馈", text)
        self.assertIn("尚无业务取数经历", text)
        self.assertNotIn("林芩澄", text)
        self.assertEqual(doc.core_properties.author, "")

    def test_word_requires_explicit_loss_acknowledgement(self):
        with self.assertRaises(engine.Rejected):
            engine.export_docx(document_bytes(), ["林芩澄"])

    def test_unsupported_textbox_stops_before_export(self):
        src = document_bytes()
        with zipfile.ZipFile(io.BytesIO(src)) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        root = engine.ET.fromstring(parts['word/document.xml'])
        engine.ET.SubElement(root.find(engine.W + 'body')[0], engine.W + 'txbxContent')
        parts['word/document.xml'] = engine.ET.tostring(root)
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w') as z:
            for n, b in parts.items(): z.writestr(n, b)
        with self.assertRaises(engine.Rejected):
            engine.inspect_docx(out.getvalue())

    def test_broken_original_relationship_stops(self):
        with zipfile.ZipFile(io.BytesIO(document_bytes())) as z:
            parts = {n: z.read(n) for n in z.namelist() if n != 'word/styles.xml'}
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w') as z:
            for n, b in parts.items(): z.writestr(n, b)
        with self.assertRaises(engine.Rejected): engine.inspect_docx(out.getvalue())

    def test_image_removes_metadata_and_preserves_unselected_pixels(self):
        im = Image.new('RGB', (100, 100), 'white')
        exif = Image.Exif(); exif[270] = 'qin.test@example.invalid'
        src = io.BytesIO(); im.save(src, format='JPEG', exif=exif)
        output, report = engine.export_image(src.getvalue(), [[5, 5, 20, 20]])
        self.assertTrue(all(report['checks'].values()))
        with Image.open(io.BytesIO(output)) as check:
            self.assertEqual(check.getpixel((10, 10)), (0, 0, 0))
            self.assertEqual(check.getpixel((30, 30)), (255, 255, 255))
            self.assertEqual(dict(check.getexif()), {})

    def test_bad_rectangles_and_multiframe_images_rejected(self):
        for value in [[], [[0, 0, 101, 20]], [[0, 0, float('nan'), 20]]]:
            with self.assertRaises(engine.Rejected): engine.rectangles(value, (100, 100))
        out = io.BytesIO()
        a = Image.new('RGB', (10, 10), 'red'); b = Image.new('RGB', (10, 10), 'blue')
        a.save(out, format='PNG', save_all=True, append_images=[b])
        with self.assertRaises(engine.Rejected): engine.image_source(out.getvalue())


if __name__ == '__main__': unittest.main()
