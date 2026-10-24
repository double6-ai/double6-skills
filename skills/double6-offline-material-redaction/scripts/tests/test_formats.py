"""扩展格式：用合成材料检查真实内容及隐藏载体，不用文案断言代替结果。"""
import io,json,sys,zipfile,tempfile,subprocess
from pathlib import Path
import unittest
from PIL import Image
from openpyxl import Workbook,load_workbook
from openpyxl.comments import Comment
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader,PdfWriter
sys.path.insert(0,str(Path(__file__).parents[1]))
import formats,engine,cli

SECRET='林测试';KEEP='业务数字120保留'

def plan(kind,**kwargs):return {'version':1,'kind':kind,'reviewed':True,'terms':[SECRET],'acknowledge_loss':True,'keep':[KEEP],**kwargs}

def workbook():
 w=Workbook();s=w.active;s.title='机构私密原名';s.append([SECRET,KEEP,120,'=1+1']);s['D1']='ordinary';s['A1'].comment=Comment('COMMENT_PRIVATE','AUTHOR_PRIVATE');s['B1'].hyperlink='https://private.invalid/private';s.append(['HIDDEN_ROW_PRIVATE']);s.row_dimensions[2].hidden=True;s['F1']='HIDDEN_COL_PRIVATE';s.column_dimensions['F'].hidden=True;h=w.create_sheet('HIDDEN_SHEET_PRIVATE');h['A1']='HIDDEN_VALUE_PRIVATE';h.sheet_state='veryHidden';w.properties.creator='AUTHOR_PRIVATE';b=io.BytesIO();w.save(b);return b.getvalue()

def presentation():
 p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Inches(1), Inches(1),Inches(8),Inches(2)).text=SECRET+' '+KEEP;s.notes_slide.notes_text_frame.text='NOTES_PRIVATE';hidden=p.slides.add_slide(p.slide_layouts[6]);hidden.shapes.add_textbox(0,0,1000000,1000000).text='HIDDEN_SLIDE_PRIVATE';hidden._element.set('show','0');p.core_properties.author='AUTHOR_PRIVATE';b=io.BytesIO();p.save(b);return b.getvalue()

def pdf_source():
 images=[]
 for k in range(2):
  im=Image.new('RGB',(400,300),'white');im.paste('red',(20,20,120,70));im.paste('green',(200,150,350,250));images.append(im)
 raw=formats.make_pdf(images);rd=PdfReader(io.BytesIO(raw));wr=PdfWriter();[wr.add_page(p)for p in rd.pages];wr.add_metadata({'/Author':'AUTHOR_PRIVATE'});wr.add_attachment('PRIVATE.txt',b'HIDDEN_PRIVATE');b=io.BytesIO();wr.write(b);return b.getvalue()

class FormatTests(unittest.TestCase):
 def test_text_markdown_unicode_and_full_source(self):
  for category in ['text','markdown']:
   original=('# '+KEEP+'\n<!-- '+SECRET+' -->\n姓名：'+SECRET+'\n`'+SECRET+'`').encode()
   output=formats.export(original,category,plan(category));self.assertNotIn(SECRET,output.decode());self.assertIn(KEEP,output.decode());self.assertTrue(formats.verify(output,category,plan(category))['checks']['utf8_readable'])
 def test_encoding_requires_choice_and_preserves_business(self):
  raw=(SECRET+' '+KEEP).encode('gb18030')
  with self.assertRaises(engine.Rejected):formats.inspect(raw,'text')
  output=formats.export(raw,'text',plan('text',input_encoding='gb18030'));self.assertIn(KEEP,output.decode())
  raw=(SECRET+' '+KEEP).encode('utf-16');self.assertIn(SECRET,formats.inspect(raw,'text')['preview'])
 def test_csv_quotes_newlines_and_formula_injection(self):
  for category,sep in [('csv',','),('tsv','\t')]:
   raw=(f'姓名{sep}内容\n{SECRET}{sep}"{KEEP}\n第二行"\n公式{sep}=1+1\n').encode();out=formats.export(raw,category,plan(category));rows=formats.csv_rows(out,category);self.assertEqual(rows[1][1],KEEP+'\n第二行');self.assertEqual(rows[2][1],"'=1+1");self.assertNotIn(SECRET,str(rows))
 def test_excel_removes_hidden_links_comments_metadata(self):
  raw=workbook();out=formats.export(raw,'excel',plan('excel'));w=load_workbook(io.BytesIO(out));self.assertEqual(w.sheetnames,['表1']);self.assertEqual(w.active['B1'].value,KEEP);self.assertEqual(w.active['C1'].value,120);self.assertTrue(w.active['B1'].alignment.wrap_text);self.assertGreaterEqual(w.active.column_dimensions['B'].width,18);self.assertFalse(w.active['A1'].comment);self.assertFalse(w.active['B1'].hyperlink)
  with zipfile.ZipFile(io.BytesIO(out))as z:
   data=b' '.join(z.read(n)for n in z.namelist());self.assertNotIn(b'PRIVATE',data);self.assertFalse(any('comment'in n for n in z.namelist()))
 def test_excel_formula_without_cached_value_stops(self):
  w=Workbook();w.active['A1']='=1+1';b=io.BytesIO();w.save(b)
  with self.assertRaises(engine.Rejected):formats.inspect(b.getvalue(),'excel')
 def test_excel_text_formula_never_executed(self):
  w=Workbook();w.active['A1']='=WEBSERVICE("private")';w.active['A1'].data_type='s';w.active['B1']=SECRET+' '+KEEP;b=io.BytesIO();w.save(b);out=formats.export(b.getvalue(),'excel',plan('excel'));r=load_workbook(io.BytesIO(out));self.assertEqual(r.active['A1'].data_type,'s')
 def test_pptx_notes_hidden_source_properties_removed(self):
  out=formats.export(presentation(),'slides',plan('slides'));p=Presentation(io.BytesIO(out));self.assertEqual(len(p.slides),1);self.assertIn(KEEP,p.slides[0].shapes[0].text)
  with zipfile.ZipFile(io.BytesIO(out))as z:
   data=b' '.join(z.read(n)for n in z.namelist());self.assertNotIn(b'PRIVATE',data);self.assertNotIn(SECRET.encode(),data)
 def test_pdf_flatten_masks_all_pages_and_drops_hidden_carriers(self):
  source=pdf_source();p={'reviewed_pages':[1,2],'pages':[{'page':1,'rectangles':[[20,20,120,70]]}],'acknowledge_loss':True}
  out=formats.export(source,'pdf',p);rd=PdfReader(io.BytesIO(out));self.assertIsNone(rd.metadata);self.assertNotIn(b'PRIVATE',out);self.assertFalse(any(page.extract_text()for page in rd.pages));self.assertTrue(formats.verify(out,'pdf',p,source)['checks']['outside_pixels_preserved'])
  a=formats.pdf_image(out,0);self.assertEqual(a.getpixel((40,40)),(0,0,0));self.assertEqual(a.getpixel((220,180)),(0,128,0))
 def test_pdf_missing_page_review_rejects(self):
  with self.assertRaises(engine.Rejected):formats.export(pdf_source(),'pdf',{'reviewed_pages':[1],'pages':[{'page':1,'rectangles':[[20,20,120,70]]}],'acknowledge_loss':True})
 def test_pdf_verify_requires_original_and_rejects_extra_objects(self):
  src=pdf_source();p={'reviewed_pages':[1,2],'pages':[{'page':1,'rectangles':[[20,20,120,70]]}],'acknowledge_loss':True};out=formats.export(src,'pdf',p)
  with self.assertRaises(engine.Rejected):formats.verify(out,'pdf',p)
  w=PdfWriter();[w.add_page(x)for x in PdfReader(io.BytesIO(out)).pages];w.add_attachment('extra.txt',b'leak');b=io.BytesIO();w.write(b)
  with self.assertRaises(engine.Rejected):formats.verify(b.getvalue(),'pdf',p,src)
 def test_bad_encoding_control_and_unmatched_stops(self):
  for data in [b'bad\x00bytes',b'\xff\xff']:
   with self.assertRaises(engine.Rejected):formats.inspect(data,'text')
  with self.assertRaises(engine.Rejected):formats.export(KEEP.encode(),'markdown',plan('markdown'))
 def test_cli_new_formats_source_and_output_protection(self):
  with tempfile.TemporaryDirectory()as tmp:
   d=Path(tmp);source=d/'原件.md';source.write_text(SECRET+' '+KEEP,encoding='utf-8');p=d/'计划.json';p.write_text(json.dumps(plan('markdown'),ensure_ascii=False),encoding='utf-8');out=d/'副本.md'
   args=[sys.executable,str(Path(cli.__file__)),'redact','--input',str(source),'--plan',str(p),'--output',str(out)]
   r=subprocess.run(args,capture_output=True);self.assertEqual(r.returncode,0,r.stdout);self.assertNotIn(SECRET.encode(),r.stdout);before=out.read_bytes();r=subprocess.run(args,capture_output=True);self.assertEqual(r.returncode,2);self.assertEqual(out.read_bytes(),before)

if __name__=='__main__':unittest.main()
