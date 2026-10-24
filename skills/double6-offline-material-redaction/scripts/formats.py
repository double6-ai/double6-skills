"""新增格式的本地重建适配；原文件与原词不进入日志。"""
import csv
import io
import json
import re
import zipfile
import zlib
from pathlib import PurePosixPath
from datetime import date, datetime, time
import unicodedata
import math
import logging
import warnings
# 第三方解析告警可能包含材料片段；只由上层返回固定错误与统计。
logging.getLogger("pypdf").addHandler(logging.NullHandler())
logging.getLogger("pypdf").propagate=False
logging.getLogger("pypdf").setLevel(logging.CRITICAL+1)
warnings.filterwarnings("ignore", module=r"openpyxl.*")
from xml.etree import ElementTree as ET

from PIL import Image, ImageChops
import engine
import cell_display
from engine import Rejected

EXTENSIONS = {'.txt':'text','.md':'markdown','.markdown':'markdown','.csv':'csv','.tsv':'tsv','.xlsx':'excel','.xlsm':'excel','.pptx':'slides','.pdf':'pdf'}
OUTPUT = {'text':'.txt','markdown':'.md','csv':'.csv','tsv':'.tsv','excel':'.xlsx','slides':'.pptx','pdf':'.pdf','word':'.docx','image':'.png'}
MIME = {'text':'text/plain; charset=utf-8','markdown':'text/markdown; charset=utf-8','csv':'text/csv; charset=utf-8','tsv':'text/tab-separated-values; charset=utf-8','excel':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','slides':'application/vnd.openxmlformats-officedocument.presentationml.presentation','pdf':'application/pdf'}
MAX_CHARS = 400_000
MAX_CELLS = 100_000
MAX_PDF_PAGES = 20
MAX_PDF_PIXELS = 40_000_000
NOTICES = {
 'text':'完整原文按所选词替换，输出UTF-8文本；不执行文本内容。',
 'markdown':'预览展示Markdown源码，含链接、代码、注释与front matter；按完整源码替换，输出UTF-8，不自动清除未选中的文字。',
 'csv':'保留行列和文字值，输出UTF-8 CSV。普通负数保持原样；等号、加号、at，或减号后面不是普通数字的文本会加单引号。保留项按完整单元格核对。',
 'tsv':'保留行列和文字值，输出UTF-8 TSV。普通负数保持原样；等号、加号、at，或减号后面不是普通数字的文本会加单引号。保留项按完整单元格核对。',
 'excel':'仅重建可见工作表、可见行列中的值；移除隐藏表/行/列、公式、图表、图片、批注、链接、宏和原属性，工作表改名。公式仅采用文件缓存值，不重新计算；缺少缓存即停止。',
 'slides':'仅重建显示幻灯片的文字与表格，按阅读顺序重排；图片、图表、动画、备注、批注、隐藏页、母版和原属性不带入，不能保留原设计。',
 'pdf':'默认原生删改PDF，保留框外文字和页面。原属性、附件、链接与批注清理；复杂隐藏层/表单会停止。只有明确选择图像化才重建扫描式PDF。无OCR，需逐页核对。',
}


def package(data):
    if len(data)>engine.MAX_BYTES: raise Rejected('文件超过20MB。')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos=z.infolist();names=[x.filename for x in infos]
            if len(names)>1500 or len(set(names))!=len(names) or sum(x.file_size for x in infos)>80*1024*1024: raise Rejected('压缩文档体积或部件数量超限。')
            if any(n.startswith(('/', '\\')) or '\\'in n or '..'in PurePosixPath(n).parts for n in names): raise Rejected('压缩文档路径无效。')
            xmls=[x for x in infos if x.filename.endswith(('.xml','.rels'))]
            if any(x.file_size>8*1024*1024 for x in xmls) or sum(x.file_size for x in xmls)>16*1024*1024: raise Rejected('文档XML体积超限。')
            parts={n:z.read(n)for n in names if not n.endswith('/')}
        for n,b in parts.items():
            if n.endswith(('.xml','.rels')): engine.parse_xml(b)
        return parts
    except (zipfile.BadZipFile,RuntimeError,ET.ParseError)as exc:raise Rejected('压缩文档无法安全读取，旧版或加密文件需单独转换。')from exc


def decode(data, encoding=None):
    if encoding not in (None,'utf-8','gb18030','utf-16'):raise Rejected('文本编码仅支持UTF-8、带BOM的UTF-16或显式GB18030。')
    codec=encoding or ('utf-16'if data.startswith((b'\xff\xfe',b'\xfe\xff'))else'utf-8-sig')
    if codec=='utf-16'and not data.startswith((b'\xff\xfe',b'\xfe\xff')):raise Rejected('UTF-16文本需带BOM，避免猜测字节顺序。')
    try:text=data.decode(codec)
    except UnicodeError as exc:raise Rejected('文本编码无法读取；请明确选择原文件编码，不自动猜测。')from exc
    if len(text)>MAX_CHARS or any(ord(c)<32 and c not in '\t\r\n'for c in text):raise Rejected('文本过大或包含不支持的控制字符。')
    return text


def replacement(preview,plan):
    terms=engine.terms_checked(plan.get('terms',[]))
    if not terms and plan.get('clean_only')is not True:raise Rejected('请补选原文；仅清理隐藏内容需明确选择。')
    if terms and plan.get('clean_only'):raise Rejected('仅清理隐藏内容与替换词不能同时使用。')
    return engine.prepare_replacement(preview, terms)


def rows_text(rows):return '\n'.join('\t'.join(''if v is None else str(v)for v in row)for row in rows)


def csv_rows(data,category,encoding=None):
    try:rows=list(csv.reader(io.StringIO(decode(data,encoding),newline=''),delimiter=','if category=='csv'else'\t',strict=True))
    except csv.Error as exc:raise Rejected('分隔文本结构无效。')from exc
    if sum(map(len,rows))>MAX_CELLS:raise Rejected('表格单元格超过10万格。')
    return rows


def excel_model(data, include_formats=False):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    package(data)
    try:
        wb=load_workbook(io.BytesIO(data),data_only=False,keep_links=False)
        values=load_workbook(io.BytesIO(data),data_only=True,keep_links=False)
        themes=cell_display.theme_colors(wb.loaded_theme)
        model=[];removed={'hidden_sheets':0,'hidden_rows':0,'hidden_columns':0,'hidden_format_cells':0,'formulas':0,'uncached_formulas':0};total=0
        for ws in wb:
            if ws.sheet_state!='visible':removed['hidden_sheets']+=1;continue
            if len(ws.conditional_formatting):raise Rejected('工作表含条件格式，可能改变显示或隐藏内容；当前停止而非猜测可见性。')
            if ws.max_row*ws.max_column>MAX_CELLS:raise Rejected('工作表范围超过10万格，已停止。')
            hidden_rows={n for n,d in ws.row_dimensions.items()if d.hidden}
            hidden_cols=set()
            for key,d in ws.column_dimensions.items():
                if d.hidden:
                    lo,hi=d.min,d.max
                    if not lo or not hi:lo=hi=range_boundaries(key+':'+key)[0]
                    if hi-lo>16384:raise Rejected('工作表隐藏列范围无效。')
                    hidden_cols.update(range(lo,hi+1))
            removed['hidden_rows']+=len(hidden_rows);removed['hidden_columns']+=len(hidden_cols);cells=[]
            for row in ws:
                for cell in row:
                    if cell.row in hidden_rows or cell.column in hidden_cols or cell.value is None:continue
                    if cell_display.hidden_format(cell.number_format):removed['hidden_format_cells']+=1;continue
                    cell_display.check_font(cell,themes,wb._colors)
                    value=cell.value
                    if cell.data_type=='f':
                        removed['formulas']+=1;value=values[ws.title][cell.coordinate].value
                        if value is None:removed['uncached_formulas']+=1;continue
                    if not isinstance(value,(str,bool,int,float,date,datetime,time)):raise Rejected('工作表含不支持的单元格值。')
                    cell_display.display(value,cell.number_format)
                    cells.append((cell.row,cell.column,value,cell.number_format))
            total+=len(cells)
            if total>MAX_CELLS:raise Rejected('工作簿总单元格超过10万格。')
            model.append(cells)
        wb.close();values.close()
    except Rejected:raise
    except Exception as exc:raise Rejected('Excel无法安全读取；旧版、加密或复杂无效文件需单独处理。')from exc
    if not model or not any(model):raise Rejected('没有支持的可见单元格内容。')
    if removed['uncached_formulas']:raise Rejected('有可见公式缺少缓存结果，不能静默丢弃；请在本机表格软件重新计算并保存后再处理。')
    preview='\n'.join('表'+str(k+1)+'\n'+'\n'.join(str(r)+','+str(c)+'\t'+cell_display.display(v,f)for r,c,v,f in cells)for k,cells in enumerate(model))
    if len(preview)>MAX_CHARS:raise Rejected('Excel可见内容过多。')
    return model if include_formats else [[(r,c,v)for r,c,v,f in cells]for cells in model],removed,preview


def slides_model(data):
    from pptx import Presentation
    parts=package(data)
    if any('vba'in n.lower()or'/embeddings/'in n for n in parts):raise Rejected('演示文稿的宏/嵌入对象尚未支持。')
    try:
        prs=Presentation(io.BytesIO(data));model=[];hidden=0
        if len(prs.slides)>100:raise Rejected('演示文稿超过100页。')
        def visit(shapes,items):
            for shape in shapes:
                if hasattr(shape,'shapes'):visit(shape.shapes,items)
                elif shape.has_text_frame:
                    text=shape.text
                    if text:items.append(('text',text))
                elif shape.has_table:items.append(('table',[[cell.text for cell in row.cells]for row in shape.table.rows]))
        for slide in prs.slides:
            if slide._element.get('show')in ('0','false'):hidden+=1;continue
            items=[];visit(slide.shapes,items)
            if not items:raise Rejected('有可见幻灯片只含图片、图表或未支持对象，重建会丢失整页，已停止。')
            model.append(items)
    except Rejected:raise
    except Exception as exc:raise Rejected('演示文稿无法安全读取。')from exc
    preview='\n'.join('第'+str(k+1)+'页\n'+'\n'.join(v if typ=='text'else rows_text(v)for typ,v in items)for k,items in enumerate(model))
    if not model or len(preview)>MAX_CHARS:raise Rejected('演示文稿无可见页或文字过多。')
    return model,{'hidden_slides':hidden},preview


def pdf_info(data):
    from pypdf import PdfReader
    import pypdfium2 as pdfium
    if not data.startswith(b'%PDF-')or len(data)>engine.MAX_BYTES:raise Rejected('PDF文件头无效或体积超限。')
    try:
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted:raise Rejected('加密PDF暂不支持，请使用有权处理的未加密副本。')
        count=len(reader.pages)
        if not 1<=count<=MAX_PDF_PAGES:raise Rejected('PDF需为1至20页；请先按业务范围拆分。')
        sizes=[];pixels=0
        with pdfium.PdfDocument(data)as doc:
            if len(doc)!=count:raise Rejected('PDF解析结果不一致。')
            for page in doc:
                w,h=page.get_size();size=(int(w*2+0.99999),int(h*2+0.99999));pixels+=size[0]*size[1]
                if min(size)<=0 or size[0]*size[1]>engine.MAX_PIXELS:raise Rejected('PDF单页尺寸超限。')
                sizes.append(size);page.close()
        if pixels>MAX_PDF_PIXELS:raise Rejected('PDF总渲染像素超过4000万，请先拆分。')
        return sizes
    except Rejected:raise
    except Exception as exc:
        if b'/Encrypt' in data:raise Rejected('加密PDF暂不支持。')from exc
        raise Rejected('PDF无法安全解析。')from exc


def pdf_image(data,index):
    import pypdfium2 as pdfium
    sizes=pdf_info(data)
    if type(index)is not int or not 0<=index<len(sizes):raise Rejected('PDF页码无效。')
    with pdfium.PdfDocument(data)as doc:
        # 表单值与批注按当前可见外观渲染，输出不复制可交互对象。
        doc.init_forms()
        page=doc[index];bitmap=page.render(scale=2,may_draw_forms=True,draw_annots=True)
        image=bitmap.to_pil().convert('RGB').copy();bitmap.close();page.close()
    if image.size!=sizes[index]:raise Rejected('PDF渲染尺寸与计划不一致。')
    return image


def pdf_boxes(plan,sizes):
    if plan.get('pdf_mode','preserve') not in {'preserve','rasterize'}:raise Rejected('PDF模式无效。')
    if plan.get('pdf_mode','preserve')=='rasterize' and plan.get('acknowledge_loss')is not True:raise Rejected('图像化PDF须明确选择并接受文字层、签名及交互移除。')
    reviewed=plan.get('reviewed_pages')
    if not isinstance(reviewed,list)or any(type(x)is not int for x in reviewed)or sorted(reviewed)!=list(range(1,len(sizes)+1)):raise Rejected('请逐页核对PDF，reviewed_pages需完整列出每一页。')
    rows=plan.get('pages',[])
    if not isinstance(rows,list)or len(rows)>len(sizes):raise Rejected('PDF框选计划无效。')
    result={};count=0
    for row in rows:
        if not isinstance(row,dict)or set(row)!={'page','rectangles'}or type(row['page'])is not int or not 1<=row['page']<=len(sizes)or row['page']in result:raise Rejected('PDF框选页码重复或无效。')
        boxes=engine.rectangles(row['rectangles'],sizes[row['page']-1]);result[row['page']]=boxes;count+=len(boxes)
    if not count:raise Rejected('PDF至少需框选一处；未补选不导出。')
    if plan.get('pdf_mode','preserve')=='rasterize' and plan.get('keep'):raise Rejected('纯图像PDF不能用文字keep判定保留，请核对框外像素与实际页面。')
    return result


def make_pdf(images):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,NumberObject,DecodedStreamObject
    writer=PdfWriter();writer.metadata=None
    for image in images:
        w,h=image.size;page=writer.add_blank_page(width=w/2,height=h/2)
        stream=DecodedStreamObject();stream.set_data(image.tobytes());stream.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(w),NameObject('/Height'):NumberObject(h),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        image_ref=writer._add_object(stream.flate_encode());content=DecodedStreamObject();content.set_data(f'q {w/2} 0 0 {h/2} 0 0 cm /Im0 Do Q'.encode('ascii'))
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im0'):image_ref})});page[NameObject('/Contents')]=writer._add_object(content)
    out=io.BytesIO();writer.write(out);return out.getvalue()


def inspect(data,category,encoding=None):
    removed={}
    if category in ('text','markdown'):preview=decode(data,encoding)
    elif category in ('csv','tsv'):preview=rows_text(csv_rows(data,category,encoding))
    elif category=='excel':_,removed,preview=excel_model(data)
    elif category=='slides':_,removed,preview=slides_model(data)
    elif category=='pdf':return {'notice':NOTICES[category],'pages':[{'page':k+1,'width':w,'height':h}for k,(w,h)in enumerate(pdf_info(data))],'manual_rectangles_required':True}
    else:raise Rejected('不支持的格式。')
    return {'preview':preview,'notice':NOTICES[category],'removed':removed,'suggestions':engine.suggest_terms(preview)}


def export(data,category,plan):
    if category=='pdf' and plan.get('pdf_mode','preserve')=='preserve':
        import pdf_preserve
        return pdf_preserve.export(data,plan)
    if category=='pdf':
        sizes=pdf_info(data);boxes=pdf_boxes(plan,sizes);images=[]
        for k in range(len(sizes)):
            im=pdf_image(data,k)
            for b in boxes.get(k+1,[]):im.paste((0,0,0),b)
            images.append(im)
        output=make_pdf(images)
    else:
        info=inspect(data,category,plan.get('input_encoding'));clean=replacement(info['preview'],plan)
        if category in ('text','markdown'):output=clean(info['preview']).encode('utf-8')
        else:
            if plan.get('acknowledge_loss')is not True:raise Rejected('请先接受该格式的清理和内容重建方式。')
            if category in ('csv','tsv'):
                rows=csv_rows(data,category,plan.get('input_encoding'));dest=io.StringIO(newline='');writer=csv.writer(dest,delimiter=','if category=='csv'else'\t',lineterminator=plan.get('_line_ending','\n'))
                writer.writerows([[safe_cell(clean(s))for s in row]for row in rows]);output=dest.getvalue().encode('utf-8-sig')
            elif category=='excel':
                from openpyxl import Workbook
                from openpyxl.packaging.core import DocumentProperties
                from openpyxl.styles import Alignment
                model,_,_=excel_model(data,include_formats=True);wb=Workbook();wb.remove(wb.active);wb.properties=DocumentProperties(creator='',lastModifiedBy='',description='')
                for k,cells in enumerate(model):
                    ws=wb.create_sheet('表'+str(k+1))
                    for row,col,value,fmt in cells:
                        shown=cell_display.display(value,fmt);replaced=clean(shown)
                        if replaced!=shown:value=replaced;fmt='General'
                        cell=ws.cell(row,col,value)
                        if isinstance(value,str):cell.data_type='s'
                        cell.number_format=fmt
                    from openpyxl.utils.cell import get_column_letter
                    widths={}
                    for row in ws:
                        for cell in row:
                            units=max((sum(2 if unicodedata.east_asian_width(ch)in 'WF'else 1 for ch in line)for line in str(cell.value or '').split('\n')),default=0)
                            widths[cell.column]=max(widths.get(cell.column,12),min(50,units+4))
                    for col,width in widths.items():ws.column_dimensions[get_column_letter(col)].width=width
                    for row in ws:
                        height=18
                        for cell in row:
                            cell.alignment=Alignment(wrap_text=True,vertical='top')
                            length=sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(ch)in 'WF'else 1 for ch in line)/max(1,widths[cell.column]-3)))for line in str(cell.value or '').split('\n'))
                            height=max(height,min(409,18*length))
                        ws.row_dimensions[row[0].row].height=height
                    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
                dest=io.BytesIO();wb.save(dest);output=dest.getvalue()
            elif category=='slides':
                from pptx import Presentation
                from pptx.util import Inches,Pt
                model,_,_=slides_model(data);prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
                for items in model:
                    for typ,value in items:
                        if typ=='text':
                            lines=[]
                            for line in clean(value).split('\n'):
                                part='';width=0
                                for ch in line:
                                    step=2 if unicodedata.east_asian_width(ch)in 'WF'else 1
                                    if width+step>92:lines.append(part);part='';width=0
                                    part+=ch;width+=step
                                lines.append(part)
                            for start in range(0,len(lines),20):
                                slide=prs.slides.add_slide(prs.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(.5),Inches(.4),Inches(12.3),Inches(6.7));box.text='\n'.join(lines[start:start+20])
                                for para in box.text_frame.paragraphs:para.font.size=Pt(18)
                        else:
                            rows=[[clean(c)for c in row]for row in value]
                            if not rows:continue
                            cols=len(rows[0])
                            if cols>10 or any(len(v)>80 or sum(2 if unicodedata.east_asian_width(ch)in 'WF'else 1 for ch in v)>max(1,int(165/cols))*3 for row in rows for v in row):raise Rejected('演示表格超过10列或单格过长，无法保证重排后可读，请先拆分。')
                            for start in range(0,len(rows),6):
                                chunk=rows[start:start+6];slide=prs.slides.add_slide(prs.slide_layouts[6]);table=slide.shapes.add_table(len(chunk),cols,Inches(.4),Inches(.4),Inches(12.5),Inches(6.5)).table
                                for rn,row in enumerate(chunk):
                                    for cn,text in enumerate(row):
                                        table.cell(rn,cn).text=text;table.cell(rn,cn).text_frame.word_wrap=True
                                        for para in table.cell(rn,cn).text_frame.paragraphs:para.font.size=Pt(14)
                    if not items:prs.slides.add_slide(prs.slide_layouts[6])
                props=prs.core_properties
                for key in ['author','last_modified_by','title','subject','keywords','comments','category']:setattr(props,key,'')
                dest=io.BytesIO();prs.save(dest);output=dest.getvalue()
            else:raise Rejected('不支持的格式。')
    if len(output)>engine.MAX_BYTES:raise Rejected('导出超过20MB，请缩小业务范围后重试。')
    verify(output,category,{**plan,'_verify_encoding':'utf-8'} if category in {'text','markdown','csv','tsv'} else plan,data)
    return output


_PLAIN_NUMBER = re.compile(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)%?')


def safe_cell(text):
    # 普通负数不是公式。只给表格软件可能执行的文本加单引号，避免把 -120 改成文本。
    if text == '':
        return text
    core = text.strip(' \t')
    if '\r' not in text and '\n' not in text and _PLAIN_NUMBER.fullmatch(core):
        return text
    head = text.lstrip(' \t')
    if text[:1] in '\t\r\n' or head[:1] in '=+-@':
        return "'" + text
    return text


def cells_preserve_keep(source_rows, output_rows, keep):
    """保留项按完整单元格比对。子串命中加了引号的格子不算保留。"""
    if source_rows is not None and (len(source_rows) != len(output_rows) or any(len(a) != len(b) for a, b in zip(source_rows, output_rows))):
        raise Rejected('保留项核对时行列结构变了。')
    for item in keep:
        exact = []
        if source_rows is not None:
            exact = [(r, c) for r, row in enumerate(source_rows) for c, value in enumerate(row) if value == item]
        if exact:
            if any(output_rows[r][c] != item for r, c in exact):
                raise Rejected('保留项“' + item + '”按完整单元格核对已变化。防公式处理只给真正的公式文本加引号；普通负数保持原样。')
            continue
        hits = [cell for row in output_rows for cell in row if item in cell]
        if not hits:
            raise Rejected('保留项“' + item + '”没有留下。')
        if all(cell == "'" + item for cell in hits):
            raise Rejected('保留项“' + item + '”只出现在加了引号的单元格里，完整值已经变了。')
    return True


def verify(data,category,plan,source=None):
    if category=='pdf':
        if plan.get('pdf_mode','preserve')=='preserve':
            import pdf_preserve
            return pdf_preserve.verify(data,plan,source)
        return verify_pdf(data,plan,source)
    checks={}
    info=inspect(data,category,plan.get('_verify_encoding',plan.get('input_encoding')));preview=info['preview'];terms=engine.terms_checked(plan.get('terms',[]))
    checks['selected_strings_absent']=not any(engine.term_remains(preview, x) for x in terms)
    keep=engine.keep_checked(plan.get('keep',[]))
    if keep and category in ('csv','tsv'):
        output_encoding=plan.get('_verify_encoding',plan.get('input_encoding'))
        if data.startswith(b'\xef\xbb\xbf') and output_encoding in (None,'utf-8'):output_encoding=None
        output_rows=csv_rows(data,category,output_encoding)
        source_rows=csv_rows(source,category,plan.get('input_encoding')) if source is not None else None
        checks['keep_phrases_preserved']=cells_preserve_keep(source_rows,output_rows,keep)
    elif keep:
        checks['keep_phrases_preserved']=all(x in preview for x in keep)
    if category in ('text','markdown','csv','tsv'):
        checks['text_encoding_readable']=True
        if category in ('csv','tsv'):checks['formula_text_neutralized']=all(safe_cell(s)==s for row in csv_rows(data,category,plan.get('_verify_encoding',plan.get('input_encoding')))for s in row)
    else:
        parts=package(data);strings=[]
        for n,b in parts.items():
            if n.endswith(('.xml','.rels')):
                root=engine.parse_xml(b);strings.append(''.join(root.itertext())+' '.join(v for el in root.iter()for v in el.attrib.values()))
        checks['selected_strings_absent']=not any(engine.term_remains(s, x) for x in terms for s in strings)
        checks['hidden_carriers_absent']=not any('/media/'in n or '/embeddings/'in n or'comment'in n.lower()or'vba'in n.lower()or n.startswith('customXml/')or'/notesSlides/'in n for n in parts)
        checks['external_links_absent']=not any(el.get('TargetMode')=='External'for n,b in parts.items()if n.endswith('.rels')for el in engine.parse_xml(b))
        checks['original_properties_absent']=not any(''.join(engine.parse_xml(parts[n]).itertext()).strip()for n in ['docProps/custom.xml']if n in parts)
        core=engine.parse_xml(parts['docProps/core.xml'])if'docProps/core.xml'in parts else []
        checks['original_properties_absent'] &= all(not (el.text or '').strip() for el in core if el.tag.rsplit('}',1)[-1]not in {'created','modified','revision','lastPrinted'})
        if category=='slides':checks['hidden_slides_absent']=not info['removed']['hidden_slides']
        if category=='excel':
            checks['formulas_hidden_absent']=not any(info['removed'].values())
    if not all(checks.values()):raise Rejected('导出复检失败；未交付副本。')
    return {'checks':checks,'not_checked':[]if keep else ['keep_phrases_preserved'],'output_sha256':engine.digest(data),'scope':'仅检查所选原词、当前支持的隐藏载体及已声明保留条件；需人工核对补选完整性和实际文件。'}


def verify_pdf(data,plan,source):
    from pypdf import PdfReader
    if source is None:raise Rejected('PDF复检必须提供原文件--source。')
    sizes=pdf_info(source);boxes=pdf_boxes(plan,sizes);expected_images=[]
    try:
        rd=PdfReader(io.BytesIO(data),strict=True)
        if rd.is_encrypted or rd.metadata or len(rd.pages)!=len(sizes)or set(rd.trailer['/Root'])!={'/Type','/Pages'}:raise Rejected('PDF输出含多余对象或页数不一致。')
        for k,page in enumerate(rd.pages):
            if set(page)!={'/Type','/Resources','/MediaBox','/Parent','/Contents'}:raise Rejected('PDF输出包含非图像页面对象。')
            if set(page['/Resources'])!={'/XObject'}or set(page['/Resources']['/XObject'])!={'/Im0'}:raise Rejected('PDF输出资源结构不符。')
            image=page['/Resources']['/XObject']['/Im0'];w,h=sizes[k]
            if set(image)!={'/Type','/Subtype','/Width','/Height','/ColorSpace','/BitsPerComponent','/Filter'}or image['/Width']!=w or image['/Height']!=h or image['/ColorSpace']!='/DeviceRGB'or image['/BitsPerComponent']!=8:raise Rejected('PDF图像结构不符。')
            raw=image.get_data()
            if len(raw)!=w*h*3:raise Rejected('PDF图像长度不符。')
            out=Image.frombytes('RGB',(w,h),raw);expected=pdf_image(source,k)
            for b in boxes.get(k+1,[]):expected.paste((0,0,0),b)
            expected_images.append(expected)
            if ImageChops.difference(out,expected).getbbox()is not None:raise Rejected('PDF遮盖或框外像素不一致。')
            expected_stream=f'q {w/2} 0 0 {h/2} 0 0 cm /Im0 Do Q'.encode('ascii')
            if page['/Contents'].get_data()!=expected_stream:raise Rejected('PDF页面指令不符。')
        # 从新构造对象重写必须是唯一当前对象集合；完整对象数校验避免挂入未引用历史载体。
        if len(rd.xref.get(0,{}))!=2+3*len(sizes):raise Rejected('PDF输出存在多余历史对象。')
        if data!=make_pdf(expected_images):raise Rejected('PDF输出字节含额外或非本工具对象。')
    except Rejected:raise
    except Exception as exc:raise Rejected('PDF实际输出复检失败。')from exc
    return {'checks':{'image_only_pdf':True,'selected_regions_solid':True,'outside_pixels_preserved':True,'original_objects_absent':True},'output_sha256':engine.digest(data),'scope':'144dpi纯图像PDF；逐页核对计划区域与框外像素，原PDF对象不复制；不自动识别敏感内容。'}
