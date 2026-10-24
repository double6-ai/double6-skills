"""本地脱敏原型：限定图片格式与可编辑 Word 干净副本。"""
from __future__ import annotations

import hashlib
import io
import math
import posixpath
import re
import urllib.parse
import zipfile
from datetime import datetime
from xml.etree import ElementTree as ET

from PIL import Image, ImageChops, ImageOps
from docx import Document

VERSION = "0.6.1"
MAX_BYTES = 20 * 1024 * 1024
MAX_PIXELS = 24_000_000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/package/2006/relationships}"


class Rejected(ValueError):
    """输入超出原型已支持范围；导出前停止。"""


class _NoDTD(ET.TreeBuilder):
    def __init__(self):
        super().__init__()
        self.depth = self.nodes = 0

    def start(self, tag, attrs):
        self.depth += 1
        self.nodes += 1
        if self.depth > 128 or self.nodes > 200_000:
            raise Rejected("Word XML结构过于复杂，已停止读取。")
        return super().start(tag, attrs)

    def end(self, tag):
        result = super().end(tag)
        self.depth -= 1
        return result

    def doctype(self, name, pubid, system):
        raise Rejected("Word含不支持的DTD声明，已停止读取。")


def parse_xml(raw):
    try:
        return ET.fromstring(raw, parser=ET.XMLParser(target=_NoDTD()))
    except ET.ParseError as exc:
        raise Rejected("Word XML结构无效，已停止读取。") from exc


def digest(data):
    return hashlib.sha256(data).hexdigest()


def image_source(data):
    if len(data) > MAX_BYTES:
        raise Rejected("文件超过 20 MB。")
    try:
        with Image.open(io.BytesIO(data)) as im:
            if im.format not in {"PNG", "JPEG", "WEBP"}:
                raise Rejected("图片只支持静态 PNG、JPEG 和 WebP。")
            if getattr(im, "n_frames", 1) != 1:
                raise Rejected("多帧图片尚未支持，请先导出需要处理的单张图片。")
            if im.width * im.height > MAX_PIXELS:
                raise Rejected("图片超过 2400 万像素。")
            im.load()
            normalized = ImageOps.exif_transpose(im).convert("RGBA")
            # 将透明区域合到白底，避免透明像素中保留原图内容。
            clean = Image.new("RGB", normalized.size, "white")
            clean.paste(normalized, mask=normalized.getchannel("A"))
            return clean
    except (OSError, Image.DecompressionBombError) as exc:
        raise Rejected("图片无法安全读取。") from exc


def png_bytes(im):
    out = io.BytesIO()
    # 从纯像素新建，避免继承 EXIF、XMP、PNG 文本或源缩略图。
    fresh = Image.frombytes("RGB", im.size, im.tobytes())
    fresh.save(out, format="PNG")
    return out.getvalue()


def rectangles(values, size):
    if not isinstance(values, list) or not 1 <= len(values) <= 200:
        raise Rejected("请先框选至少一处需要遮盖的区域，最多 200 处。")
    result = []
    for value in values:
        if not isinstance(value, list) or len(value) != 4:
            raise Rejected("框选区域格式无效。")
        if any(isinstance(n, bool) or not isinstance(n, (int, float)) or not math.isfinite(n) for n in value):
            raise Rejected("框选坐标无效。")
        x0, y0, x1, y1 = value
        box = (math.floor(min(x0, x1)), math.floor(min(y0, y1)), math.ceil(max(x0, x1)), math.ceil(max(y0, y1)))
        if not (0 <= box[0] < box[2] <= size[0] and 0 <= box[1] < box[3] <= size[1]):
            raise Rejected("框选区域超出图片边界。")
        result.append(box)
    return result


def export_image(data, selected):
    source = image_source(data)
    boxes = rectangles(selected, source.size)
    target = source.copy()
    for box in boxes:
        target.paste((0, 0, 0), box)
    output = png_bytes(target)
    if len(output) > MAX_BYTES:
        raise Rejected("导出PNG超过20MB，请先缩小原图后重新框选。")
    with Image.open(io.BytesIO(output)) as reopened:
        reopened.load()
        metadata_clean = not reopened.info and not reopened.getexif()
        masks_solid = all(reopened.crop(box).getbbox() is None for box in boxes)
        outside = ImageChops.difference(source, reopened.convert("RGB"))
        for box in boxes:
            outside.paste((0, 0, 0), box)
        outside_preserved = outside.getbbox() is None
    checks = {"metadata_removed": metadata_clean, "selected_regions_solid": masks_solid, "outside_pixels_preserved": outside_preserved}
    if not all(checks.values()):
        raise Rejected("图片导出复检未通过，未提供下载。")
    return output, {"version": VERSION, "format": "png", "checks": checks, "regions": len(boxes), "sha256": digest(output), "scope": "仅确认框选区域、单帧像素和元数据；未自动证明所有敏感内容已选全。"}


def read_package(data):
    if len(data) > MAX_BYTES:
        raise Rejected("文件超过 20 MB。")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
        if len(entries) > 1000 or sum(x.file_size for x in entries) > 80 * 1024 * 1024:
            raise Rejected("Word 包体过大或结构过于复杂。")
        xml_entries = [x for x in entries if x.filename.endswith((".xml", ".rels"))]
        if any(x.file_size > 8 * 1024 * 1024 for x in xml_entries) or sum(x.file_size for x in xml_entries) > 16 * 1024 * 1024:
            raise Rejected("Word XML体积超过原型范围。")
        names = [x.filename for x in entries]
        if len(names) != len(set(names)):
            raise Rejected("Word 包内有重名部件。")
        if any(n.startswith(("/", "\\")) or "\\" in n or ".." in n.split("/") for n in names):
            raise Rejected("Word 包内路径无效。")
        if "word/document.xml" not in names:
            raise Rejected("请选择 .docx 文件。旧版 .doc 与加密文件尚未支持。")
        if any("vba" in n.lower() or n.startswith("word/embeddings/") for n in names):
            raise Rejected("含宏或嵌入文件的 Word 尚未支持，已停止导出。")
        parts = {n: archive.read(n) for n in names if not n.endswith("/")}
        if any(b"<!DOCTYPE" in b.upper() or b"<!ENTITY" in b.upper() for n, b in parts.items() if n.endswith((".xml", ".rels"))):
            raise Rejected("Word 含不支持的 XML 声明。")
        missing = missing_targets(parts)
        if missing:
            raise Rejected("Word 原文件存在缺失部件，请先修复后重试。")
        return parts
    except (zipfile.BadZipFile, RuntimeError, ET.ParseError, OSError) as exc:
        raise Rejected("Word 文件无法安全读取。") from exc


def missing_targets(parts):
    missing = []
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        base = posixpath.dirname(name).removesuffix("/_rels")
        if name == "_rels/.rels":
            base = ""
        for rel in parse_xml(data):
            if rel.get("TargetMode") == "External":
                continue
            target = urllib.parse.unquote(rel.get("Target", ""))
            target = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
            if target not in parts:
                missing.append(target)
    return missing


def visible_text(node):
    if node.tag in {W + "del", W + "moveFrom", W + "drawing", W + "pict"}:
        return ""
    if node.tag == W + "r":
        props = node.find(W + "rPr")
        if props is not None and any(props.find(W + x) is not None and props.find(W + x).get(W + "val", "true") not in {"0", "false", "off"} for x in ["vanish", "webHidden"]):
            return ""
    if node.tag == W + "t":
        return node.text or ""
    if node.tag == W + "tab":
        return "\t"
    if node.tag in {W + "br", W + "cr"}:
        return "\n"
    return "".join(visible_text(child) for child in node)


def extract_blocks(container):
    result = []
    for node in container:
        if node.tag == W + "p":
            result.append({"kind": "paragraph", "text": visible_text(node)})
        elif node.tag == W + "tbl":
            if any(x.find(".//" + W + "tbl") is not None for x in node.findall(".//" + W + "tc")):
                raise Rejected("嵌套表格尚未支持。")
            rows = []
            for row in node.findall(W + "tr"):
                rows.append(["\n".join(visible_text(p) for p in cell.findall(W + "p")) for cell in row.findall(W + "tc")])
            if rows and max(map(len, rows), default=0) <= 30:
                result.append({"kind": "table", "rows": rows})
            else:
                raise Rejected("表格结构尚未支持。")
        elif node.tag != W + "sectPr":
            raise Rejected("含内容控件或其它未支持的正文对象，已停止导出。")
    return result


def inspect_docx(data):
    parts = read_package(data)
    roots = {n: parse_xml(b) for n, b in parts.items() if n.startswith("word/") and n.endswith(".xml")}
    # 样式继承可能令正文不可见；尚未实现完整Word样式级联，检测到时停止。
    styles = [root for root in roots.values() if root.tag == W + "styles"]
    if any(el.tag in {W + "vanish", W + "webHidden"} and el.get(W + "val", "true").lower() not in {"0", "false", "off"} for style_root in styles for el in style_root.iter()):
        raise Rejected("包含由样式控制的隐藏文字，当前版本尚未支持，已停止导出。")
    unsupported = {W + x for x in ["txbxContent", "altChunk", "object", "footnoteReference", "endnoteReference", "sdt"]}
    if any(el.tag in unsupported for root in roots.values() for el in root.iter()):
        raise Rejected("文本框、嵌入对象、脚注或尾注尚未支持，已停止导出。")
    body = roots["word/document.xml"].find(W + "body")
    if body is None:
        raise Rejected("Word 缺少正文。")
    main = extract_blocks(body)
    headers, footers = [], []
    # 只复制正文实际引用的默认页眉页脚；未引用的历史部件不会重新变成正文。
    rels = parse_xml(parts.get("word/_rels/document.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'))
    by_id = {el.get("Id"): el for el in rels}
    referenced = {"header": set(), "footer": set()}
    office_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for ref in body.iter():
        label = "header" if ref.tag == W + "headerReference" else "footer" if ref.tag == W + "footerReference" else None
        if label is None:
            continue
        if ref.get(W + "type", "default") != "default":
            raise Rejected("首页或奇偶页专用页眉页脚尚未支持，已停止导出。")
        rel = by_id.get(ref.get(office_id))
        if rel is None or not rel.get("Type", "").endswith("/" + label) or rel.get("TargetMode") == "External":
            raise Rejected("页眉页脚引用无效，已停止导出。")
        target = posixpath.normpath(posixpath.join("word", urllib.parse.unquote(rel.get("Target", "")))).lstrip("/")
        if target not in roots or roots[target].tag != W + ("hdr" if label == "header" else "ftr"):
            raise Rejected("页眉页脚部件无效，已停止导出。")
        referenced[label].add(target)
    for name in sorted(referenced["header"]):
        headers.extend(extract_blocks(roots[name]))
    for name in sorted(referenced["footer"]):
        footers.extend(extract_blocks(roots[name]))
    blocks = headers + main + footers
    plain = "\n".join(b["text"] if b["kind"] == "paragraph" else "\n".join("\t".join(row) for row in b["rows"]) for b in blocks)
    if len(plain) > 400_000:
        raise Rejected("Word 文字量超过原型范围。")
    counts = {"images": sum(n.startswith("word/media/") for n in parts), "comment_parts": sum("comment" in n.lower() for n in parts), "deleted_or_moved_fragments": sum(el.tag in {W + "del", W + "moveFrom"} for root in roots.values() for el in root.iter()), "headers_and_footers": len(headers) + len(footers)}
    suggestions = sorted(set(re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}|(?<!\d)1\d{10}(?!\d)", plain)))
    return {"blocks": blocks, "preview": plain, "suggestions": suggestions, "removed": counts, "notice": "导出为可编辑的文字和表格副本；页眉页脚移到正文，图片、批注、删除修订、隐藏文字、超链接地址及原属性移除。字体、分页和合并单元格样式不保留。"}


def terms_checked(values):
    if not isinstance(values, list) or len(values) > 200:
        raise Rejected("补选内容最多 200 项。")
    if any(not isinstance(v, str) or not 1 <= len(v.strip()) <= 200 for v in values):
        raise Rejected("每项补选内容应为 1–200 个字符。")
    return sorted(set(v.strip() for v in values), key=lambda v: (-len(v), v))


def keep_checked(values):
    if not isinstance(values,list) or len(values)>200 or any(not isinstance(v,str)or not 1<=len(v.strip())<=500 for v in values):
        raise Rejected('关键保留项需为最多200项、每项1–500字的列表。')
    return list(dict.fromkeys(v.strip()for v in values))


def export_docx(data, terms, acknowledge_loss=False, clean_only=False, keep=None):
    if acknowledge_loss is not True:
        raise Rejected("请先确认干净副本的图片和版式处理方式。")
    plan = inspect_docx(data)
    selected = terms_checked(terms)
    if not selected and clean_only is not True:
        raise Rejected("请补选正文原文；只清理隐藏内容时需明确选择该模式。")
    if selected and clean_only:
        raise Rejected("仅清理隐藏内容模式不能同时填写替换词。")
    if any(v not in plan["preview"] for v in selected):
        raise Rejected("有补选原文未在预览中找到，请核对准确原文后重试；隐藏部件无需列为正文替换词。")
    replacements = {v: "【已隐去" + str(n + 1) + "】" for n, v in enumerate(selected)}
    # 单次匹配，避免后来输入的词再次命中替换标记。
    regex = re.compile("|".join(re.escape(v) for v in selected)) if selected else None
    def clean(text):
        return regex.sub(lambda m: replacements[m.group()], text) if regex else text
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(W + "eastAsia", "Microsoft YaHei")
    for block in plan["blocks"]:
        if block["kind"] == "paragraph":
            doc.add_paragraph(clean(block["text"]))
        else:
            rows = block["rows"]
            table = doc.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.style = "Table Grid"
            for n, row in enumerate(rows):
                for k, text in enumerate(row):
                    table.cell(n, k).text = clean(text)
    props = doc.core_properties
    for field in ["author", "last_modified_by", "title", "subject", "keywords", "comments", "category", "content_status", "identifier", "language", "version"]:
        setattr(props, field, "")
    props.created = props.modified = datetime(2000, 1, 1)
    out = io.BytesIO()
    doc.save(out)
    output = strip_template_extras(out.getvalue())
    checks = verify_docx(output, selected)
    required=keep_checked(keep if keep is not None else [])
    if required:checks['keep_phrases_preserved']=all(v in inspect_docx(output)['preview']for v in required)
    if not all(checks.values()):
        raise Rejected("Word 导出复检未通过，未提供下载。请检查补选是否与替换标记冲突。")
    return output, {"version": VERSION, "format": "docx", "checks": checks, "not_checked":[]if required else ['keep_phrases_preserved'], "selected_terms": len(selected), "matched_terms": sum(v in plan["preview"] for v in selected), "removed": plan["removed"], "sha256": digest(output), "scope": "复检覆盖所选原文、已声明保留项、干净副本结构及列出的隐藏载体；姓名等未补选内容仍需本人核对。"}


def strip_template_extras(data):
    """清除新建模板自带的书目 XML/缩略图，绝不从输入包复制部件。"""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    removed = {n for n in parts if n.startswith("customXml/") or n == "docProps/thumbnail.jpeg"}
    for name in removed:
        parts.pop(name)
    for name, raw in list(parts.items()):
        if name.endswith(".rels"):
            root = parse_xml(raw)
            base = "" if name == "_rels/.rels" else posixpath.dirname(name).removesuffix("/_rels")
            for rel in list(root):
                target = posixpath.normpath(posixpath.join(base, rel.get("Target", ""))).lstrip("/")
                if target in removed:
                    root.remove(rel)
            # 部分原生读取器要求 OPC 根沿用常见的默认命名空间写法。
            ET.register_namespace("", root.tag[1:].split("}")[0])
            parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        elif name == "[Content_Types].xml":
            root = parse_xml(raw)
            for el in list(root):
                if el.get("PartName", "").lstrip("/") in removed:
                    root.remove(el)
            ET.register_namespace("", root.tag[1:].split("}")[0])
            parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, raw in parts.items():
            z.writestr(name, raw)
    return out.getvalue()


def verify_docx(data, terms):
    parts = read_package(data)
    strings = []
    external = False
    history = False
    for name, raw in parts.items():
        if name.endswith((".xml", ".rels")):
            root = parse_xml(raw)
            strings.append(urllib.parse.unquote("".join(root.itertext()) + " ".join(v for el in root.iter() for v in el.attrib.values())))
            external |= any(el.get("TargetMode") == "External" for el in root.iter())
            history |= any(el.tag in {W + x for x in ["del", "ins", "moveFrom", "moveTo", "vanish", "webHidden"]} for el in root.iter())
    doc = Document(io.BytesIO(data))
    props = doc.core_properties
    return {"selected_strings_absent": not any(v in s for s in strings for v in terms), "media_embeddings_comments_absent": not any(n.startswith(("word/media/", "word/embeddings/", "customXml/")) or "comment" in n.lower() for n in parts), "external_relationships_absent": not external, "revision_hidden_markup_absent": not history, "original_properties_absent": not any(getattr(props, f) for f in ["author", "last_modified_by", "title", "subject", "keywords", "comments"]), "internal_relationships_valid": not missing_targets(parts), "editable_document_reopened": doc._element.body is not None}
