"""Build T6-T11 regression templates (permissive sources) -> templates/, updates SOURCES.json."""
import hashlib, json, shutil, subprocess, zipfile, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from lxml import etree
ROOT = Path(__file__).resolve().parents[1] / "templates"; SRC = ROOT.parent / "src"
LO = Path("/usr/lib/libreoffice/share/template/common/presnt")
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
sources = json.load(open(ROOT / "SOURCES.json"))

ZH = {"title": ["年度经营分析报告", "核心业务回顾", "市场拓展策略", "风险与应对", "下一步计划"],
      "sub": "经营管理部 ｜ 2026年10月",
      "body": [("业务收入同比增长百分之十二，主要来自华东与华南市场", 0), ("新客户贡献占比提升到三成", 1), ("毛利率保持稳定，成本管控成效明显", 0), ("采购成本下降约四个百分点", 1), ("数字化渠道成为新的增长点", 0)]}
EN = {"title": ["Annual Business Review", "Core Business Highlights", "Market Expansion Strategy", "Risks and Mitigations", "Next Steps"],
      "sub": "Strategy Office · October 2026",
      "body": [("Revenue grew 12% year over year, led by the East and South regions", 0), ("New customers now contribute roughly a third", 1), ("Gross margin held steady thanks to cost discipline", 0), ("Procurement costs down about four points", 1), ("Digital channels became a new growth engine", 0)]}

def put(ph, content):
    tf = ph.text_frame
    if isinstance(content, list):
        for i, (t, l) in enumerate(content):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = t; p.level = l
    else:
        tf.text = content

def mixed_box(slide, lang, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h); tb.name = "KPI Callout" if lang == "en" else "指标说明"
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = "Revenue: " if lang == "en" else "营业收入："; r1.font.bold = True; r1.font.size = Pt(18)
    r2 = p.add_run(); r2.text = "12.5M USD" if lang == "en" else "12.5亿元"; r2.font.size = Pt(18)
    from pptx.dml.color import RGBColor
    r2.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    return tb

def transition(slide, kind="fade"):
    el = slide._element
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    for ac in el.findall(f"{{{MC}}}AlternateContent"):
        if ac.find(f".//{{{P}}}transition") is not None:
            return  # keep the template's own (p14 Choice + Fallback) transition
    old = el.find(f"{{{P}}}transition")
    if old is not None:
        for c in list(old): old.remove(c)
        etree.SubElement(old, f"{{{P}}}{kind}")
        return
    tr = etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med"><p:{kind}/></p:transition>')
    anchor = el.find(f"{{{P}}}clrMapOvr")
    if anchor is None: anchor = el.find(f"{{{P}}}cSld")
    anchor.addnext(tr)

def lo_template(name, out, lang, extra_layout=None):
    d = SRC / "lo2"; d.mkdir(exist_ok=True)
    shutil.copy(LO / f"{name}.otp", d / f"{name}.odp")
    subprocess.run(["soffice", "--headless", "--convert-to", "pptx", "--outdir", str(d), str(d / f"{name}.odp")], check=True, capture_output=True, timeout=300)
    prs = Presentation(d / f"{name}.pptx")
    S = ZH if lang == "zh" else EN
    content_layout = None
    for i, s in enumerate(prs.slides):
        for ph in s.placeholders:
            t = ph.placeholder_format.type
            if t == 1: put(ph, S["title"][i % len(S["title"])])
            elif t == 4: put(ph, S["sub"])
            elif t in (7, 2): put(ph, S["body"]); content_layout = s.slide_layout
        s.notes_slide.notes_text_frame.text = ("讲者备注：本页示例备注" if lang == "zh" else "Speaker notes: sample note for this slide")
    if content_layout is not None:
        s = prs.slides.add_slide(content_layout)
        for ph in s.placeholders:
            t = ph.placeholder_format.type
            if t == 1: put(ph, S["title"][3])
            elif t in (7, 2): put(ph, S["body"][:3])
        W, H = prs.slide_width, prs.slide_height
        mixed_box(s, lang, Emu(int(W * 0.55)), Emu(int(H * 0.78)), Emu(int(W * 0.4)), Emu(int(H * 0.1)))
        s.notes_slide.notes_text_frame.text = "第二条备注" if lang == "zh" else "Second sample note"
    transition(prs.slides[0], "fade")
    if len(prs.slides) > 1: transition(prs.slides[1], "push")
    prs.save(ROOT / out)
    sources[out] = {"origin": f"LibreOffice bundled Impress template {name}.otp ({LO}), converted to PPTX by soffice; sample {lang} text written into its placeholders, one extra content slide from its own layout, mixed-run callout, speaker notes and transitions added via python-pptx",
                    "license": "MPL-2.0 (LibreOffice core extras/source/templates)", "source_sha256": sha(LO / f"{name}.otp"),
                    "features": ["LO-exported master/layout", "placeholders", "speaker notes", "transitions", "mixed-format runs", lang]}

def poi_layouts():
    src = SRC / "poi" / "layouts.pptx"; prs = Presentation(src)
    S = EN
    for i, s in enumerate(prs.slides):
        for ph in s.placeholders:
            t = ph.placeholder_format.type
            if t == 1 or t == 3: put(ph, S["title"][i % 5])
            elif t == 4: put(ph, S["sub"])
            elif t in (2, 7) and ph.has_text_frame: put(ph, S["body"][: 2 + i % 3])
        s.notes_slide.notes_text_frame.text = f"Presenter note {i + 1}"
    mixed_box(prs.slides[1], "en", Emu(int(prs.slide_width * 0.55)), Emu(int(prs.slide_height * 0.85)), Emu(int(prs.slide_width * 0.4)), Emu(int(prs.slide_height * 0.08)))
    transition(prs.slides[0], "fade")
    prs.save(ROOT / "T10-en-poi-layouts.pptx")
    sources["T10-en-poi-layouts.pptx"] = {"origin": "https://github.com/apache/poi/raw/trunk/test-data/slideshow/layouts.pptx (PowerPoint-authored, 10 slides on the 11 Office layouts); English sample text, notes, transition and a mixed-run callout added via python-pptx",
                                          "license": "Apache-2.0 (Apache POI test-data)", "source_sha256": sha(src),
                                          "features": ["PowerPoint-authored Office theme", "all built-in layouts", "two-content/comparison", "notes", "mixed-format runs"]}

def poi_smartart():
    src = SRC / "poi" / "SmartArt.pptx"; prs = Presentation(src)
    lay = prs.slide_layouts[0]
    # put a title slide first by adding and moving it
    s = prs.slides.add_slide(lay)
    for ph in s.placeholders:
        if ph.placeholder_format.type in (1, 3): put(ph, "组织结构与职责")
        elif ph.placeholder_format.type == 4: put(ph, "人力资源部 ｜ 2026年10月")
    lst = prs.slides._sldIdLst; ids = list(lst); lst.remove(ids[-1]); lst.insert(0, ids[-1])
    lay2 = prs.slide_layouts[1]; s2 = prs.slides.add_slide(lay2)
    for ph in s2.placeholders:
        if ph.placeholder_format.type == 1: put(ph, "部门职责说明")
        elif ph.placeholder_format.type in (2, 7): put(ph, ZH["body"])
    transition(prs.slides[1], "fade")
    prs.save(ROOT / "T11-zh-poi-smartart.pptx")
    sources["T11-zh-poi-smartart.pptx"] = {"origin": "https://github.com/apache/poi/raw/trunk/test-data/slideshow/SmartArt.pptx (PowerPoint-authored SmartArt diagram); title and bullet slides from its own layouts with Chinese text added via python-pptx",
                                           "license": "Apache-2.0 (Apache POI test-data)", "source_sha256": sha(src),
                                           "features": ["SmartArt (dgm parts) left untouched", "non_text_content_unedited warning", "CJK", "transition"]}

if __name__ == "__main__" and "--lo-only" in __import__("sys").argv:
    for a, b, c in (("Vivid", "T6-en-lo-vivid.pptx", "en"), ("Focus", "T7-zh-lo-focus.pptx", "zh"), ("Piano", "T8-zh-lo-piano.pptx", "zh"), ("Inspiration", "T9-en-lo-inspiration.pptx", "en")):
        lo_template(a, b, c)
    for k in list(sources):
        if (ROOT / k).exists(): sources[k]["sha256"] = sha(ROOT / k)
    json.dump(sources, open(ROOT / "SOURCES.json", "w"), ensure_ascii=False, indent=2)
    raise SystemExit
if __name__ == "__main__":
    lo_template("Vivid", "T6-en-lo-vivid.pptx", "en")
    lo_template("Focus", "T7-zh-lo-focus.pptx", "zh")
    lo_template("Piano", "T8-zh-lo-piano.pptx", "zh")
    lo_template("Inspiration", "T9-en-lo-inspiration.pptx", "en")
    poi_layouts(); poi_smartart()
    for k in list(sources):
        if (ROOT / k).exists(): sources[k]["sha256"] = sha(ROOT / k)
    json.dump(sources, open(ROOT / "SOURCES.json", "w"), ensure_ascii=False, indent=2)
    print("\n".join(sorted(sources)))
