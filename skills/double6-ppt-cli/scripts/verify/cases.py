"""Deterministic template-fill verification cases (built with python-pptx; no external files)."""
from __future__ import annotations

from pathlib import Path

from lxml import etree

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _transition(slide, kind="fade"):
    el = slide._element
    anchor = el.find(f"{{{P}}}clrMapOvr")
    if anchor is None:
        anchor = el.find(f"{{{P}}}cSld")
    anchor.addnext(etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med"><p:{kind}/></p:transition>'))


def _bullets(tf, items):
    for i, (t, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = t, lvl


def build_case_en(path: Path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly Business Review"
    s.placeholders[1].text = "Strategy Office"
    s.notes_slide.notes_text_frame.text = "Sample cover notes"
    _transition(s)
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Executive summary"
    _bullets(s.placeholders[1].text_frame, [("Revenue headline", 0), ("Supporting detail", 1), ("Second headline", 0)])
    tb = s.shapes.add_textbox(Emu(700000), Emu(5900000), Emu(6000000), Emu(500000))
    tb.name = "KPI Callout"
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "Revenue: "; r.font.bold = True; r.font.size = Pt(18)
    r = p.add_run(); r.text = "12.5M USD"; r.font.size = Pt(18)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Pipeline"
    t = s.shapes.add_table(3, 3, Emu(700000), Emu(1700000), Emu(6000000), Emu(1200000)).table
    for i, row in enumerate([["Stage", "Deals", "Value"], ["Lead", "10", "1.0M"], ["Won", "3", "0.4M"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    cd = CategoryChartData(); cd.categories = ["Q1", "Q2"]; cd.add_series("Revenue", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(7000000), Emu(1700000), Emu(4500000), Emu(3000000), cd)
    prs.save(path)


def plan_en(lib):
    def slot(slide, pred):
        return next(x for x in lib["slides"][slide - 1]["slots"] if pred(x))
    body = slot(2, lambda x: (x.get("placeholder") or [None])[0] == "body" or x.get("shape_name", "").startswith("Content"))
    long_bullets = "\n".join([
        "Revenue reached $48.2M, up 18% year over year and 6% above plan",
        "\tGrowth driven by three new strategic accounts in the DACH and Nordics regions",
        "Gross margin improved to 64.1% as discounting discipline held across all segments",
        "\tAverage discount on new business fell from 22% to 14%",
        "Pipeline coverage for next quarter stands at 3.4x with healthy late-stage mix",
        "\tPublic sector deals remain dependent on budget approvals",
        "Hiring plan on track: 42 of 45 roles filled, attrition below 8%",
    ])
    tbl = lib["slides"][2]["tables"][0]["table_id"]
    ch = lib["slides"][2]["charts"][0]["chart_id"]
    return {"schema": "template_fill_pptx_plan.v1", "status": "confirmed", "slides": [
        {"source_slide": 1, "replacements": [{"slot_id": slot(1, lambda x: x["role"] == "title_candidate")["slot_id"], "text": "Q3 FY2026 Quarterly Business Review"}],
         "notes": "Welcome everyone. Keep this under one minute."},
        {"source_slide": 2, "replacements": [{"slot_id": body["slot_id"], "text": long_bullets},
                                             {"shape_name": "KPI Callout", "text": "Revenue: 48.2M USD"}],
         "notes": "Pause after the margin line.", "transition": {"effect": "push", "duration": 0.6}},
        {"source_slide": 3, "replacements": [], "notes": "keep",
         "table_edits": [{"table_id": tbl, "cells": [{"row": 1, "col": 1, "text": "14"}, {"row": 2, "col": 2, "text": "0.9M"}]}],
         "chart_edits": [{"chart_id": ch, "categories": ["Q1", "Q2", "Q3"], "series": [{"name": "Revenue", "values": [41.0, 44.5, 48.2]}]}]},
        {"source_slide": 2, "replacements": [{"slot_id": body["slot_id"], "text": "Close three public sector deals\nHire two account executives\nLaunch the partner program", "levels": [0, 0, 0]}],
         "notes": "Priorities for next quarter."},
    ]}


def build_case_zh(path: Path):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "项目汇报"
    s.placeholders[1].text = "部门 ｜ 日期"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "现状与问题"
    _bullets(s.placeholders[1].text_frame, [("要点一", 0), ("说明", 1), ("要点二", 0)])
    for i, x in enumerate((700000, 4300000, 7900000)):
        tb = s.shapes.add_textbox(Emu(x), Emu(5000000), Emu(3300000), Emu(900000))
        tb.name = f"卡片{i + 1}"
        tb.text_frame.word_wrap = True
        tb.text_frame.text = "卡片说明文字"
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.size = Pt(16)
    s.notes_slide.notes_text_frame.text = "示例备注"
    _transition(s, "push")
    prs.save(path)


def plan_zh(lib):
    sl = lib["slides"][1]["slots"]
    body = next(x for x in sl if x.get("shape_name", "").startswith("Content") or (x.get("placeholder") or [None])[0] == "body")
    title = next(x for x in sl if x["role"] == "title_candidate")
    text = "\n".join([
        "现状：12条产线使用7套独立系统，数据口径不统一，月度经营分析需要人工汇总3天以上",
        "\t质量问题平均追溯时间达到36小时，客户投诉响应慢",
        "痛点：设备综合效率（OEE）仅为61%，低于行业标杆约15个百分点",
        "\t停机原因依赖纸质台账记录，无法开展根因分析",
        "目标：18个月内实现全厂数据贯通，OEE提升至75%，交付周期缩短20%",
    ])
    cards = ["统一采集12条产线、400余台设备的实时数据，建立主数据标准。",
             "边缘网关接入PLC与传感器，异常提前48小时预警，工单自动派发。",
             "以批次码串联来料、加工、检验与发运，追溯时间缩短到15分钟。"]
    reps = [{"slot_id": title["slot_id"], "text": "产线数据孤岛导致交付周期长、质量追溯难"}, {"slot_id": body["slot_id"], "text": text}]
    reps += [{"shape_name": f"卡片{i + 1}", "text": c} for i, c in enumerate(cards)]
    return {"schema": "template_fill_pptx_plan.v1", "status": "confirmed", "slides": [
        {"source_slide": 1, "replacements": [], "notes": "开场"},
        {"source_slide": 2, "replacements": reps, "notes": "重点讲OEE差距。"},
    ]}


CASES = [("en_qbr", build_case_en, plan_en), ("zh_factory", build_case_zh, plan_zh)]
