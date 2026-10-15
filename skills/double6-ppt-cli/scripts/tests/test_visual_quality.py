"""Phase-3 visual quality suite: rubric, deterministic checks, gate wiring, template-fill fidelity."""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from double6_ppt_cli.visual_quality import (
    RUBRIC_DIMENSIONS,
    compute_visual_metrics,
    contrast_ratio,
    evaluate_rubric,
)

SKILL = Path(__file__).resolve().parents[2]
W, H = 12192000, 6858000


def _deck(path: Path, *, decorated: bool = True, chart_colors: str | None = None, sizes=(34, 18), text_rgb="1E2733") -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = s.shapes.add_textbox(Emu(600000), Emu(500000), Emu(9000000), Emu(800000))
    run = title.text_frame.paragraphs[0].add_run(); run.text = "Revenue grew 18%"
    run.font.size = Pt(sizes[0]); run.font.color.rgb = RGBColor.from_string(text_rgb)
    for i, sz in enumerate(sizes[1:]):
        b = s.shapes.add_textbox(Emu(600000), Emu(1500000 + i * 400000), Emu(5000000), Emu(400000))
        r = b.text_frame.paragraphs[0].add_run(); r.text = f"body {i}"; r.font.size = Pt(sz)
        r.font.color.rgb = RGBColor.from_string(text_rgb)
    if decorated:
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(600000), Emu(4200000), Emu(10000000), Emu(1800000))
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    if chart_colors is not None:
        cd = CategoryChartData(); cd.categories = ["A", "B"]; cd.add_series("S", (1, 2))
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(6500000), Emu(1500000), Emu(4000000), Emu(2500000), cd).chart
        if chart_colors:
            ser = ch.plots[0].series[0]
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(chart_colors)
    prs.save(path)


def _page(path: Path, *, bottom: bool = True, text_bg=(255, 255, 255)) -> None:
    img = Image.new("RGB", (1600, 900), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle((80, 60, 1260, 170), fill=text_bg)
    for y in range(80, 150, 12):
        d.line((90, y, 900, y), fill=(30, 39, 51), width=2)
    for y in range(205, 255, 12):
        d.line((90, y, 700, y), fill=(30, 39, 51), width=2)
    if bottom:
        d.rectangle((80, 550, 1390, 790), fill=(31, 58, 95))
    img.save(path)


class RubricTests(unittest.TestCase):
    def test_eight_dimensions_without_restraint(self):
        self.assertEqual(len(RUBRIC_DIMENSIONS), 8)
        self.assertNotIn("restraint", RUBRIC_DIMENSIONS)

    def test_pass_requires_mean_and_floor(self):
        ok = evaluate_rubric({d: 4 for d in RUBRIC_DIMENSIONS})
        self.assertTrue(ok["pass"])
        low = evaluate_rubric({**{d: 5 for d in RUBRIC_DIMENSIONS}, "polish": 2})
        self.assertFalse(low["pass"]); self.assertEqual(low["below_min"], ["polish"])
        mean = evaluate_rubric({d: 3 for d in RUBRIC_DIMENSIONS})
        self.assertFalse(mean["pass"])
        self.assertFalse(evaluate_rubric({"hierarchy": 5})["pass"])
        with self.assertRaises(ValueError):
            evaluate_rubric({d: 6 for d in RUBRIC_DIMENSIONS})

    def test_contrast_ratio(self):
        self.assertAlmostEqual(contrast_ratio((0, 0, 0), (255, 255, 255)), 21.0, places=1)


class MetricTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, deck_kw=None, page_kw=None, palette=None):
        deck, page = self.root / "d.pptx", self.root / "p.png"
        _deck(deck, **(deck_kw or {})); _page(page, **(page_kw or {}))
        return compute_visual_metrics(deck, [page], palette=palette)

    def codes(self, r):
        return {f["code"] for f in r["findings"] if f["severity"] == "error"}

    def test_clean_page_passes(self):
        r = self._run()
        self.assertEqual(r["status"], "pass", r["findings"])

    def test_text_only_and_top_heavy_detected(self):
        r = self._run({"decorated": False}, {"bottom": False})
        self.assertTrue({"no_non_text_element", "top_heavy"} <= self.codes(r))

    def test_font_size_budget(self):
        r = self._run({"sizes": (40, 34, 28, 24, 20, 18, 16, 14)})
        self.assertIn("too_many_font_sizes", self.codes(r))

    def test_low_contrast(self):
        r = self._run({"text_rgb": "BBBBBB"})
        self.assertIn("low_contrast", self.codes(r))

    def test_office_default_chart_only_with_palette(self):
        palette = {"primary": "#1F3A5F", "accent": "#3C7A78", "chart_series": ["#3C7A78", "#1F3A5F"]}
        self.assertIn("office_default_chart_colors", self.codes(self._run({"chart_colors": ""}, palette=palette)))
        self.assertNotIn("office_default_chart_colors", self.codes(self._run({"chart_colors": ""})))
        self.assertNotIn("office_default_chart_colors", self.codes(self._run({"chart_colors": "3C7A78"}, palette=palette)))

    def test_cjk_orphan(self):
        from double6_ppt_cli.visual_quality import _font_for
        font = _font_for("Noto Sans CJK SC", False, 18)
        if font is None:
            self.skipTest("no CJK font")
        text = "季度收入同比增长主要来自新客户"
        deck, page = self.root / "o.pptx", self.root / "o.png"
        _deck(deck); _page(page)
        prs = Presentation(str(deck)); s = prs.slides[0]
        width_pt = font.getlength(text[:-1]) / 4 + 0.5 + 14.4
        box = s.shapes.add_textbox(Emu(600000), Emu(2500000), Emu(int(width_pt * 12700)), Emu(900000))
        box.text_frame.word_wrap = True
        r = box.text_frame.paragraphs[0].add_run(); r.text = text; r.font.size = Pt(18); r.font.name = "Noto Sans CJK SC"
        prs.save(deck)
        self.assertIn("orphan_line", self.codes(compute_visual_metrics(deck, [page])))
        box.width = Emu(int((font.getlength(text) / 4 + 20) * 12700)); prs.save(deck)
        self.assertNotIn("orphan_line", self.codes(compute_visual_metrics(deck, [page])))


class BriefFixturesTests(unittest.TestCase):
    def test_six_reference_briefs(self):
        briefs = json.loads((SKILL / "assets" / "visual-briefs" / "briefs.json").read_text(encoding="utf-8"))
        self.assertEqual(len(briefs["briefs"]), 6)
        langs = {b["lang"] for b in briefs["briefs"]}
        self.assertEqual(langs, {"zh", "en"})
        from double6_ppt_cli.design_presets import names
        for b in briefs["briefs"]:
            self.assertIn(b["design"], names())
            self.assertGreaterEqual(len(b["slides"]), 5)

    def test_quality_reference_exists_and_has_no_restraint_rule(self):
        text = (SKILL / "references" / "visual-quality.md").read_text(encoding="utf-8")
        for dim in RUBRIC_DIMENSIONS:
            self.assertIn(f"`{dim}`", text)
        self.assertIn("盲评", text)
        self.assertNotIn("restraint", text.lower())


class GateWiringTests(unittest.TestCase):
    def test_generate_acceptance_requires_rubric(self):
        import inspect
        from double6_ppt_cli import visual_review, visual_policy
        src = inspect.getsource(visual_review.record_visual_review)
        self.assertIn("visual_rubric_required", src)
        self.assertIn("visual_quality_failed", src)
        self.assertIn('visual_quality', inspect.getsource(visual_policy.resolve_visual_gate))


class TemplateFillFidelityRegression(unittest.TestCase):
    """Template-fill must change only the text it was asked to change; design stays byte-identical."""

    def test_design_shapes_and_run_styles_survive(self):
        from double6_ppt_cli.pm6_template import analyze_pptx, apply_plan, scaffold_plan
        from double6_ppt_cli.pm6_template.ooxml import NS
        from double6_ppt_cli.pm6_template.roundtrip import PM6Error, pm6_root
        try:
            pm6_root()
        except PM6Error:
            self.skipTest("vendored PM6 core unavailable")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            src = root / "designed.pptx"
            prs = Presentation(); prs.slide_width, prs.slide_height = W, H
            s = prs.slides.add_slide(prs.slide_layouts[6])
            band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(W), Emu(900000)); band.name = "Band"
            band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0x22, 0x34, 0x4F); band.line.fill.background()
            dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(600000), Emu(1300000), Emu(300000), Emu(300000)); dot.name = "Accent Dot"
            dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor(0xB0, 0x8D, 0x57)
            t = s.shapes.add_textbox(Emu(1100000), Emu(1200000), Emu(9000000), Emu(700000)); t.name = "Headline"
            r = t.text_frame.paragraphs[0].add_run(); r.text = "Old headline"
            r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = RGBColor(0x16, 0x18, 0x1C); r.font.name = "Georgia"
            prs.save(src)
            lib = analyze_pptx(src)
            slot = next(x for x in lib["slides"][0]["slots"] if x["name"] == "Headline")
            plan = scaffold_plan(lib)
            s1 = dict(plan["slides"][0]); s1["replacements"] = [{"slot_id": slot["slot_id"], "text": "New headline"}]
            out = root / "out.pptx"
            apply_plan(src, dict(plan, slides=[s1]), out, mode="native")

            def shapes(path):
                tree = Presentation(str(path)).slides[0].shapes
                return {sh.name: sh for sh in tree}
            a, b = shapes(src), shapes(out)
            for name in ("Band", "Accent Dot"):
                self.assertEqual(etree.tostring(a[name]._element.spPr), etree.tostring(b[name]._element.spPr), name)
            for name in a:
                self.assertEqual((a[name].left, a[name].top, a[name].width, a[name].height),
                                 (b[name].left, b[name].top, b[name].width, b[name].height), name)
            self.assertEqual(b["Headline"].text_frame.text, "New headline")
            nr = b["Headline"].text_frame.paragraphs[0].runs[0]
            self.assertEqual((nr.font.size, nr.font.bold, str(nr.font.color.rgb), nr.font.name), (Pt(32), True, "16181C", "Georgia"))
            self.assertTrue(NS)


class Phase4ChecksTests(unittest.TestCase):
    def test_orphan_rules(self):
        from double6_ppt_cli.visual_quality import _is_orphan, _last_line, _font_for
        self.assertTrue(_is_orphan("增速远高于门店渠道的 4%", "4%"))
        self.assertTrue(_is_orphan("增速远高于门店渠道的好", "好"))
        self.assertFalse(_is_orphan("增速远高于门店渠道", "门店渠道"))
        self.assertTrue(_is_orphan("Revenue grew across every region in", "in"))
        self.assertFalse(_is_orphan("Revenue grew across every region", "every region"))
        font = _font_for("Noto Sans CJK SC", False, 18)
        if font is not None:
            text = "即时零售增速远高于门店渠道的4%"
            avail = font.getlength("即时零售增速远高于门店渠道的") + 2
            self.assertEqual(_last_line(text, font, avail), "4%")  # digit runs never split

    def test_cjk_latin_space(self):
        from double6_ppt_cli.visual_quality import CJK_LATIN_SPACE
        self.assertTrue(CJK_LATIN_SPACE.search("11 月 11 日"))
        self.assertTrue(CJK_LATIN_SPACE.search("前 20% 门店"))
        self.assertFalse(CJK_LATIN_SPACE.search("11月11日开售"))
        self.assertFalse(CJK_LATIN_SPACE.search("Q3 2026 review"))

    def test_chart_series_distinct(self):
        from double6_ppt_cli.visual_quality import delta_e, MIN_SERIES_DELTA_E
        self.assertLess(delta_e("#1F3A5F", "#2F6E9E"), 25)
        self.assertGreater(delta_e("#F26B38", "#0F2A4A"), MIN_SERIES_DELTA_E)
        import itertools
        from double6_ppt_cli.design_presets import PRESETS, resolve
        for name in PRESETS:
            cs = resolve(name)["chart_series"]
            self.assertGreaterEqual(min(delta_e(a, b) for a, b in itertools.combinations(cs, 2)), 25, name)

    def test_indistinct_series_flagged(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "c.pptx"; page = Path(tmp) / "c.png"
            prs = Presentation(); prs.slide_width, prs.slide_height = W, H
            s = prs.slides.add_slide(prs.slide_layouts[6])
            card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(600000), Emu(4200000), Emu(10000000), Emu(1800000))
            cd = CategoryChartData(); cd.categories = ["A", "B"]; cd.add_series("S1", (1, 2)); cd.add_series("S2", (2, 3))
            ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(600000), Emu(600000), Emu(5000000), Emu(3000000), cd).chart
            for ser, col in zip(ch.plots[0].series, ("1F3A5F", "2F6E9E")):
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
            prs.save(deck); _page(page)
            codes = {f["code"] for f in compute_visual_metrics(deck, [page])["findings"]}
            self.assertIn("chart_series_indistinct", codes)


class ImageSupportTests(unittest.TestCase):
    def test_treatments(self):
        from double6_ppt_cli.images import apply_treatments, parse_treatments, placeholder_image
        src = placeholder_image((800, 600), ["#102030", "#E0A040"])
        out = apply_treatments(src, parse_treatments("cover=640x360;focus=0.5,0.3;duotone=#16181C,#B08D57;tint=#000000@0.2"))
        self.assertEqual(out.size, (640, 360)); self.assertEqual(out.mode, "RGB")
        rounded = apply_treatments(src, parse_treatments("cover=200x200;circle"))
        self.assertEqual(rounded.mode, "RGBA"); self.assertEqual(rounded.getpixel((0, 0))[3], 0)
        with self.assertRaises(Exception):
            parse_treatments("sepia=1")

    def test_credits_gate_and_placeholder(self):
        from double6_ppt_cli.common import D6PPTError
        from double6_ppt_cli.images import add_image, check_image_credits, credits_markdown
        from double6_ppt_cli.runs import init_run
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "b.md"; src.write_text("# x\n", encoding="utf-8")
            run = Path(tmp) / "run"
            init_run("generate", src, run, design_profile="investor-pitch", verify_tier="portable")
            project = run / "authoring" / "project"
            (project / "svg_output" / "01.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="../images/hero.jpg" x="0" y="0" width="10" height="10"/></svg>', encoding="utf-8")
            with self.assertRaises(D6PPTError) as cm:
                check_image_credits(project)
            self.assertEqual(cm.exception.code, "image_credit_missing")
            rec = add_image(run, "hero", placeholder="gradient", colours=["#2B1C8F", "#FF5A3C"], size="640x360",
                            treatments="cover=640x360")
            self.assertTrue(rec["placeholder"])
            (project / "svg_output" / "01.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="../images/hero.jpg" x="0" y="0" width="10" height="10"/></svg>', encoding="utf-8")
            report = check_image_credits(project)
            self.assertEqual(report["placeholders"], ["hero.jpg"])
            self.assertIn("placeholder", credits_markdown(project))
            with self.assertRaises(D6PPTError):
                add_image(run, "x", url="https://example.com/a.jpg")  # no license metadata
