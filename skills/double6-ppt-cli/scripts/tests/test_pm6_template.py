"""Regression tests for the experimental PPT Master v6.7 template-fill backend (exp/pm6-template).

Unit tests run everywhere (lxml; python-pptx for fixtures). Roundtrip integration tests need a
PPT Master v6.x checkout: set D6PPT_PM6_ROOT (and optionally D6PPT_PM6_PYTHON) or they are skipped.
"""
from __future__ import annotations

import hashlib
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree

from double6_ppt_cli.pm6_template import analyze_pptx, check_plan, scaffold_plan
from double6_ppt_cli.pm6_template import textfit
from double6_ppt_cli.pm6_template.ooxml import NS, Package, find_by_id, q
from double6_ppt_cli.pm6_template.roundtrip import _drop_orphan_slide_rels, pm6_root, PM6Error
from double6_ppt_cli.pm6_template.style import fit_context
from double6_ppt_cli.pm6_template.txbody import apply_fit, set_text

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Pt
    HAVE_PPTX = True
except Exception:  # pragma: no cover
    HAVE_PPTX = False

A = NS["a"]


def _ctx(w_pt=200.0, h_pt=60.0, sz=18.0, autofit="none", wrap="square", levels=(0,), latin="Liberation Sans", ea="Noto Sans CJK SC"):
    lvl = {"sz": sz, "marL": 0.0, "indent": 0.0, "lnSpc": None, "spcBef": None, "spcAft": None,
           "latin": latin, "ea": ea, "bold": False, "bullet": False, "algn": None}
    return {"box": {"x_pt": 0, "y_pt": 0, "w_pt": w_pt, "h_pt": h_pt}, "slide": {"w_pt": 960, "h_pt": 540},
            "body": {"lIns": 7.2, "tIns": 3.6, "rIns": 7.2, "bIns": 3.6, "wrap": wrap, "anchor": "t", "autofit": autofit, "vert": None},
            "levels": {"0": dict(lvl), "1": dict(lvl, sz=sz - 2, marL=20.0), "2": dict(lvl, sz=sz - 4, marL=40.0)},
            "template_paragraph_levels": list(levels), "placeholder": None}


def _txbody(xml_paras: str, body_extra: str = "") -> etree._Element:
    return etree.fromstring(f'<p:txBody xmlns:p="{NS["p"]}" xmlns:a="{A}"><a:bodyPr wrap="square">{body_extra}</a:bodyPr><a:lstStyle/>{xml_paras}</p:txBody>')


class TextFitTest(unittest.TestCase):
    def test_cjk_wrap_respects_width_and_kinsoku(self):
        ctx = _ctx(w_pt=150)
        report = textfit.fit(ctx, "边缘推理平台的延迟降低了四成，冷启动时间缩短到两秒以内，成本下降三成。")
        lines = report["lines"][0]
        self.assertGreater(len(lines), 1)
        m = textfit.Metrics(ctx["levels"]["0"])
        for line in lines:
            self.assertLessEqual(m.width(line), 150 - 14.4 + 18.5, line)  # kinsoku may hang one glyph
            self.assertFalse(line[:1] in textfit.NO_LINE_START, f"line starts with closing punctuation: {line}")

    def test_latin_breaks_on_words(self):
        report = textfit.fit(_ctx(w_pt=160), "Revenue reached forty eight million dollars this quarter")
        for line in report["lines"][0]:
            self.assertFalse(line.startswith(" ") or line.endswith(" "))
        self.assertEqual(" ".join(report["lines"][0]), "Revenue reached forty eight million dollars this quarter")

    def test_fit_statuses(self):
        self.assertEqual(textfit.fit(_ctx(w_pt=400, h_pt=80), "short")["status"], "fits")
        long = "word " * 60
        self.assertEqual(textfit.fit(_ctx(w_pt=300, h_pt=60), long)["status"], "overflow")
        auto = textfit.fit(_ctx(w_pt=300, h_pt=70, autofit="norm"), "word " * 14)
        self.assertIn(auto["status"], {"autofit", "fits"})
        grow = textfit.fit(_ctx(w_pt=300, h_pt=20, autofit="shape"), "word " * 30)
        self.assertEqual(grow["status"], "grow")
        self.assertGreater(grow["new_h_pt"], 20)

    def test_autofit_never_below_floor(self):
        r = textfit.fit(_ctx(w_pt=200, h_pt=40, autofit="norm"), "词" * 400)
        self.assertEqual(r["status"], "overflow")
        self.assertGreaterEqual(r["font_scale"], textfit.MIN_FONT_SCALE - 1e-9)

    def test_heuristic_engine_without_fontconfig(self):
        orig = textfit.font_file
        try:
            textfit.font_file = lambda *a, **k: None  # type: ignore
            textfit._font.cache_clear()
            r = textfit.fit(_ctx(w_pt=120), "中文中文中文中文中文中文中文 and latin words")
            self.assertEqual(r["engine"], "heuristic")
            self.assertGreater(r["line_count"], 1)
        finally:
            textfit.font_file = orig
            textfit._font.cache_clear()

    def test_level_mapping(self):
        ctx = _ctx(levels=(0, 1, 0, 1))
        self.assertEqual([l for l, _ in textfit.map_levels(ctx, "a\nb\nc")], [0, 1, 0])
        self.assertEqual([l for l, _ in textfit.map_levels(ctx, "a\n\tb\nc")], [0, 1, 0])
        self.assertEqual([l for l, _ in textfit.map_levels(ctx, "a\nb\nc", [0, 0, 0])], [0, 0, 0])
        self.assertTrue(textfit.levels_are_implicit(ctx, "a\nb"))
        self.assertFalse(textfit.levels_are_implicit(ctx, "a\n\tb"))


class TxBodyTest(unittest.TestCase):
    PARAS = ('<a:p><a:pPr marL="285750" indent="-285750"><a:buChar char="•"/></a:pPr><a:r><a:rPr sz="2000" b="1"/><a:t>L0</a:t></a:r></a:p>'
             '<a:p><a:pPr lvl="1" marL="628650" indent="-285750"><a:buChar char="•"/></a:pPr><a:r><a:rPr sz="1800"/><a:t>L1</a:t></a:r><a:endParaRPr sz="1800"/></a:p>')

    def test_paragraph_prototypes_and_order(self):
        body = _txbody(self.PARAS)
        set_text(body, "one\ntwo\nthree")
        ps = body.findall(q("a", "p"))
        self.assertEqual([("".join(p.xpath(".//a:t/text()", namespaces=NS))) for p in ps], ["one", "two", "three"])
        self.assertEqual([p.find(q("a", "pPr")).get("lvl") for p in ps], [None, "1", "1"])
        self.assertEqual(ps[0].find("./a:r/a:rPr", NS).get("b"), "1")
        for p in ps:  # schema order: pPr, r*, endParaRPr
            tags = [etree.QName(c).localname for c in p]
            if "endParaRPr" in tags:
                self.assertEqual(tags[-1], "endParaRPr")

    def test_explicit_levels_pick_matching_prototype(self):
        body = _txbody(self.PARAS)
        set_text(body, "a\n\tb\nc")
        ps = body.findall(q("a", "p"))
        self.assertEqual([p.find(q("a", "pPr")).get("lvl") for p in ps], [None, "1", None])
        self.assertEqual(ps[2].find("./a:r/a:rPr", NS).get("sz"), "2000")

    def test_soft_break_and_empty_line(self):
        body = _txbody(self.PARAS)
        set_text(body, "x\vy\n")
        ps = body.findall(q("a", "p"))
        self.assertEqual(len(ps), 2)
        self.assertEqual(len(ps[0].findall(q("a", "br"))), 1)
        self.assertEqual(ps[1].findall(q("a", "r")), [])

    def test_explicit_shrink_writes_sizes(self):
        body = _txbody(self.PARAS)
        set_text(body, "a\nb", size_scale=0.9, resolved_sizes={0: 20.0, 1: 18.0})
        szs = [p.find("./a:r/a:rPr", NS).get("sz") for p in body.findall(q("a", "p"))]
        self.assertEqual(szs, ["1800", "1620"])

    def test_apply_fit_norm_autofit_position(self):
        sp = etree.fromstring(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:txBody><a:bodyPr><a:prstTxWarp prst="textNoShape"/><a:spAutoFit/></a:bodyPr></p:txBody></p:sp>')
        apply_fit(sp, {"status": "autofit", "font_scale": 0.85, "ln_spc_reduction": 0.1, "autofit_mode": "norm"})
        kids = [etree.QName(c).localname for c in sp.find("./p:txBody/a:bodyPr", NS)]
        self.assertEqual(kids, ["prstTxWarp", "normAutofit"])
        norm = sp.find(".//a:normAutofit", NS)
        self.assertEqual((norm.get("fontScale"), norm.get("lnSpcReduction")), ("85000", "10000"))


@unittest.skipUnless(HAVE_PPTX, "python-pptx needed for fixtures")
class LibraryAndCheckTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.deck = Path(cls.tmp.name) / "fixture.pptx"
        build_fixture(cls.deck)
        cls.lib = analyze_pptx(cls.deck)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def slot(self, slide, name):
        return next(s for s in self.lib["slides"][slide - 1]["slots"] if s["name"] == name)

    def test_placeholder_box_inherited_from_layout(self):
        title = self.slot(1, "Title 1")
        prs = Presentation(str(self.deck))
        layout_title = prs.slides[0].slide_layout.placeholders[0]
        self.assertAlmostEqual(title["fit_context"]["box"]["w_pt"], layout_title.width / 12700, places=1)
        self.assertEqual(title["role"], "title_candidate")

    def test_tables_charts_detected(self):
        s2 = self.lib["slides"][1]
        self.assertEqual(len(s2["tables"]), 1)
        self.assertEqual(s2["tables"][0]["merge_slaves"], [[1, 1]])
        self.assertEqual(s2["charts"][0]["categories"], ["A", "B"])

    def _plan(self, reps=None, slides=None, **slide_extra):
        plan = scaffold_plan(self.lib)
        plan["slides"] = slides or [plan["slides"][0]]
        if reps is not None:
            plan["slides"][0]["replacements"] = reps
        plan["slides"][0].update(slide_extra)
        return plan

    def codes(self, report, status="ERROR"):
        return [r["code"] for r in report["results"] if r["status"] == status]

    def test_scaffold_passes(self):
        report = check_plan(self.lib, scaffold_plan(self.lib))
        self.assertEqual(report["summary"]["error"], 0, report)

    def test_cross_slide_slot_rejected(self):
        body = self.slot(2, "Body")
        report = check_plan(self.lib, self._plan([{"slot_id": body["slot_id"], "text": "x"}]))
        self.assertIn("cross_slide_slot", self.codes(report))

    def test_stale_old_text_rejected(self):
        body = self.slot(1, "Body")
        report = check_plan(self.lib, self._plan([{"slot_id": body["slot_id"], "old_text": "nope", "text": "x"}]))
        self.assertIn("stale_old_text", self.codes(report))

    def test_overflow_is_error_and_autofit_is_warn(self):
        small = self.slot(1, "Tiny Label")
        report = check_plan(self.lib, self._plan([{"slot_id": small["slot_id"], "text": "far too long for this tiny label " * 4}]))
        self.assertIn("text_fit", self.codes(report))
        body = self.slot(1, "Body")
        report = check_plan(self.lib, self._plan([{"slot_id": body["slot_id"], "text": "\n".join(["A fairly long bullet line that wraps"] * 9)}]))
        fit = next(r for r in report["results"] if r["code"] == "text_fit")
        self.assertIn(fit["status"], {"WARN", "ERROR"})
        self.assertIn(fit["fit"], {"autofit", "overflow"})

    def test_implicit_levels_warn(self):
        body = self.slot(1, "Body")
        report = check_plan(self.lib, self._plan([{"slot_id": body["slot_id"], "text": "a\nb"}]))
        self.assertIn("implicit_list_levels", self.codes(report, "WARN"))
        report = check_plan(self.lib, self._plan([{"slot_id": body["slot_id"], "text": "a\nb", "levels": [0, 0]}]))
        self.assertNotIn("implicit_list_levels", self.codes(report, "WARN"))

    def test_table_merge_slave_and_chart_shape(self):
        lib2 = self.lib["slides"][1]
        plan = scaffold_plan(self.lib)
        plan["slides"] = [plan["slides"][1]]
        plan["slides"][0]["table_edits"] = [{"table_id": lib2["tables"][0]["table_id"], "cells": [{"row": 1, "col": 1, "text": "x"}]}]
        plan["slides"][0]["chart_edits"] = [{"chart_id": lib2["charts"][0]["chart_id"], "categories": ["A", "B"], "series": [{"name": "S", "values": [1]}]}]
        codes = self.codes(check_plan(self.lib, plan))
        self.assertIn("table_cell_is_merge_slave", codes)
        self.assertIn("chart_edit", codes)


@unittest.skipUnless(HAVE_PPTX, "python-pptx needed for fixtures")
class OrphanRelTest(unittest.TestCase):
    def test_drops_only_unreferenced_slide_rels(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "d.pptx"
            prs = Presentation()
            s1, s2 = prs.slides.add_slide(prs.slide_layouts[6]), prs.slides.add_slide(prs.slide_layouts[6])
            box = s1.shapes.add_textbox(0, 0, 100, 100)
            box.click_action.target_slide = s2
            s1.part.relate_to(s2.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide")
            prs.save(deck)
            pkg = Package.open(deck)
            # strip the click action (what the wrapper does) leaving its relationship behind
            root = pkg.xml(pkg.slide_parts()[0])
            for h in root.iter(q("a", "hlinkClick")):
                h.getparent().remove(h)
            pkg.touch(pkg.slide_parts()[0])
            stripped = Path(tmp) / "s.pptx"
            pkg.save(stripped)
            out, dropped = _drop_orphan_slide_rels(stripped, Path(tmp) / "o.pptx")
            self.assertGreaterEqual(len(dropped), 1)
            self.assertEqual({d["target"] for d in dropped}, {"slide2.xml"})


def build_fixture(path: Path) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Fixture title"
    s.shapes.title.name = "Title 1"
    body = s.placeholders[1]
    body.name = "Body"
    tf = body.text_frame
    for i, (t, lvl) in enumerate([("Point one", 0), ("Detail", 1), ("Point two", 0)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = t, lvl
    lab = s.shapes.add_textbox(Emu(500000), Emu(5800000), Emu(900000), Emu(300000))
    lab.name = "Tiny Label"
    lab.text_frame.text = "label"
    lab.text_frame.word_wrap = True
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Data"
    b2 = s2.shapes.add_textbox(Emu(500000), Emu(1500000), Emu(3000000), Emu(800000))
    b2.name = "Body"
    b2.text_frame.text = "slide two body"
    tbl = s2.shapes.add_table(2, 2, Emu(500000), Emu(2500000), Emu(4000000), Emu(800000)).table
    tbl.cell(0, 0).text = "h"
    tbl.cell(1, 0).merge(tbl.cell(1, 1))
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S", (1, 2))
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(5000000), Emu(2500000), Emu(3000000), Emu(2000000), cd)
    prs.save(path)


def _pm6_available() -> bool:
    try:
        pm6_root()
        return True
    except PM6Error:
        return False


@unittest.skipUnless(HAVE_PPTX and _pm6_available(), "set D6PPT_PM6_ROOT to a PPT Master v6.x checkout")
class RoundtripIntegrationTest(unittest.TestCase):
    """End-to-end: repeat a slide, rewrite a multi-paragraph body; check identity + package contract."""

    @classmethod
    def setUpClass(cls):
        from double6_ppt_cli.pm6_template import apply_plan
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        cls.deck = root / "fixture.pptx"
        build_fixture(cls.deck)
        cls.lib = analyze_pptx(cls.deck)
        body = next(s for s in cls.lib["slides"][0]["slots"] if s["name"] == "Body")
        cls.body_id = body["drawingml_id"]
        plan = scaffold_plan(cls.lib)
        s1 = dict(plan["slides"][0])
        s1["replacements"] = [{"slot_id": body["slot_id"], "text": "Alpha headline\n\tAlpha detail\nBeta headline\n\tBeta detail"}]
        s1b = dict(plan["slides"][0])
        s1b["replacements"] = [{"slot_id": body["slot_id"], "text": "Gamma\nDelta", "levels": [0, 0]}]
        s2 = dict(plan["slides"][1])
        t = cls.lib["slides"][1]
        s2["table_edits"] = [{"table_id": t["tables"][0]["table_id"], "cells": [{"row": 0, "col": 0, "text": "Header"}]}]
        s2["chart_edits"] = [{"chart_id": t["charts"][0]["chart_id"], "categories": ["X", "Y", "Z"], "series": [{"name": "New", "values": [3, 4, 5]}]}]
        cls.plan = dict(plan, slides=[s1, s2, s1b])
        cls.outputs = {}
        for mode in ("native", "svg", "svg-restore"):
            out = root / f"out-{mode}.pptx"
            cls.outputs[mode] = (out, apply_plan(cls.deck, cls.plan, out, mode=mode))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _body(self, mode, slide):
        pkg = Package.open(self.outputs[mode][0])
        return pkg, find_by_id(pkg.xml(pkg.slide_parts()[slide - 1]), self.body_id)

    def test_native_keeps_identity_paragraphs_and_placeholder(self):
        pkg, hits = self._body("native", 1)
        self.assertEqual(len(hits), 1)
        self.assertEqual(etree.QName(hits[0]).localname, "sp")
        self.assertIsNotNone(hits[0].find("./p:nvSpPr/p:nvPr/p:ph", NS))
        ps = hits[0].findall("./p:txBody/a:p", NS)
        self.assertEqual(len(ps), 4)
        self.assertEqual([p.find(q("a", "pPr")).get("lvl") if p.find(q("a", "pPr")) is not None else None for p in ps], [None, "1", None, "1"])
        _, hits3 = self._body("native", 3)
        ps3 = hits3[0].findall("./p:txBody/a:p", NS)
        self.assertEqual(["".join(p.xpath(".//a:t/text()", namespaces=NS)) for p in ps3], ["Gamma", "Delta"])

    def test_native_masters_layouts_theme_byte_identical(self):
        with zipfile.ZipFile(self.deck) as a, zipfile.ZipFile(self.outputs["native"][0]) as b:
            for name in a.namelist():
                if name.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/")):
                    self.assertEqual(hashlib.sha256(a.read(name)).digest(), hashlib.sha256(b.read(name)).digest(), name)

    def test_native_table_and_chart(self):
        pkg = Package.open(self.outputs["native"][0])
        root = pkg.xml(pkg.slide_parts()[1])
        self.assertIn("Header", "".join(root.xpath(".//a:tbl//a:t/text()", namespaces=NS)))
        lib = analyze_pptx(self.outputs["native"][0])
        self.assertEqual(lib["slides"][1]["charts"][0]["categories"], ["X", "Y", "Z"])

    def test_export_used_clone_passthrough(self):
        summary = self.outputs["native"][1]["export_summary"]
        self.assertIn("rebuilt=0", summary)

    def test_pure_svg_route_loses_multi_paragraph_identity(self):
        """Documents the upstream limitation this prototype works around (fails if upstream fixes it)."""
        _, hits = self._body("svg", 1)
        native_like = len(hits) == 1 and etree.QName(hits[0]).localname == "sp" and len(hits[0].findall("./p:txBody/a:p", NS)) == 4
        self.assertFalse(native_like, "upstream now keeps multi-paragraph frames native; revisit svg mode")

    def test_svg_restore_recovers_identity(self):
        _, hits = self._body("svg-restore", 1)
        self.assertEqual(len(hits), 1)
        self.assertIsNotNone(hits[0].find("./p:nvSpPr/p:nvPr/p:ph", NS))
        self.assertEqual(len(hits[0].findall("./p:txBody/a:p", NS)), 4)


class BackendSwitchTest(unittest.TestCase):
    def test_invalid_backend_rejected(self):
        from double6_ppt_cli.common import D6PPTError
        from double6_ppt_cli.template_workflow import template_backend
        old = os.environ.get("D6PPT_TEMPLATE_BACKEND")
        try:
            os.environ["D6PPT_TEMPLATE_BACKEND"] = "nope"
            with self.assertRaises(D6PPTError):
                template_backend()
            os.environ["D6PPT_TEMPLATE_BACKEND"] = "pm6"
            self.assertEqual(template_backend(), "pm6")
        finally:
            if old is None:
                os.environ.pop("D6PPT_TEMPLATE_BACKEND", None)
            else:
                os.environ["D6PPT_TEMPLATE_BACKEND"] = old


if __name__ == "__main__":
    unittest.main()
