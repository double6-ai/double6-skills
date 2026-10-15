"""Visual-quality gate: 8-dimension rubric + deterministic checks on the rendered deck.

The rubric is scored by the visual model (or a human) from the rendered pages; the automated
checks run on the PPTX and the rendered PNG pages from the current render manifest.
Colour expectations follow the selected upstream PPT Master style; the only colour checks here
are objective (text contrast, Office-default chart colours when a deck palette exists).
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree
from PIL import Image

RUBRIC_DIMENSIONS = {
    "hierarchy": "层级：标题/重点/正文主次清楚，一眼看到结论",
    "color": "配色：遵循所选 style 的 Color Behavior，颜色有角色、对比足够、与 deck 调色一致",
    "whitespace": "留白与平衡：留白有意图，无上重下空、无拥挤",
    "alignment": "对齐与网格：边距、栅格、基线一致",
    "imagery": "图像/图标：视觉元素服务含义，风格统一",
    "data_viz": "图表/表格：直接标注、突出重点、配色跟随 deck（无图表页按版式元素评）",
    "consistency": "一致性：跨页 motif、字号角色、页眉页脚一致",
    "polish": "完成度：无溢出、孤字、压线、错位，整体专业",
}
RUBRIC_PASS_MEAN = 3.5
RUBRIC_PASS_MIN = 3.0

OFFICE_DEFAULT_SERIES = {
    "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47",  # Office 2013+
    "4F81BD", "C0504D", "9BBB59", "8064A2", "4BACC6", "F79646",  # Office 2007-2010
    "156082", "E97132", "196B24", "0F9ED5", "A02B93", "4EA72E",  # Office 2023+
}
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
EMU_PER_PT = 12700
MAX_FONT_SIZES = 7


def evaluate_rubric(scores: dict[str, Any]) -> dict[str, Any]:
    missing = [d for d in RUBRIC_DIMENSIONS if d not in scores]
    unknown = [d for d in scores if d not in RUBRIC_DIMENSIONS]
    values = {}
    for dim in RUBRIC_DIMENSIONS:
        if dim in scores:
            v = float(scores[dim])
            if not 1.0 <= v <= 5.0:
                raise ValueError(f"rubric score out of range 1..5: {dim}={v}")
            values[dim] = v
    mean = round(sum(values.values()) / len(values), 3) if values else 0.0
    low = [d for d, v in values.items() if v < RUBRIC_PASS_MIN]
    passed = not missing and not unknown and mean >= RUBRIC_PASS_MEAN and not low
    return {"scores": values, "mean": mean, "below_min": low, "missing": missing, "unknown": unknown,
            "pass": passed, "thresholds": {"mean": RUBRIC_PASS_MEAN, "each_min": RUBRIC_PASS_MIN}}


# --------------------------------------------------------------------------- helpers

def _hex_to_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _lum(rgb: tuple[int, int, int]) -> float:
    def ch(c: float) -> float:
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = rgb[:3]
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast_ratio(a: tuple[int, int, int], b: tuple[int, int, int]) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _is_cjk(ch: str) -> bool:
    o = ord(ch)
    return 0x2E80 <= o <= 0x9FFF or 0xF900 <= o <= 0xFAFF or 0xFF00 <= o <= 0xFFEF or 0x3000 <= o <= 0x303F


CJK_LATIN_SPACE = re.compile(r"[\u3400-\u9fff\uff00-\uffef] +[0-9A-Za-z%$¥€]|[0-9A-Za-z%] +[\u3400-\u9fff]")
_TOKEN = re.compile(r"[\x21-\x7E\u2013\u2014\u2103\u00B7]+\s*|\s+|.")


def _last_line(text: str, font, avail: float) -> str | None:
    """Greedy line-break simulation (Latin words and digit runs never split); None if one line."""
    lines, cur = [], ""
    for tok in _TOKEN.findall(text):
        if cur and font.getlength((cur + tok).rstrip()) > avail:
            lines.append(cur)
            cur = tok.lstrip()
        else:
            cur += tok
    lines.append(cur)
    return lines[-1].strip() if len(lines) > 1 else None


def _is_orphan(text: str, tail: str) -> bool:
    if not tail:
        return False
    cjk_text = any(_is_cjk(c) for c in text)
    punct = "，。、；：！？）》」”’,.;:!?)"
    core = tail.strip(punct)
    if cjk_text:
        if core and all(_is_cjk(c) for c in core) and len(core) <= 1:
            return True
        if core and all(ord(c) < 0x2E80 for c in core) and len(core) <= 6:
            return True  # e.g. "4%", "Q4", "AI"
        return False
    words = tail.split()
    return len(words) == 1 and len(core) <= 6


def _lab(rgb: tuple[int, int, int]) -> tuple[float, float, float]:
    def f(c: float) -> float:
        c = c / 255
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (f(c) for c in rgb[:3])
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    g_ = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116  # noqa: E731
    return 116 * g_(y) - 16, 500 * (g_(x) - g_(y)), 200 * (g_(y) - g_(z))


def delta_e(a: str, b: str) -> float:
    la, lb = _lab(_hex_to_rgb(a)), _lab(_hex_to_rgb(b))
    return sum((p - q) ** 2 for p, q in zip(la, lb)) ** 0.5


MIN_SERIES_DELTA_E = 25.0


def _theme_colors(z: zipfile.ZipFile) -> dict[str, str]:
    out: dict[str, str] = {}
    names = [n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n)]
    if not names:
        return out
    root = etree.fromstring(z.read(sorted(names)[0]))
    scheme = root.find(".//a:clrScheme", NS)
    if scheme is None:
        return out
    for child in scheme:
        name = etree.QName(child).localname
        srgb = child.find("a:srgbClr", NS)
        sys_clr = child.find("a:sysClr", NS)
        if srgb is not None:
            out[name] = srgb.get("val", "").upper()
        elif sys_clr is not None:
            out[name] = (sys_clr.get("lastClr") or "").upper()
    return out


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # group
            yield from _iter_shapes(shape.shapes)


def _font_for(family: str | None, bold: bool, size_pt: float):
    from PIL import ImageFont
    try:
        from .pm6_template.textfit import font_file
        found = font_file(family or "Noto Sans CJK SC", True, bold) or font_file("Noto Sans CJK SC", True, bold)
    except Exception:  # pragma: no cover - font discovery is best effort
        found = None
    if not found:
        return None
    path, index = found
    try:
        return ImageFont.truetype(path, max(1, round(size_pt * 4)), index=index)  # 4x for precision
    except OSError:
        return None


# --------------------------------------------------------------------------- checks

MIN_BODY_PT = 14.0  # 18.7 px on the 1280x720 canvas
SMALL_TEXT_MIN_UNITS = 50  # width units (CJK = 2): footers, kickers and labels stay exempt
SMALL_TEXT_EXEMPT = re.compile(r"^(Photo|Source|Sources|Note|Notes|来源|图片|注)[:：\s]", re.I)


def compute_visual_metrics(pptx: Path, pages: list[Path], palette: dict[str, Any] | None = None) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(pptx))
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    findings: list[dict[str, Any]] = []
    per_slide: list[dict[str, Any]] = []
    deck_sizes: set[float] = set()

    def add(code: str, slide: int | None, message: str, severity: str = "error", **evidence: Any) -> None:
        findings.append({"code": code, "slide": slide, "severity": severity, "message": message, "evidence": evidence})

    images = [Image.open(p).convert("RGB") for p in pages] if pages else []

    for index, slide in enumerate(prs.slides, 1):
        non_text = 0
        img = images[index - 1] if index - 1 < len(images) else None
        for shape in _iter_shapes(slide.shapes):
            area = (int(shape.width or 0) * int(shape.height or 0)) / float(slide_w * slide_h)
            has_text = bool(getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip())
            if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False) or shape.shape_type == 13:
                non_text += 1
            elif not has_text and shape.shape_type != 6 and area < 0.95:
                xml = etree.tostring(shape._element).decode("utf-8", "ignore")
                if re.search(r"<a:(solidFill|gradFill|blipFill)|<a:ln[^>]*>\s*<a:solidFill|<a:custGeom|prst=\"(line|straightConnector1)\"", xml):
                    non_text += 1
            elif has_text:
                xml = etree.tostring(shape._element).decode("utf-8", "ignore")
                sppr = xml.split("<p:txBody", 1)[0]
                if re.search(r"<a:(solidFill|gradFill)", sppr) and area < 0.95:
                    non_text += 1  # filled card / band carrying text
            if not has_text:
                continue
            frame = shape.text_frame
            box = None
            if img is not None and shape.width and shape.height:
                sx, sy = img.width / slide_w, img.height / slide_h
                left, top = int(shape.left or 0), int(shape.top or 0)
                box = (max(0, int(left * sx)), max(0, int(top * sy)),
                       min(img.width, int((left + shape.width) * sx) + 1), min(img.height, int((top + shape.height) * sy) + 1))
            for para in frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    size = run.font.size.pt if run.font.size is not None else None
                    if size:
                        deck_sizes.add(round(size, 1))
                    color = None
                    try:
                        if run.font.color is not None and run.font.color.type is not None and run.font.color.rgb is not None:
                            color = str(run.font.color.rgb)
                    except (AttributeError, ValueError):
                        color = None
                    if color and box and box[2] > box[0] and box[3] > box[1]:
                        samples = []
                        bw = box[2] - box[0]
                        for seg in range(3):  # left / middle / right: catches gradients and photos behind text
                            sb = (box[0] + seg * bw // 3, box[1], box[0] + (seg + 1) * bw // 3, box[3])
                            if sb[2] - sb[0] < 2:
                                continue
                            crop = img.crop(sb).resize((max(1, (sb[2] - sb[0]) // 2), max(1, (sb[3] - sb[1]) // 2)))
                            cols = crop.getcolors(crop.width * crop.height) or []
                            # ignore pixels close to the text colour itself (glyph ink)
                            tr = _hex_to_rgb(color)
                            cols = [c for c in cols if sum(abs(c[1][k] - tr[k]) for k in range(3)) > 60] or cols
                            if cols:
                                samples.append(max(cols, key=lambda c: c[0])[1])
                        colors = samples
                        if colors:
                            bg = min(colors, key=lambda c: contrast_ratio(_hex_to_rgb(color), c))
                            ratio = contrast_ratio(_hex_to_rgb(color), bg)
                            large = (size or 0) >= 18 or ((size or 0) >= 14 and bool(run.font.bold))
                            need = 3.0 if large else 4.5
                            if ratio + 1e-6 < need:
                                add("low_contrast", index, f"Text '{run.text[:20]}' contrast {ratio:.2f}:1 < {need}:1",
                                    shape=shape.name, color=color, background="#%02X%02X%02X" % bg[:3], ratio=round(ratio, 2))
                        break  # one sample per paragraph is enough
                # Orphans: a wrapped paragraph whose last line is a lone CJK char, a short Latin/digit
                # fragment after CJK text ("4%"), or a single short word in Latin text.
                if text and frame.word_wrap is not False and shape.width:
                    run0 = next((r for r in para.runs if r.text.strip()), None)
                    size = (run0.font.size.pt if run0 is not None and run0.font.size is not None else None)
                    if size:
                        font = _font_for(run0.font.name, bool(run0.font.bold), size)
                        if font is not None:
                            body = frame._txBody.find(f"{{{NS['a']}}}bodyPr")
                            ins = sum(int(body.get(k)) if body is not None and body.get(k) else 91440 for k in ("lIns", "rIns"))
                            avail = (int(shape.width) - ins) / EMU_PER_PT * 4
                            tail = _last_line(text, font, avail)
                            if tail is not None and _is_orphan(text, tail):
                                add("orphan_line", index, f"Paragraph '{text[:16]}…' leaves '{tail}' alone on its last line",
                                    shape=shape.name, tail=tail)
                # Body copy must stay readable from the back of a room: long text below MIN_BODY_PT.
                if text and not SMALL_TEXT_EXEMPT.match(text.strip()):
                    run0 = next((r for r in para.runs if r.text.strip()), None)
                    psize = run0.font.size.pt if run0 is not None and run0.font.size is not None else None
                    units = sum(2 if _is_cjk(c) else 1 for c in text)
                    if psize and psize + 1e-6 < MIN_BODY_PT and units >= SMALL_TEXT_MIN_UNITS:
                        add("small_body_text", index, f"Body text '{text[:20]}…' is {psize:g}pt (< {MIN_BODY_PT:g}pt)",
                            shape=shape.name, size_pt=psize)
                # CJK typography: no inserted spaces between CJK and Latin/digits.
                if text and CJK_LATIN_SPACE.search(text):
                    add("cjk_latin_space", index, f"Space between CJK and Latin/digits in '{text[:24]}'", shape=shape.name)
        page_metrics: dict[str, Any] = {"slide": index, "non_text_elements": non_text}
        if non_text < 1:
            add("no_non_text_element", index, "Page has no non-text visual element (card, band, icon, chart, image, rule…)")
        if img is not None:
            small = img.resize((160, 90))
            px = list(small.getdata())
            corners = [px[0], px[159], px[-160], px[-1]]
            bg = max(set(corners), key=corners.count)
            ink = [(i % 160, i // 160) for i, c in enumerate(px) if sum(abs(c[k] - bg[k]) for k in range(3)) > 30]
            ink_ratio = len(ink) / len(px)
            page_metrics["ink_ratio"] = round(ink_ratio, 3)
            if ink:
                cy = sum(y for _, y in ink) / len(ink) / 90
                cx = sum(x for x, _ in ink) / len(ink) / 160
                bottom = sum(1 for _, y in ink if y >= 54) / len(ink)
                page_metrics.update(center_x=round(cx, 3), center_y=round(cy, 3), bottom_40pct_ink_share=round(bottom, 3))
                if bottom < 0.05 and ink_ratio < 0.5:
                    add("top_heavy", index, f"Bottom 40% of the page carries only {bottom:.0%} of the content (top-heavy layout)",
                        bottom_share=round(bottom, 3))
            if ink_ratio < 0.04:
                add("near_empty_page", index, f"Only {ink_ratio:.1%} of the page differs from the background", ink_ratio=round(ink_ratio, 3))
        per_slide.append(page_metrics)

    if len(deck_sizes) > MAX_FONT_SIZES:
        add("too_many_font_sizes", None, f"Deck uses {len(deck_sizes)} explicit font sizes (> {MAX_FONT_SIZES})",
            sizes=sorted(deck_sizes))

    with zipfile.ZipFile(pptx) as z:
        theme = _theme_colors(z)
        deck_palette = {str(v).lstrip("#").upper() for k, v in (palette or {}).items() if isinstance(v, str)}
        for v in (palette or {}).get("chart_series", []) or []:
            deck_palette.add(str(v).lstrip("#").upper())
        for name in sorted(n for n in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n)):
            root = etree.fromstring(z.read(name))
            series_colors = []
            for ser in root.iter(f"{{{NS['c']}}}ser"):
                sppr = ser.find("c:spPr", NS)
                if sppr is None:
                    series_colors.append(None)
                    continue
                srgb = sppr.find(".//a:solidFill/a:srgbClr", NS)
                scheme = sppr.find(".//a:solidFill/a:schemeClr", NS)
                if srgb is not None:
                    series_colors.append(srgb.get("val", "").upper())
                elif scheme is not None:
                    series_colors.append(theme.get(scheme.get("val", ""), "").upper() or None)
                else:
                    series_colors.append(None)
            resolved = [c if c else theme.get(f"accent{i + 1}", "") for i, c in enumerate(series_colors)]
            office = [c for c in resolved if c in OFFICE_DEFAULT_SERIES and c not in deck_palette]
            present = [c for c in resolved if c]
            close = [(x, y) for i, x in enumerate(present) for y in present[i + 1:] if x != y and delta_e(x, y) < MIN_SERIES_DELTA_E]
            same = len(present) != len(set(present)) and len(present) > 1
            if close or same:
                add("chart_series_indistinct", None, f"{name} series colours are too similar to tell apart {close or present}",
                    chart=name, colors=resolved)
            if palette and office:
                add("office_default_chart_colors", None, f"{name} uses Office default series colours {office} instead of the deck palette",
                    chart=name, colors=resolved)

    errors = [f for f in findings if f["severity"] == "error"]
    return {
        "status": "pass" if not errors else "fail",
        "error_count": len(errors),
        "warning_count": len(findings) - len(errors),
        "font_sizes": sorted(deck_sizes),
        "slides": per_slide,
        "findings": findings,
        "checks": ["low_contrast", "top_heavy", "near_empty_page", "no_non_text_element", "too_many_font_sizes",
                   "office_default_chart_colors", "orphan_line", "cjk_latin_space", "chart_series_indistinct", "small_body_text"],
    }
