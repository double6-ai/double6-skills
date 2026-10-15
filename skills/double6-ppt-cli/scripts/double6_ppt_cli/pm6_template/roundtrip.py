"""apply_plan on top of PPT Master v6.7 Edit-Native roundtrip.

Pipeline
  1. ``pptx_to_svg.py --roundtrip`` imports the template into a disposable workspace.
  2. ``page_plan.json`` expresses select / reorder / repeat (upstream clones slide parts; masters,
     layouts and untouched pages stay byte-identical).
  3. Text strategy (``D6PPT_PM6_TEXT_MODE``):
       native       (default) leave SVG text untouched so every object restores natively, then
                    rewrite the native txBody by shape id after export (ids, bullets, placeholders,
                    wrap and autofit stay native; fit decided with real font metrics).
       svg          pure upstream: write wrapped lines into the SVG and let the exporter rebuild
                    (evaluation baseline; multi-paragraph frames are split into groups upstream).
       svg-restore  svg + post-export restore: swap each rebuilt object back to the original native
                    element (same id) and apply the native txBody rewrite.
  4. ``svg_to_pptx.py --roundtrip`` exports; tables/charts are then filled natively by id.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from lxml import etree

from .library import CHART_RE, SLOT_RE, TABLE_RE
from .ooxml import NS, Package, cnvpr, find_by_id, iter_shapes, q  # noqa: F401
from .style import fit_context
from .textfit import Metrics, fit, map_levels
from .txbody import apply_fit, set_text
from .notes_transitions import apply_notes, apply_transition

SVG = "http://www.w3.org/2000/svg"


class PM6Error(RuntimeError):
    def __init__(self, message: str, code: str = "pm6_template_error", details: Any = None):
        super().__init__(message)
        self.code = code
        self.details = details


PM6_LOCKED_TAG = "v6.7.0"
PM6_LOCKED_COMMIT = "b4efe3ddf237a97a42880164759162f8e2412f2f"


def vendor_root() -> Path:
    return Path(__file__).resolve().parents[3] / "vendor" / "ppt-master-core"


def pm6_root() -> Path:
    """Vendored PPT Master core, version-locked through vendor/ppt-master-core/BOM.json.

    ``D6PPT_PM6_ROOT`` is a developer override only (e.g. to test a newer upstream); it is
    recorded in the receipt and never needed in normal use."""
    env = os.environ.get("D6PPT_PM6_ROOT")
    if env:
        c = Path(env)
        for cand in (c, c / "skills" / "ppt-master"):
            if (cand / "scripts" / "svg_to_pptx.py").is_file() and (cand / "scripts" / "pptx_to_svg.py").is_file():
                return cand
        raise PM6Error(f"D6PPT_PM6_ROOT does not contain PPT Master scripts: {env}", "pm6_missing")
    vroot = vendor_root()
    bom_path = vroot / "BOM.json"
    try:
        bom = json.loads(bom_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise PM6Error("vendored PPT Master BOM missing or unreadable", "pm6_missing", {"bom": str(bom_path)}) from exc
    if bom.get("tag") != PM6_LOCKED_TAG or bom.get("commit") != PM6_LOCKED_COMMIT:
        raise PM6Error("vendored PPT Master version does not match the template-fill lock", "pm6_version_mismatch",
                       {"expected": [PM6_LOCKED_TAG, PM6_LOCKED_COMMIT], "found": [bom.get("tag"), bom.get("commit")]})
    root = vroot / "skills" / "ppt-master"
    for name in ("pptx_to_svg.py", "svg_to_pptx.py"):
        if not (root / "scripts" / name).is_file():
            raise PM6Error(f"vendored PPT Master is missing scripts/{name}", "pm6_missing")
    return root


def _python() -> str:
    return os.environ.get("D6PPT_PM6_PYTHON") or sys.executable


def _run(args: list[str], cwd: Path, log: list[dict[str, Any]]) -> str:
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    proc = subprocess.run([_python(), *args], cwd=cwd, capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=900, env=env)
    log.append({"argv": args, "returncode": proc.returncode, "stdout_tail": proc.stdout[-4000:], "stderr_tail": proc.stderr[-4000:]})
    if proc.returncode != 0:
        raise PM6Error(f"PPT Master step failed: {args[0]}: {(proc.stderr or proc.stdout)[-800:]}", "pm6_step_failed", log[-1])
    return proc.stdout + proc.stderr


# ----------------------------------------------------------------------------- SVG mode helpers

def _svg_text_edit(svg_path: Path, shape_id: int, text: str, ctx: dict[str, Any] | None, levels: list[int] | None = None) -> dict[str, Any]:
    """Pure-upstream text edit: canonical positioned lines, width-aware breaking with real metrics."""
    tree = etree.parse(str(svg_path))
    root = tree.getroot()
    groups = root.xpath(f"//svg:g[@data-pptx-source-ref='slide:{shape_id}']", namespaces={"svg": SVG})
    if len(groups) != 1:
        return {"status": "svg_target_missing"}
    g = groups[0]
    texts = g.xpath(".//svg:text", namespaces={"svg": SVG})
    first = texts[0] if texts else None
    fx, fy, fw, fh = [float(v) for v in g.get("data-pptx-frame").split()]
    size = float(first.get("font-size", "18")) if first is not None else 18.0
    attrs = {k: v for k, v in (first.attrib.items() if first is not None else []) if k in ("font-size", "fill", "font-weight", "font-family")}
    anchor = first.get("text-anchor", "start") if first is not None else "start"
    for t in texts:
        t.getparent().remove(t)
    report = fit(ctx, text, levels=levels) if ctx else None
    scale = report["font_scale"] if report and report["status"] in ("autofit", "shrunk", "overflow") else 1.0
    lines: list[tuple[int, str]] = []
    if report:
        for (lvl, _), wrapped in zip(map_levels(ctx, text, levels), report["lines"]):
            lines.extend((lvl, w) for w in wrapped)
    else:
        lines = [(0, s) for s in text.split("\n")]
    px_per_pt = 4.0 / 3.0
    ins_l = (ctx["body"]["lIns"] if ctx else 7.2) * px_per_pt
    ins_t = (ctx["body"]["tIns"] if ctx else 3.6) * px_per_pt
    fs = size * scale
    # Upstream canonical form: ONE <text> per logical paragraph, positioned <tspan> rows (x + dy) per
    # wrapped line. A semantic shape accepts at most one paragraph-based text component, so a
    # multi-paragraph frame is expressed as separate paragraphs only by font-size/marker changes;
    # we therefore keep one <text> and put every wrapped line (all paragraphs) in positioned rows.
    node = etree.SubElement(g, f"{{{SVG}}}text")
    x0 = fx + ins_l if anchor == "start" else (fx + fw / 2 if anchor == "middle" else fx + fw - ins_l)
    node.set("x", f"{x0:.2f}")
    node.set("y", f"{fy + ins_t + fs:.2f}")
    node.set("text-anchor", anchor)
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for k, v in attrs.items():
        node.set(k, v)
    node.set("font-size", f"{fs:.2f}")
    for idx, (lvl, line) in enumerate(lines):
        mar = (ctx["levels"].get(str(lvl), ctx["levels"]["0"])["marL"] * px_per_pt) if ctx and anchor == "start" else 0
        row = etree.SubElement(node, f"{{{SVG}}}tspan")
        row.set("x", f"{x0 + mar:.2f}")
        row.set("dy", "0" if idx == 0 else f"{fs * 1.2:.2f}")
        row.text = line or " "
    tree.write(str(svg_path), xml_declaration=False, encoding="utf-8")
    return {"status": "svg_edited", "lines": len(lines), "font_scale": scale, "fit": report and report["status"]}


def _svg_href_edit(svg_path: Path, mapping: dict[str, str]) -> None:
    raw = svg_path.read_text(encoding="utf-8")
    for old, new in mapping.items():
        raw = raw.replace(f'href="{old}" data-pptx-source-href="{old}"', f'href="{new}"')
    svg_path.write_text(raw, encoding="utf-8")


# ----------------------------------------------------------------------------- native fill

def _fill_text(pkg: Package, slide_part: str, element, text: str, receipts: list, base: dict, levels: list[int] | None = None) -> None:
    ctx = fit_context(pkg, slide_part, element)
    report = fit(ctx, text, levels=levels) if ctx else {"status": "unknown", "font_scale": 1.0}
    body = element.find(q("p", "txBody"))
    sizes = None
    scale = None
    if report["status"] in ("shrunk",) or (report["status"] == "overflow" and report.get("autofit_mode") != "norm"):
        scale = report["font_scale"]
        sizes = {int(k): v["sz"] for k, v in ctx["levels"].items()}
    set_text(body, text, ctx["template_paragraph_levels"] if ctx else None, size_scale=scale, resolved_sizes=sizes, levels=levels)
    apply_fit(element, report)
    receipts.append({**base, "fit": report["status"], "font_scale": report.get("font_scale"), "ratio": report.get("ratio"),
                     "line_count": report.get("line_count"), "need_h_pt": report.get("need_h_pt"), "avail_h_pt": report.get("avail_h_pt"),
                     "engine": report.get("engine")})


def _fill_table(frame, edit: dict[str, Any]) -> int:
    rows = frame.findall(".//a:tbl/a:tr", NS)
    n = 0
    for cell in edit.get("cells", []):
        tc = rows[int(cell["row"])].findall(q("a", "tc"))[int(cell["col"])]
        body = tc.find(q("a", "txBody"))
        if body is None:
            body = etree.SubElement(tc, q("a", "txBody"))
            etree.SubElement(body, q("a", "bodyPr"))
            etree.SubElement(body, q("a", "lstStyle"))
        set_text(body, str(cell.get("text", "")))
        n += 1
    return n


def _fill_charts(output: Path, chart_jobs: list[tuple[int, int, dict[str, Any]]]) -> list[dict[str, Any]]:
    """Use python-pptx only to compute new chart + workbook parts, then splice just those parts."""
    if not chart_jobs:
        return []
    from pptx import Presentation  # local import: optional dependency of this path
    from pptx.chart.data import CategoryChartData

    prs = Presentation(str(output))
    receipts = []
    for slide_no, shape_id, edit in chart_jobs:
        slide = prs.slides[slide_no - 1]
        shape = next((s for s in slide.shapes if s.shape_id == shape_id and s.has_chart), None)
        if shape is None:
            raise PM6Error(f"chart {shape_id} not found on output slide {slide_no}", "chart_target_missing")
        data = CategoryChartData()
        data.categories = list(edit["categories"])
        for s in edit["series"]:
            data.add_series(s["name"], list(s["values"]))
        shape.chart.replace_data(data)
        receipts.append({"plan_slide": slide_no, "drawingml_id": shape_id, "chart_part": str(shape.chart.part.partname)})
    tmp = output.with_suffix(".charts.pptx")
    prs.save(str(tmp))
    src = Package.open(output)
    new = Package.open(tmp)
    for name, blob in new.entries.items():
        if (name.startswith("ppt/charts/") or name.startswith("ppt/embeddings/")) and src.entries.get(name) != blob:
            src.entries[name] = blob
    if new.entries["[Content_Types].xml"] != src.entries["[Content_Types].xml"]:
        src.entries["[Content_Types].xml"] = new.entries["[Content_Types].xml"]
    src.save(output)
    tmp.unlink()
    return receipts


def _restore_native(pkg: Package, template: Package, slide_part: str, source_part: str, shape_id: int) -> str:
    """svg-restore: replace whatever upstream rebuilt for ``shape_id`` with the original native element."""
    out_root = pkg.xml(slide_part)
    src_matches = find_by_id(template.xml(source_part), shape_id)
    if len(src_matches) != 1:
        return "source_missing"
    original = src_matches[0]
    name = cnvpr(original).get("name")
    hits = find_by_id(out_root, shape_id)
    if len(hits) == 1:
        # upstream kept the id but rebuilt the element (new spPr/pPr, placeholder link dropped)
        target = hits[0]
    else:
        named = [e for e in iter_shapes(out_root) if cnvpr(e) is not None and cnvpr(e).get("name") == name]
        if len(named) != 1:
            return "rebuilt_target_ambiguous"
        target = named[0]
    target.getparent().replace(target, deepcopy(original))
    pkg.touch(slide_part)
    return "restored"


def _drop_orphan_slide_rels(template: Path, dest: Path) -> tuple[Path, list[dict[str, str]]]:
    """Remove slide->slide relationships no longer referenced by the slide XML.

    The Double6 wrapper strips confirmed navigation actions (hlinkClick) before the backend runs but
    leaves their relationships; PM6 page-plan cloning refuses any slide relationship that points at
    a repeated slide. Unreferenced relationships carry no behaviour, so dropping them is lossless.
    """
    pkg = Package.open(template)
    dropped = []
    for part in pkg.slide_parts():
        rels_part = part.rsplit("/", 1)[0] + "/_rels/" + part.rsplit("/", 1)[1] + ".rels"
        if rels_part not in pkg.entries:
            continue
        xml = pkg.entries[part].decode("utf-8")
        root = pkg.xml(rels_part)
        for rel in list(root):
            if (rel.get("Type") or "").endswith("/slide") and f'"{rel.get("Id")}"' not in xml:
                root.remove(rel)
                dropped.append({"slide_part": part, "relationship_id": rel.get("Id"), "target": rel.get("Target")})
                pkg.touch(rels_part)
    if not dropped:
        return template, []
    pkg.save(dest)
    return dest, dropped


# ----------------------------------------------------------------------------- selectors

def _resolve_target(tpl: Package, src_part: str, item: dict[str, Any], kind: str) -> int | None:
    """v4.8 selector forms: ``slot_id``/``table_id``/``chart_id``, else ``shape_id``, else unique ``shape_name``."""
    regex, key = {"slot": (SLOT_RE, "slot_id"), "table": (TABLE_RE, "table_id"), "chart": (CHART_RE, "chart_id")}[kind]
    if item.get(key):
        m = regex.fullmatch(str(item[key]))
        return int(m.group(2)) if m else None
    if item.get("shape_id") not in (None, ""):
        try:
            return int(item["shape_id"])
        except (TypeError, ValueError):
            return None
    if item.get("shape_name"):
        hits = [cnvpr(e) for e in iter_shapes(tpl.xml(src_part)) if cnvpr(e) is not None and cnvpr(e).get("name") == item["shape_name"]]
        if len(hits) == 1:
            return int(hits[0].get("id"))
    return None


def _rep_text(rep: dict[str, Any]) -> str:
    if isinstance(rep.get("paragraphs"), list):
        return "\n".join(str(x) for x in rep["paragraphs"])
    return str(rep.get("text") if rep.get("text") is not None else "")


# ----------------------------------------------------------------------------- entry point

def apply_plan(template: Path, plan: dict[str, Any], output: Path, transition: str = "keep",
               transition_duration: float = 0.5, *, mode: str | None = None, keep_workspace: Path | None = None) -> dict[str, Any]:
    mode = mode or os.environ.get("D6PPT_PM6_TEXT_MODE", "native")
    if mode not in ("native", "svg", "svg-restore"):
        raise PM6Error(f"unknown text mode {mode}", "pm6_mode_invalid")
    root = pm6_root()
    scripts = root / "scripts"
    template = Path(template).resolve()
    output = Path(output).resolve()
    log: list[dict[str, Any]] = []
    slides = plan.get("slides") or []
    tpl = Package.open(template)
    src_parts = tpl.slide_parts()
    receipt: dict[str, Any] = {"schema": "double6_pm6_apply_receipt.v1", "mode": mode, "pm6_root": str(root),
                               "text": [], "tables": [], "charts": [], "svg_edits": [], "restored": [],
                               "notes": [], "transitions": []}
    with tempfile.TemporaryDirectory(prefix="d6-pm6-") as tmp:
        ws = Path(tmp) / "ws"
        source_input, dropped = _drop_orphan_slide_rels(template, Path(tmp) / "input.pptx")
        receipt["dropped_orphan_slide_relationships"] = dropped
        _run([str(scripts / "pptx_to_svg.py"), str(source_input), "-o", str(ws), "--inheritance-mode", "both", "--roundtrip"], Path(tmp), log)
        svg_dir = ws / "authoring-svg-flat"
        pages, used = [], set()
        for i, item in enumerate(slides, 1):
            src = int(item["source_slide"])
            if not 1 <= src <= len(src_parts):
                raise PM6Error(f"source_slide {src} out of range", "source_slide_out_of_range")
            canonical = f"slide_{src:02d}.svg"
            if src in used:
                name = f"slide_{src:02d}_p{i:02d}.svg"
                shutil.copy(svg_dir / canonical, svg_dir / name)
                pages.append({"source_slide": src, "svg": name})
            else:
                used.add(src)
                pages.append({"source_slide": src, "svg": canonical})
        identity = [p["source_slide"] for p in pages] == list(range(1, len(src_parts) + 1))
        if not identity:
            plan_pages = [{"source_slide": p["source_slide"], **({"svg": p["svg"]} if p["svg"] != f"slide_{p['source_slide']:02d}.svg" else {})} for p in pages]
            (ws / "page_plan.json").write_text(json.dumps({"schema": "ppt-master.roundtrip-page-plan.v1", "pages": plan_pages}, indent=1), encoding="utf-8")
        if mode in ("svg", "svg-restore"):
            for i, (item, page) in enumerate(zip(slides, pages), 1):
                src_part = src_parts[page["source_slide"] - 1]
                for rep in item.get("replacements", []):
                    sid = _resolve_target(tpl, src_part, rep, "slot")
                    if sid is None:
                        continue
                    if _rep_text(rep) == str(rep.get("old_text", "\0")):
                        continue
                    hits = find_by_id(tpl.xml(src_part), sid)
                    ctx = fit_context(tpl, src_part, hits[0]) if len(hits) == 1 and hits[0].find(q("p", "txBody")) is not None else None
                    r = _svg_text_edit(svg_dir / page["svg"], sid, _rep_text(rep), ctx, rep.get("levels"))
                    receipt["svg_edits"].append({"plan_slide": i, "drawingml_id": sid, **r})
        export_args = [str(scripts / "svg_to_pptx.py"), str(ws), "--roundtrip"]
        out_text = _run(export_args, Path(tmp), log)
        summary = re.search(r"Round-trip export summary: (.*)", out_text)
        receipt["export_summary"] = summary.group(1).strip() if summary else None
        qg = re.search(r"quality_gate=(\S+)", out_text)
        receipt["quality_gate"] = qg.group(1) if qg else None
        exports = sorted((ws / "exports").glob("*.pptx"), key=lambda p: p.stat().st_mtime)
        if not exports:
            raise PM6Error("roundtrip export produced no pptx", "pm6_export_missing")
        exported = exports[-1]
        if keep_workspace:
            shutil.rmtree(keep_workspace, ignore_errors=True)
            shutil.copytree(ws, keep_workspace)
        pkg = Package.open(exported)
        out_parts = pkg.slide_parts()
        if len(out_parts) != len(slides):
            raise PM6Error("output slide count differs from plan", "pm6_slide_count_mismatch")
        chart_jobs = []
        auto_advance = False
        for i, (item, part) in enumerate(zip(slides, out_parts), 1):
            src_part = src_parts[int(item["source_slide"]) - 1]
            root_xml = pkg.xml(part)
            for rep in item.get("replacements", []):
                sid = _resolve_target(tpl, src_part, rep, "slot")
                if sid is None:
                    raise PM6Error(f"replacement target not found on plan slide {i}", "replacement_target_not_found",
                                   {k: rep.get(k) for k in ("slot_id", "shape_id", "shape_name")})
                text = _rep_text(rep)
                slot_id = rep.get("slot_id") or f"s{int(item['source_slide']):02d}_sh{sid}"
                base = {"plan_slide": i, "slot_id": slot_id, "drawingml_id": sid}
                if mode == "svg":
                    continue
                if mode == "svg-restore" and text != str(rep.get("old_text", "\0")):
                    receipt["restored"].append({**base, "result": _restore_native(pkg, tpl, part, src_part, sid)})
                    root_xml = pkg.xml(part)
                hits = [h for h in find_by_id(root_xml, sid) if h.tag == q("p", "sp")]
                if not hits and text == str(rep.get("old_text", "\0")):
                    receipt["text"].append({**base, "fit": "unchanged"})
                    continue
                if len(hits) != 1:
                    raise PM6Error(f"slot {slot_id} not unique on output slide {i}", "pm6_slot_identity_lost", {"hits": len(hits)})
                if text == str(rep.get("old_text", "\0")) and "old_text" in rep:
                    receipt["text"].append({**base, "fit": "unchanged"})
                    continue
                _fill_text(pkg, part, hits[0], text, receipt["text"], base, rep.get("levels"))
                pkg.touch(part)
            for edit in item.get("table_edits", []) or []:
                tid = _resolve_target(tpl, src_part, edit, "table")
                hits = find_by_id(root_xml, tid) if tid is not None else []
                if len(hits) != 1:
                    raise PM6Error(f"table {edit.get('table_id')} not found on output slide {i}", "pm6_table_identity_lost")
                receipt["tables"].append({"plan_slide": i, "table_id": edit.get("table_id") or f"s{int(item['source_slide']):02d}_tbl{tid}", "cells": _fill_table(hits[0], edit)})
                pkg.touch(part)
            for edit in item.get("chart_edits", []) or []:
                cid = _resolve_target(tpl, src_part, edit, "chart")
                if cid is not None:
                    chart_jobs.append((i, cid, edit))
            receipt["notes"].append({"plan_slide": i, **apply_notes(pkg, part, i, item)})
            tr = apply_transition(pkg, part, item, transition, transition_duration)
            auto_advance = auto_advance or tr.get("auto_advance", False)
            receipt["transitions"].append({"plan_slide": i, **tr})
        if auto_advance:
            from pptx_transitions import set_package_use_timings
            for part_name in list(pkg.dirty):
                pkg.save_part(part_name)
            set_package_use_timings(pkg.entries)
            pkg._xml.clear()
        output.parent.mkdir(parents=True, exist_ok=True)
        pkg.save(output)
        receipt["charts"] = _fill_charts(output, chart_jobs)
        from pptx_transitions import validate_pptx_transition_package
        try:
            validate_pptx_transition_package(output, require_use_timings=auto_advance)
        except ValueError as exc:
            raise PM6Error(f"transition package validation failed: {exc}", "pm6_transition_invalid") from exc
    receipt["log"] = [{k: v for k, v in row.items() if k != "stdout_tail"} | {"stdout_tail": row["stdout_tail"][-600:]} for row in log]
    output.with_name(output.stem + ".pm6_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=1), encoding="utf-8")
    return receipt
